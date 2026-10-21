"""PPTX内へ文章役割を保存する。図形名と表示文字列には手を加えない。"""
from pptx.oxml.ns import qn


ROLES = frozenset({"natural", "body", "compact"})
_EXTENSION = "{8AF73BB6-5E31-4F0F-A399-3D98F3B95FD9}"
_TAG = "{urn:pptxdsl:text-layout:v1}textLayout"


def _extension(text_frame):
    extensions = text_frame._txBody.bodyPr.find(qn("a:extLst"))
    if extensions is not None:
        for extension in extensions:
            if extension.get("uri") == _EXTENSION:
                return extension
    return None


def set_text_role(text_frame, role):
    if role not in ROLES:
        raise ValueError(f"未対応の文章役割です: {role}")
    body = text_frame._txBody.bodyPr
    extension = _extension(text_frame)
    if extension is None:
        extensions = body.find(qn("a:extLst"))
        if extensions is None:
            extensions = body.makeelement(qn("a:extLst"), {})
            body.append(extensions)
        extension = body.makeelement(qn("a:ext"), {"uri": _EXTENSION})
        extensions.append(extension)
    layout = extension.find(_TAG)
    if layout is None:
        layout = body.makeelement(_TAG, {})
        extension.append(layout)
    layout.set("role", role)


def text_role(text_frame, default="body"):
    extension = _extension(text_frame)
    layout = extension.find(_TAG) if extension is not None else None
    if layout is None:
        return default
    role = layout.get("role")
    if role not in ROLES:
        raise ValueError("PPTXに保存された文章役割が不正です")
    return role
