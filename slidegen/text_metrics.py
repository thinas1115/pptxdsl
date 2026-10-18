"""保存済みPPTXのrun・段落・余白から文字の行ボックスを推定する。

Officeの描画エンジンではない。対応フォント以外を黙って別書体で測定しない。
"""
from dataclasses import dataclass
from itertools import groupby
import math

from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn

from slidegen.font_config import FontConfigError, font_family
from slidegen.text_roles import text_role
from slidegen.textfit import _BREAK_AFTER, _BREAK_BEFORE, _is_protected_token, _natural_tokens
from slidegen.textfit import boundary_spacing_in, text_width_in, width_safety, wrapper_for_role


EMU = 914400


class TextMetricsError(ValueError):
    """保存書式を本プロジェクトの測定契約で再現できない場合。"""


@dataclass(frozen=True)
class Style:
    size: float
    weight: str
    family: str


@dataclass(frozen=True)
class Character:
    text: str
    style: Style


@dataclass(frozen=True)
class TextLine:
    text: str
    rect: tuple[float, float, float, float]
    advance: float
    paragraph: int
    max_size: float


@dataclass(frozen=True)
class TextLayout:
    lines: tuple[TextLine, ...]
    height: float
    inner: tuple[float, float, float, float]
    role: str

    @property
    def bounds(self):
        rects = [line.rect for line in self.lines if line.text.strip()]
        if not rects:
            return None
        return (min(r[0] for r in rects), min(r[1] for r in rects),
                max(r[2] for r in rects), max(r[3] for r in rects))


def _property(elements, attribute, default=None):
    return next((element.get(attribute) for element in elements
                 if element is not None and element.get(attribute) is not None), default)


def _default_properties(tf, paragraph):
    """通常図形の段落→text bodyのリストstyle→masterのotherStyleを読む。"""
    p_pr = paragraph._p.pPr
    level = paragraph.level + 1
    properties = [p_pr]
    style = tf._txBody.find(qn("a:lstStyle"))
    if style is not None:
        properties.extend((style.find(qn(f"a:lvl{level}pPr")), style.find(qn("a:defPPr"))))
    master = tf.part.slide.slide_layout.slide_master
    tx_styles = master.element.find(qn("p:txStyles"))
    if tx_styles is not None:
        other = tx_styles.find(qn("p:otherStyle"))
        if other is not None:
            properties.append(other.find(qn(f"a:lvl{level}pPr")))
    return [element for element in properties if element is not None]


def _style(r_pr, properties, char, scale):
    defaults = [p.find(qn("a:defRPr")) for p in properties]
    elements = [r_pr, *defaults]
    size = float(_property(elements, "sz", "1800")) / 100 * scale
    bold = _property(elements, "b", "0") in {"1", "true"}
    tag = "a:ea" if ord(char) >= 0x2E80 else "a:latin"
    family = None
    for element in elements:
        face = element.find(qn(tag)) if element is not None else None
        if face is not None and face.get("typeface"):
            family = face.get("typeface")
            break
    if family is None:
        # 明示書体のない他社PPTXを、実行OSの既定書体で測定したことにはしない。
        raise TextMetricsError("書体を明示していない文字があります。対応書体を指定してから検査してください")
    weight = "bold" if bold else "regular"
    if family in {"Consolas", "Noto Sans Mono CJK JP"}:
        weight = "code"
        if bold:
            raise TextMetricsError("コード書体の太字は測定対象外です")
    expected = font_family(weight)
    if family != expected:
        raise TextMetricsError(f"保存書体{family!r}と測定書体{expected!r}が一致しません。対応する同じ書体で検査してください")
    if not math.isfinite(size) or size <= 0:
        raise TextMetricsError("文字サイズは正数にしてください")
    return Style(size, weight, family)


def _characters(paragraph, properties, scale):
    result = [[]]
    for element in paragraph._p:
        if element.tag not in {qn("a:r"), qn("a:fld"), qn("a:br")}:
            continue
        if element.tag == qn("a:br"):
            result.append([])
            continue
        text = element.find(qn("a:t"))
        if text is None:
            continue
        r_pr = element.find(qn("a:rPr"))
        for char in text.text or "":
            result[-1].append(Character(char, _style(r_pr, properties, char, scale)))
    return result


def _width(chars):
    groups = [(style, list(group)) for style, group in groupby(chars, key=lambda char: char.style)]
    width = sum(text_width_in("".join(char.text for char in group), style.size, style.weight)
                for style, group in groups)
    for (left, a), (right, b) in zip(groups, groups[1:]):
        width += boundary_spacing_in(a[-1].text + b[0].text, max(left.size, right.size), left.weight) * width_safety(left.weight)
    return width


def _strip(chars, *, left=False):
    result = list(chars)
    while result and (result[0] if left else result[-1]).text.isspace():
        result.pop(0 if left else -1)
    return result


def _wrap_mixed(chars, width, role, first_width=None):
    """run境界を維持した幅でトークン配置する。大きいrunを全文へ代用しない。"""
    plain = "".join(char.text for char in chars)
    tokens, offset = [], 0
    for text in _natural_tokens(plain):
        tokens.append(chars[offset:offset + len(text)])
        offset += len(text)
    lines, current = [], []
    def available():
        return first_width if not lines and first_width is not None else width
    for token in tokens:
        if current and _width(current + token) > available():
            lines.append(_strip(current))
            current = _strip(token, left=True)
        else:
            current += token
        if current and _width(current) > available():
            text = "".join(char.text for char in current)
            if not _is_protected_token(text):
                part = []
                for char in current:
                    if part and _width(part + [char]) > available() and char.text not in _BREAK_BEFORE:
                        if part[-1].text in _BREAK_AFTER and len(part) > 1:
                            lines.append(part[:-1])
                            part = part[-1:]
                        else:
                            lines.append(part)
                            part = []
                    part.append(char)
                current = part
    if current or not lines:
        lines.append(_strip(current))
    # 文章役割ごとの最終行比率を、runの実測幅でも維持する。
    ratio = {"natural": 0.28, "body": 0.18, "compact": 0.0}[role]
    if len(lines) >= 2 and _width(lines[-1]) < width * ratio:
        joined = lines[-2] + lines[-1]
        text = "".join(char.text for char in joined)
        boundaries, position = [], 0
        for token in _natural_tokens(text):
            position += len(token)
            boundaries.append(position)
        for split in reversed(boundaries):
            left, right = _strip(joined[:split]), _strip(joined[split:], left=True)
            if not left or not right or right[0].text in _BREAK_BEFORE or left[-1].text in _BREAK_AFTER:
                continue
            left_width = first_width if len(lines) == 2 and first_width is not None else width
            if _width(left) <= left_width and width * ratio <= _width(right) <= width:
                lines[-2:] = [left, right]
                break
    return lines


def _wrap(chars, width, role, enabled, first_width=None):
    if not enabled or not chars:
        return [chars]
    styles = {char.style for char in chars}
    if len(styles) != 1 or first_width not in (None, width):
        return _wrap_mixed(chars, width, role, first_width)
    style = chars[0].style
    plain = "".join(char.text for char in chars)
    # 均一書式ではrendererと全く同じ役割別wrapperを使用する。
    texts = wrapper_for_role(role)(plain, width, style.size, style.weight)
    return [[Character(char, style) for char in text] for text in texts]


def _spacing(properties, tag, size, default):
    for prop in properties:
        setting = prop.find(qn(tag))
        if setting is None:
            continue
        points = setting.find(qn("a:spcPts"))
        if points is not None:
            value = float(points.get("val")) / 100 / 72
            if not math.isfinite(value) or value < 0:
                raise TextMetricsError("段落間隔は有限の非負数にしてください")
            return value
        percent = setting.find(qn("a:spcPct"))
        if percent is not None:
            value = size / 72 * float(percent.get("val")) / 100000
            if not math.isfinite(value) or value < 0:
                raise TextMetricsError("段落間隔は有限の非負数にしてください")
            return value
    return default


def _layout_text_frame(tf, box, *, box_is_inner=False, default_role="body"):
    """boxは通常図形の外形、表セルだけは既にmarginを除いた内側を渡す。"""
    try:
        role = text_role(tf, default_role)
    except ValueError as error:
        raise TextMetricsError(str(error)) from None
    inner = box if box_is_inner else (box[0] + tf.margin_left / EMU, box[1] + tf.margin_top / EMU,
                                     box[2] - tf.margin_right / EMU, box[3] - tf.margin_bottom / EMU)
    if not all(math.isfinite(value) for value in inner):
        raise TextMetricsError("文字領域の位置と余白は有限の値にしてください")
    if inner[2] <= inner[0] or inner[3] <= inner[1]:
        raise TextMetricsError("文字領域の幅または高さが余白より小さくなっています")
    body = tf._txBody.bodyPr
    autofit = body.find(qn("a:normAutofit"))
    scale = float(autofit.get("fontScale", "100000")) / 100000 if autofit is not None else 1.0
    reduction = float(autofit.get("lnSpcReduction", "0")) / 100000 if autofit is not None else 0.0
    if not 0 < scale <= 1 or not 0 <= reduction < 1:
        raise TextMetricsError("自動縮小の倍率または行間削減率が不正です")
    pending, cursor = [], 0.0
    for pi, paragraph in enumerate(tf.paragraphs):
        properties = _default_properties(tf, paragraph)
        indent_left = float(_property(properties, "marL", "0")) / EMU
        indent_right = float(_property(properties, "marR", "0")) / EMU
        indent = float(_property(properties, "indent", "0")) / EMU
        if not all(math.isfinite(value) for value in (indent_left, indent_right, indent)):
            raise TextMetricsError("段落の字下げは有限の値にしてください")
        width = inner[2] - inner[0] - indent_left - indent_right
        if width <= 0:
            raise TextMetricsError("段落の字下げ後に文字領域が残りません")
        paragraphs = _characters(paragraph, properties, scale)
        available_chars = [char for line in paragraphs for char in line]
        default_fonts = [prop.find(qn("a:defRPr")) for prop in properties]
        inherited_size = float(_property(default_fonts, "sz", "1100")) / 100 * scale
        default_size = max((char.style.size for char in available_chars), default=inherited_size)
        if not math.isfinite(default_size) or default_size <= 0:
            raise TextMetricsError("空段落を含む文字サイズは正数にしてください")
        before = _spacing(properties, "a:spcBef", default_size, 0.0)
        after = _spacing(properties, "a:spcAft", default_size, 0.0)
        cursor += before
        alignment = _property(properties, "algn", "l")
        if alignment not in {"l", "ctr", "r"}:
            raise TextMetricsError(f"段落配置{alignment!r}は測定対象外です")
        first = True
        for manual in paragraphs:
            first_width = width - (indent if first else 0)
            if first_width <= 0:
                raise TextMetricsError("最初の行の字下げ後に文字領域が残りません")
            for chars in _wrap(manual, width, role, tf.word_wrap is not False, first_width):
                size = max((char.style.size for char in chars), default=default_size)
                advance = _spacing(properties, "a:lnSpc", size, size / 72) * (1 - reduction)
                if not math.isfinite(advance) or advance < 0:
                    raise TextMetricsError("行間は有限の非負数にしてください")
                leading_left = indent_left + (indent if first else 0)
                line_width = _width(chars)
                usable = width - (indent if first else 0)
                x = inner[0] + leading_left
                if alignment == "ctr":
                    x += (usable - line_width) / 2
                elif alignment == "r":
                    x += usable - line_width
                # 行送りと描画行の高さを分ける。最終行の後ろの余白を衝突矩形にしない。
                pending.append(("".join(char.text for char in chars), x,
                                cursor + max(0.0, advance - size / 72) / 2, line_width,
                                size / 72, advance, pi, size))
                cursor += advance
                first = False
        cursor += after
    top = inner[1]
    if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
        top += (inner[3] - inner[1] - cursor) / 2
    elif tf.vertical_anchor == MSO_ANCHOR.BOTTOM:
        top += inner[3] - inner[1] - cursor
    elif tf.vertical_anchor not in (None, MSO_ANCHOR.TOP):
        raise TextMetricsError("縦方向の均等配置は測定対象外です")
    lines = tuple(TextLine(text, (x, top + y, x + width, top + y + height), advance, pi, size)
                  for text, x, y, width, height, advance, pi, size in pending)
    return TextLayout(lines, cursor, inner, role)


def layout_text_frame(tf, box, *, box_is_inner=False, default_role="body"):
    """保存値の異常を診断結果へ変換し、対応書体の設定エラーとは区別する。"""
    try:
        return _layout_text_frame(tf, box, box_is_inner=box_is_inner, default_role=default_role)
    except (TextMetricsError, FontConfigError):
        raise
    except (AttributeError, KeyError, TypeError, ValueError, OverflowError):
        raise TextMetricsError("PPTXの文字書式を読み取れません。サイズ・行間・余白を確認してください") from None
