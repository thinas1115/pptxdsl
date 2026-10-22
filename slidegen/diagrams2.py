"""図解系スライド第2弾: プロセスタイムライン・工程表・2軸マップ。"""

from math import hypot
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from slidegen.generate import (ACCENT, CORAL, GRAY, LIGHT, NAVY, RULE, SURFACE, TEXT, WHITE,
                      ZEBRA, ContentArea, add_rect, add_text, header, note_line)
from slidegen.diagrams import add_arrow
from slidegen.diagrams3 import plain_line, route
from slidegen.layout_fit import FitError, ensure_within, fit_text_or_raise, select_fit
from slidegen.textfit import line_height_in, text_width_in
from slidegen.timeline_layout import plan_program_roadmap
from slidegen.layout_settings import ROADMAP

PROGRAM_LINE_PT = ROADMAP.line_pt
MATRIX_LABEL_DISTANCE = 0.45


def _rects_overlap(a, b, gap=0.04):
    return not (a[2] + gap <= b[0] or b[2] + gap <= a[0]
                or a[3] + gap <= b[1] or b[3] + gap <= a[1])


def _matrix_label_positions(
        points, ox, top, aw, ah, *, label_w=2.0, gap=0.04, barriers=()):
    """点ラベルを周囲8方向へ自動配置し、衝突しない矩形を返す。"""
    point_xy = [(ox + point["x"] * aw, top + (1 - point["y"]) * ah)
                for point in points]
    point_boxes = [(x - 0.17, y - 0.17, x + 0.17, y + 0.17)
                   for x, y in point_xy]
    placed = {}
    leaders = {}
    order = sorted(
        range(len(points)),
        key=lambda i: (
            -sum(abs(point_xy[i][0] - x) < 2.2 and abs(point_xy[i][1] - y) < 0.7
                 for x, y in point_xy),
            -int(bool(points[i].get("emph"))),
            i,
        ),
    )
    label_h = 0.31
    for index in order:
        px, py = point_xy[index]
        manual = (points[index].get("lx"), points[index].get("ly"))
        weight = "bold" if points[index].get("emph") else "regular"
        size = _fit_single_line("matrix", "points.name", points[index]["name"], label_w - 0.08, 11, 9,
                                bold=bool(points[index].get("emph")))
        width = max(0.40, min(label_w, text_width_in(points[index]["name"], size, weight) + 0.08))
        candidates = [
            (-width / 2, -label_h - 0.24), (-width / 2, 0.24),
            (0.32, -label_h / 2), (-width - 0.32, -label_h / 2),
            (0.26, -label_h - 0.20), (-width - 0.26, -label_h - 0.20),
            (0.26, 0.20), (-width - 0.26, 0.20),
        ]
        if manual[0] is not None or manual[1] is not None:
            # 非公開の旧内部呼び出しだけは従来の枠幅を維持する。JSONでは指定不可。
            width = label_w
            candidates = [(-label_w / 2 + (manual[0] or 0.0),
                           manual[1] if manual[1] is not None else -0.36)]
        for dx, dy in candidates:
            rect = (px + dx, py + dy, px + dx + width, py + dy + label_h)
            inside = (rect[0] >= ox + 0.04 and rect[2] <= ox + aw - 0.04
                      and rect[1] >= top + 0.04 and rect[3] <= top + ah - 0.04)
            if not inside:
                continue
            if any(rect[0] < barrier < rect[2] for barrier in barriers[:1]):
                continue
            if any(rect[1] < barrier < rect[3] for barrier in barriers[1:2]):
                continue
            if any(_rects_overlap(rect, other, gap)
                   for other_index, other in enumerate(point_boxes)
                   if other_index != index or manual == (None, None)):
                continue
            if any(_rects_overlap(rect, other, gap) for other in placed.values()):
                continue
            nearest = _matrix_nearest_edge((px, py), rect)
            if hypot(nearest[0] - px, nearest[1] - py) > MATRIX_LABEL_DISTANCE:
                continue
            leader = _matrix_leader_box((px, py), rect)
            if leader and (any(_rects_overlap(leader, box, 0.015) for box in placed.values())
                           or any(_rects_overlap(leader, box, 0.015)
                                  for j, box in enumerate(point_boxes) if j != index)
                           or any(_rects_overlap(leader, box, 0.015) for box in leaders.values())):
                continue
            if any(_rects_overlap(rect, box, 0.015) for box in leaders.values()):
                continue
            placed[index] = rect
            if leader:
                leaders[index] = leader
            break
        else:
            raise FitError(
                f"matrix.points[{index}]: {points[index]['name']!r} のラベルを"
                "他の点と重ならずに配置できません。点を減らすか座標の粒度を"
                "見直してください。")
    return point_xy, placed


def _matrix_nearest_edge(point, rect):
    """点からラベル枠への最短接続点。図形を生成せず計算する。"""
    return (min(max(point[0], rect[0]), rect[2]),
            min(max(point[1], rect[1]), rect[3]))


def _matrix_leader_box(point, rect):
    """引出線の短い通路。矩形で保守的に他の点・文字・引出線との交差を避ける。"""
    ex, ey = _matrix_nearest_edge(point, rect)
    distance = hypot(ex - point[0], ey - point[1])
    if distance <= 0.26:
        return None
    ux, uy = (ex - point[0]) / distance, (ey - point[1]) / distance
    sx, sy = point[0] + ux * 0.17, point[1] + uy * 0.17
    tx, ty = ex - ux * 0.025, ey - uy * 0.025
    return min(sx, tx), min(sy, ty), max(sx, tx), max(sy, ty)


def fit_matrix_labels(points, ox, top, aw, ah, *, barriers=()):
    """ラベル間隔を圧縮し、その後ラベル幅を縮めて配置を試す。"""
    last_error = None
    for stage, width, gap in (
            ("standard", 2.0, 0.05),
            ("gap", 2.0, 0.015),
            ("element", 1.65, 0.015)):
        try:
            point_xy, labels = _matrix_label_positions(
                points, ox, top, aw, ah, label_w=width, gap=gap,
                barriers=barriers)
            return stage, width, point_xy, labels
        except FitError as exc:
            last_error = exc
    raise FitError(
        f"matrix: 最小ラベル幅でも衝突を解消できません。{last_error}")


def _grid_line(slide, x1, y1, x2, y2):
    """表の罫線を本文より先に描き、後続のラベルマスクを有効にする。"""
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = RULE
    line.line.width = Pt(0.6)
    line.shadow.inherit = False
    line.name = "roadmap:grid"
    return line


def _activity_line(slide, x1, y, x2, color):
    """工程表の作業期間を描く。矩形ではなく固定pt幅で見え方を揃える。"""
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2), Inches(y))
    line.line.color.rgb = color
    line.line.width = Pt(PROGRAM_LINE_PT)
    line.shadow.inherit = False
    return line


def _fit_single_line(renderer, field, text, width, max_pt, min_pt, *, bold=False):
    """工程表ラベルを折り返さず、幅へ収まるサイズまで縮小する。"""
    size = max_pt
    weight = "bold" if bold else "regular"
    while size >= min_pt - 0.01:
        if text_width_in(text, size, weight) <= width:
            return size
        size -= 0.5
    raise FitError(
        f"{renderer}.{field}: 1行ラベルが最小フォント{min_pt:g}ptでも"
        f"幅{width:.2f}inに収まりません。文言を短くするか期間を広げてください。"
    )


# ---- 番号付きプロセスタイムライン ----
def _process_attribute(step):
    attribute = step.get("attribute")
    if attribute:
        return attribute["label"], attribute["value"]
    if step.get("actor"):
        return "OWNER", step["actor"]
    return None


def s_process(slide, spec, page):
    if "flow" in spec:
        return _s_process_flow(slide, spec, page)
    area = header(slide, spec["kicker"], spec["title"], spec.get("lead"))
    if spec.get("note") and area.shifted:
        area = ContentArea(area.top, area.bottom - 0.30, area.shifted)
    y = area.map_y
    steps = spec["steps"]
    n = len(steps)
    if not 3 <= n <= 6:
        raise FitError(
            "process: 工程は3〜6件までです。工程を統合または分割してください。")
    left, usable_w = 0.78, 11.72
    w = usable_w / n
    line_y = y(2.42)
    add_rect(slide, left + w / 2, line_y, usable_w - w, 0.035, RULE)
    for i, st in enumerate(steps):
        x = left + i * w
        cx = x + w / 2
        emph = i in spec.get("emph", [])
        color = ACCENT if emph else NAVY
        sp = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(cx - 0.31), Inches(line_y - 0.29),
            Inches(0.62), Inches(0.62))
        sp.fill.solid()
        sp.fill.fore_color.rgb = color
        sp.line.fill.background()
        sp.shadow.inherit = False
        add_text(slide, cx - 0.31, line_y - 0.2, 0.62, 0.32, f"{i + 1:02d}",
                 11.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        name_size, name_lines = fit_text_or_raise(
            "process", f"steps[{i}].name", st["name"], w - 0.2, 0.42,
            15.5, min_pt=12.5, weight="bold", spacing=1.1)
        name_box = add_text(slide, x + 0.1, y(2.94), w - 0.2, 0.42,
                 "\n".join(name_lines), name_size,
                 bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        name_box.name = f"process:name:{i}"
        name_box.text_frame.word_wrap = False
        attribute = _process_attribute(st)
        desc_h = 1.35 if attribute else 1.95
        size, desc_lines = fit_text_or_raise(
            "process", f"steps[{i}].desc", st["desc"],
            # LibreOffice 26.2はwrap=Falseでも和欧混在行を再折り返す場合がある。
            # 計画幅と実際の文字枠の間に余裕を残し、説明が補足欄へ伸びるのを防ぐ。
            w - 0.34 - 0.12, desc_h, 13.5, min_pt=11.5, spacing=1.2)
        desc_h = len(desc_lines) * line_height_in(size, 1.2) + 0.03
        desc_box = add_text(slide, x + 0.17, y(3.56), w - 0.34, desc_h,
                 "\n".join(desc_lines), size,
                 color=TEXT, align=PP_ALIGN.CENTER, spacing=1.2)
        desc_box.name = f"process:desc:{i}"
        desc_box.text_frame.word_wrap = False
        if attribute:
            label, value = attribute
            attribute_y = y(3.56) + desc_h + 0.20
            label_size, label_lines = fit_text_or_raise(
                "process", f"steps[{i}].attribute.label", label,
                w - 0.3, 0.22, 10.5, min_pt=9.5,
                weight="bold", spacing=1.05)
            label_box = add_text(slide, x + 0.15, attribute_y, w - 0.3, 0.22,
                     "\n".join(label_lines), label_size, bold=True, color=TEXT,
                     align=PP_ALIGN.CENTER)
            label_box.name = f"process:attribute-label:{i}"
            label_box.text_frame.word_wrap = False
            value_size, value_lines = fit_text_or_raise(
                "process", f"steps[{i}].attribute.value", value,
                w - 0.3, 0.3, 11.5, min_pt=9.5, weight="bold", spacing=1.1)
            value_box = add_text(slide, x + 0.15, attribute_y + 0.27, w - 0.3, 0.3,
                     "\n".join(value_lines), value_size,
                     bold=True, color=color, align=PP_ALIGN.CENTER)
            value_box.name = f"process:attribute-value:{i}"
            value_box.text_frame.word_wrap = False
            ensure_within("process.attribute", attribute_y + 0.57 - area.top,
                          area.height, guidance="工程説明または補足を短くしてください。")
    ensure_within(
        "process", y(5.86) - area.top, area.height,
        guidance="工程説明を短くしてください。")
    if spec.get("note"):
        note_line(slide, spec["note"])


def _fit_process_flow(area_h, levels):
    max_rows = max(len(level) for level in levels)

    def candidates():
        values = [
            ("standard", 1.08, 0.36, 12.5),
            ("gap", 1.02, 0.22, 12.5),
            ("element", 0.88, 0.18, 11.0),
        ]
        for stage, node_h, gap_y, font in values:
            used = max_rows * node_h + max(0, max_rows - 1) * gap_y + 0.42
            yield stage, {"node_h": node_h, "gap_y": gap_y, "font": font}, used

    return select_fit(
        "process.flow", area_h, candidates(),
        guidance="工程ノードを減らすか、分岐フローを複数スライドへ分割してください。",
    )


def _s_process_flow(slide, spec, page):
    """levelsとedgesから分岐・合流・戻り線を持つ工程図を描画する。"""
    area = header(slide, spec["kicker"], spec["title"], spec.get("lead"))
    if spec.get("note") and area.shifted:
        area = ContentArea(area.top, area.bottom - 0.30, area.shifted)
    flow = spec["flow"]
    nodes, levels, edges = flow["nodes"], flow["levels"], flow["edges"]
    fitted = _fit_process_flow(area.height, levels)
    node_h = fitted.values["node_h"]
    gap_y = fitted.values["gap_y"]
    body_size = fitted.values["font"]
    left, usable_w = 0.72, 11.90
    level_count = len(levels)
    # 分岐ラベルを箱に重ねず置けるよう、列数に応じて必要な間隔を確保する。
    gap_x = 0.90 if level_count <= 5 else 0.52
    node_w = (usable_w - gap_x * (level_count - 1)) / level_count
    if node_w < 1.55:
        raise FitError(
            f"process.flow: {level_count}列では箱幅が{node_w:.2f}inとなり"
            "最小1.55inを下回ります。工程階層を減らしてください。")

    max_rows = max(len(level) for level in levels)
    rows_h = max_rows * node_h + max(0, max_rows - 1) * gap_y
    feedback_reserve = 0.64 if any(
        edge.get("kind") == "feedback" for edge in edges) else 0.18
    free_h = max(0.0, area.height - rows_h - feedback_reserve)
    top_pad = min(0.72, max(0.42, free_h * 0.30))

    rects = {}
    for level_index, level in enumerate(levels):
        top = area.top + top_pad
        x = left + level_index * (node_w + gap_x)
        for row_index, node_id in enumerate(level):
            y = top + row_index * (node_h + gap_y)
            rects[node_id] = (x, y, node_w, node_h)

    feedback_count = 0
    routed = []
    max_node_bottom = max(y + h for _x, y, _w, h in rects.values())
    level_of = {node_id: index for index, level in enumerate(levels)
                for node_id in level}
    for edge in edges:
        source, target = edge["from"], edge["to"]
        sx, sy, sw, sh = rects[source]
        tx, ty, tw, th = rects[target]
        if level_of[target] > level_of[source]:
            start = (sx + sw, sy + sh / 2)
            end = (tx, ty + th / 2)
            trunk = (start[0] + end[0]) / 2
            points = [start, (trunk, start[1]), (trunk, end[1]), end]
            if abs(start[1] - end[1]) < 0.03:
                points = [start, end]
        else:
            lane_y = max_node_bottom + 0.42 + feedback_count * 0.16
            if lane_y > area.bottom - 0.16:
                raise FitError(
                    "process.flow: 差戻し線を本文領域へ配置できません。"
                    "分岐数を減らすか工程を複数スライドへ分割してください。")
            feedback_count += 1
            start = (sx + sw / 2, sy + sh)
            end = (tx + tw / 2, ty + th)
            points = [start, (start[0], lane_y), (end[0], lane_y), end]
        route(slide, points, dash="dash" if edge.get("kind") == "feedback" else None,
              width=1.25)
        routed.append((edge, points))

    for node_id, (x, y, w, h) in rects.items():
        node = nodes[node_id]
        style = node.get("style", "standard")
        fill = LIGHT if style == "decision" else SURFACE
        line = ACCENT if style in {"accent", "decision"} else RULE
        add_rect(slide, x, y, w, h, fill, line=line)
        title_size, _ = fit_text_or_raise(
            "process.flow", f"nodes.{node_id}.name", node["name"],
            w - 0.28, 0.32, 14, min_pt=11.5, weight="bold", spacing=1.1)
        add_text(slide, x + 0.14, y + 0.13, w - 0.28, 0.32, node["name"],
                 title_size, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        desc = node.get("desc")
        if desc:
            desc_h = 0.34 if node.get("actor") else h - 0.52
            desc_size, lines = fit_text_or_raise(
                "process.flow", f"nodes.{node_id}.desc", desc,
                w - 0.30, desc_h, body_size, min_pt=9.5, spacing=1.12)
            add_text(slide, x + 0.15, y + 0.50, w - 0.30, desc_h,
                     "\n".join(lines), desc_size, color=TEXT,
                     align=PP_ALIGN.CENTER, spacing=1.12)
        if node.get("actor"):
            actor_size, _ = fit_text_or_raise(
                "process.flow", f"nodes.{node_id}.actor", node["actor"],
                w - 0.30, 0.24, 9.5, min_pt=8.0, weight="bold", spacing=1.05)
            add_text(slide, x + 0.15, y + h - 0.27, w - 0.30, 0.22,
                     node["actor"], actor_size, bold=True, color=ACCENT,
                     align=PP_ALIGN.CENTER)

    placed_labels = []
    for edge_index, (edge, points) in enumerate(routed):
        if not edge.get("label"):
            continue
        if edge.get("kind") == "feedback":
            segments = list(zip(points[:-1], points[1:]))
            start, end = max(
                segments,
                key=lambda segment: abs(segment[1][0] - segment[0][0])
                + abs(segment[1][1] - segment[0][1]),
            )
            label_x, label_y, label_w = (
                (start[0] + end[0]) / 2, (start[1] + end[1]) / 2 + 0.20, 1.35)
        else:
            sx, _sy, sw, _sh = rects[edge["from"]]
            tx, _ty, _tw, _th = rects[edge["to"]]
            end = points[-1]
            label_w = min(1.20, tx - (sx + sw) - 0.06)
            # 縦の分岐線が延びていない側へ避け、接続経路を一切マスクしない。
            label_x = (sx + sw + tx) / 2
            label_y = end[1] + (0.20 if end[1] > points[0][1] + 0.03 else -0.20)
        size = _fit_single_line("process.flow", f"edges[{edge_index}].label",
                                edge["label"], label_w, 10.5, 9.5)
        label_w = text_width_in(edge["label"], size) + 0.02
        label_h = line_height_in(size, 1.1) + 0.03
        bounds = (label_x - label_w / 2, label_y - label_h / 2,
                  label_x + label_w / 2, label_y + label_h / 2)
        obstacles = [(x, yy, x + w, yy + h) for x, yy, w, h in rects.values()]
        obstacles.extend(placed_labels)
        for _other, path in routed:
            obstacles.extend((min(a[0], b[0]), min(a[1], b[1]),
                              max(a[0], b[0]), max(a[1], b[1]))
                             for a, b in zip(path, path[1:]))
        if (bounds[1] < area.top or bounds[3] > area.bottom
                or any(_rects_overlap(bounds, obstacle, 0.025) for obstacle in obstacles)):
            raise FitError("process.flow: 分岐条件を線や工程に重ねず配置できません。"
                           "条件を短くするかフローを分割してください。")
        label_box = add_text(slide, bounds[0], bounds[1], label_w, label_h,
                             edge["label"], size, color=TEXT, spacing=1.1,
                             align=PP_ALIGN.CENTER, role="compact")
        label_box.name = f"process:condition:{edge_index}"
        label_box.text_frame.word_wrap = False
        placed_labels.append(bounds)
    if spec.get("note"):
        note_line(slide, spec["note"])


# ---- 複数テーマ・複数作業のプログラム工程表 ----
def s_program_roadmap(slide, spec, page):
    area = header(slide, spec["kicker"], spec["title"], spec.get("lead"))
    plan = plan_program_roadmap(spec, area)
    _draw_program_roadmap(slide, spec, plan)
    if spec.get("note"):
        note_line(slide, spec["note"])


def _draw_program_roadmap(slide, spec, plan):
    periods, tracks = spec["periods"], spec["tracks"]
    values = plan.fitted.values
    label_x, track_col_w, grid_w = ROADMAP.label_x, ROADMAP.track_width, ROADMAP.grid_width
    grid_x, period_w = plan.grid_x, plan.period_width
    top, header_h, grid_h = plan.top, values["header_h"], plan.grid_height
    add_rect(slide, label_x, top, track_col_w, header_h, NAVY)
    add_text(slide, label_x + 0.18, top + 0.07, track_col_w - 0.36, 0.28,
             "テーマ", values["period_pt"], bold=True, color=WHITE)
    add_rect(slide, grid_x, top, grid_w, header_h, NAVY)
    for index, period in enumerate(periods):
        size, _ = fit_text_or_raise(
            "program_roadmap", f"periods[{index}]", period,
            period_w - 0.04, 0.28, values["period_pt"],
            min_pt=ROADMAP.period_min_pt, weight="bold", spacing=1.0)
        period_box = add_text(slide, grid_x + index * period_w, top + 0.07,
                 period_w, 0.28, period, size, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, spacing=1.0)
        period_box.name = f"roadmap:period:{index}"

    for index, (row_top, row_h) in enumerate(zip(plan.row_tops, plan.row_heights)):
        add_rect(slide, label_x, row_top, track_col_w + grid_w, row_h,
                 SURFACE if index % 2 == 0 else ZEBRA)
    body_top = top + header_h
    for index in range(len(periods) + 1):
        _grid_line(slide, grid_x + index * period_w, body_top,
                   grid_x + index * period_w, top + grid_h)
    for index, row_top in enumerate(plan.row_tops):
        _grid_line(slide, label_x, row_top,
                   label_x + track_col_w + grid_w, row_top)
        track = tracks[index]
        row_h = plan.row_heights[index]
        has_goal = bool(track.get("goal"))
        number_y = row_top + (0.07 if has_goal else 0.11)
        name_y = row_top + (0.05 if has_goal else 0.09)
        add_text(slide, label_x + 0.12, number_y, 0.38, 0.25,
                 f"{index + 1:02d}", values["track_pt"], bold=True,
                 color=ACCENT)
        track_size = _fit_single_line(
            "program_roadmap", f"tracks[{index}].name", track["name"],
            track_col_w - 0.66, values["track_pt"], ROADMAP.track_min_pt, bold=True)
        track_box = add_text(slide, label_x + 0.52, name_y,
                 track_col_w - 0.66, min(0.28, row_h - 0.14),
                 track["name"], track_size, bold=True, color=NAVY,
                 spacing=1.0, wrap=False)
        track_box.name = f"roadmap:track:{index}"
        if has_goal:
            goal_size = _fit_single_line(
                "program_roadmap", f"tracks[{index}].goal", track["goal"],
                track_col_w - 0.66, min(values["activity_pt"], 9.5), ROADMAP.activity_min_pt)
            add_text(slide, label_x + 0.52, row_top + 0.29,
                     track_col_w - 0.66, 0.18, track["goal"], goal_size,
                     color=TEXT, spacing=1.0, wrap=False)

        for activity_plan in plan.activities[index]:
            placement = activity_plan.placement
            activity = placement.activity
            lane_top = activity_plan.lane_top
            x1, x2 = activity_plan.x1, activity_plan.x2
            color = CORAL if activity.get("emph") else ACCENT
            activity_line = _activity_line(slide, x1, lane_top + 0.035, x2, color)
            activity_line.name = f"roadmap:span:{index}:{placement.index}"
            text_h = values["lane_pitch"] - 0.07
            activity_label_x = activity_plan.label_x
            activity_label_w = activity_plan.label_width
            activity_size = _fit_single_line(
                "program_roadmap",
                f"tracks[{index}].activities[{placement.index}].label",
                activity["label"], activity_label_w - 0.04,
                values["activity_pt"], ROADMAP.activity_min_pt,
                bold=bool(activity.get("emph")))
            label = add_text(
                slide, activity_label_x + 0.02, lane_top + 0.07,
                activity_label_w - 0.04, text_h, activity["label"], activity_size,
                bold=bool(activity.get("emph")), color=TEXT,
                align=PP_ALIGN.CENTER, spacing=1.0, wrap=False)
            label.fill.solid()
            label.fill.fore_color.rgb = SURFACE if index % 2 == 0 else ZEBRA
            label.name = f"roadmap:activity:{index}:{placement.index}"
        milestone = track.get("milestone")
        if milestone:
            marker = plan.milestones[index]
            marker_x, marker_y = marker.x, marker.y
            diamond_size = 0.15
            diamond = slide.shapes.add_shape(
                MSO_SHAPE.DIAMOND,
                Inches(marker_x - diamond_size / 2),
                Inches(marker_y - diamond_size / 2),
                Inches(diamond_size), Inches(diamond_size))
            diamond.fill.solid()
            diamond.fill.fore_color.rgb = CORAL
            diamond.line.fill.background()
            diamond.shadow.inherit = False
            diamond.name = f"roadmap:milestone:{index}"
            milestone_w, milestone_x = marker.label_width, marker.label_x
            milestone_size = _fit_single_line(
                "program_roadmap", f"tracks[{index}].milestone.label",
                milestone["label"], milestone_w - 0.04,
                min(values["activity_pt"], 9.5), ROADMAP.activity_min_pt)
            milestone_label = add_text(
                slide, milestone_x + 0.02, marker_y + 0.07,
                milestone_w - 0.04, values["lane_pitch"] - 0.07,
                milestone["label"], milestone_size, color=TEXT,
                align=PP_ALIGN.CENTER, spacing=1.0, wrap=False)
            milestone_label.fill.solid()
            milestone_label.fill.fore_color.rgb = (
                SURFACE if index % 2 == 0 else ZEBRA)
            milestone_label.name = f"roadmap:milestone-label:{index}"
    _grid_line(slide, label_x, plan.bottom,
               label_x + track_col_w + grid_w, plan.bottom)


# ---- 2軸ポジショニングマップ ----
def s_matrix(slide, spec, page):
    area = header(slide, spec["kicker"], spec["title"], spec.get("lead"))
    if spec.get("note") and area.shifted:
        area = ContentArea(area.top, area.bottom - 0.30, area.shifted)
    y = area.map_y
    if not 1 <= len(spec["points"]) <= 8:
        raise FitError(
            "matrix: 点は1〜8件までです。点をまとめるかスライドを分割してください。")
    ox, oy = 1.68, y(6.24)
    aw, ah = 10.5, oy - y(2.22)
    mid_x, mid_y = ox + aw / 2, oy - ah / 2

    quadrants = spec.get("quadrants")
    if quadrants:
        # Explicit region names make the highlighted quadrant semantic. Existing
        # matrix specs without them retain the neutral scatter-plot treatment.
        add_rect(slide, ox, oy - ah, aw, ah, SURFACE, line=RULE)
        add_rect(slide, mid_x, oy - ah, aw / 2, ah / 2, LIGHT)
        qlabels = [
            (ox + 0.18, oy - 0.42, 2.2, quadrants[0], PP_ALIGN.LEFT, GRAY),
            (ox + aw - 2.38, oy - 0.42, 2.2, quadrants[1], PP_ALIGN.RIGHT, GRAY),
            (ox + 0.18, oy - ah + 0.18, 2.2, quadrants[2], PP_ALIGN.LEFT, GRAY),
            (ox + aw - 2.38, oy - ah + 0.18, 2.2, quadrants[3], PP_ALIGN.RIGHT, ACCENT),
        ]
        for idx, (x, y, w, label, align, color) in enumerate(qlabels):
            label_size, _ = fit_text_or_raise(
                "matrix", f"quadrants[{idx}]", label, w, 0.3, 10.5,
                min_pt=8.5, weight="bold", spacing=1.1)
            add_text(slide, x, y, w, 0.3, label, label_size, bold=True,
                     color=color, align=align)
    else:
        target_size, _ = fit_text_or_raise(
            "matrix", "target_label", spec["target_label"], 2.0, 0.3, 10.5,
            min_pt=8.5, weight="bold", spacing=1.1)
        add_text(slide, mid_x + 0.18, oy - ah + 0.16, 2.0, 0.3,
                 spec["target_label"], target_size, bold=True, color=ACCENT)

    add_arrow(slide, ox, oy, ox + aw, oy, width=1.75)
    add_arrow(slide, ox, oy, ox, oy - ah, width=1.75)
    x_axis_size, _ = fit_text_or_raise(
        "matrix", "x_axis", spec["x_axis"], 2.4, 0.3, 11.5,
        min_pt=9.5, weight="bold", spacing=1.1)
    y_axis_size, _ = fit_text_or_raise(
        "matrix", "y_axis", spec["y_axis"], 3.0, 0.3, 11.5,
        min_pt=9.5, weight="bold", spacing=1.1)
    add_text(slide, ox + aw - 2.5, oy + 0.16, 2.4, 0.3, spec["x_axis"], x_axis_size,
             bold=True, color=TEXT, align=PP_ALIGN.RIGHT)
    add_text(slide, ox - 0.02, oy - ah - 0.38, 3.0, 0.3, spec["y_axis"], y_axis_size,
             bold=True, color=TEXT)
    _stage, label_w, point_xy, labels = fit_matrix_labels(
        spec["points"], ox, oy - ah, aw, ah,
        barriers=(mid_x, mid_y) if quadrants else ())
    for index, (p, (px, py)) in enumerate(zip(spec["points"], point_xy)):
        r = 0.15
        sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(px - r), Inches(py - r),
                                    Inches(2 * r), Inches(2 * r))
        sp.fill.solid()
        sp.fill.fore_color.rgb = CORAL if p.get("emph") else ACCENT
        sp.line.color.rgb = WHITE
        sp.line.width = Pt(1.0)
        sp.shadow.inherit = False
        sp.name = f"matrix:point:{index}"
    for index, p in enumerate(spec["points"]):
        left, top, right, bottom = labels[index]
        point_size = _fit_single_line("matrix", "points.name", p["name"], right - left - 0.06,
                                      11, 9, bold=bool(p.get("emph")))
        px, py = point_xy[index]
        ex, ey = _matrix_nearest_edge((px, py), labels[index])
        distance = hypot(ex - px, ey - py)
        if distance > 0.26:
            ux, uy = (ex - px) / distance, (ey - py) / distance
            leader = plain_line(slide, px + ux * 0.17, py + uy * 0.17,
                                ex - ux * 0.025, ey - uy * 0.025,
                                color=GRAY, width=0.7)
            leader.name = f"matrix:leader:{index}"
        label = add_text(slide, left, top, right - left, bottom - top,
                 p["name"], point_size,
                 bold=bool(p.get("emph")), color=CORAL if p.get("emph") else TEXT,
                 align=PP_ALIGN.CENTER, wrap=False)
        label.name = f"matrix:label:{index}"
    ensure_within(
        "matrix", oy - area.top, area.height,
        guidance="軸ラベルまたは点の名称を短くしてください。")
    if spec.get("note"):
        note_line(slide, spec["note"])
