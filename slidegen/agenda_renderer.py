"""章名・補足・開始ページを読み順に揃える目次renderer。"""

from dataclasses import dataclass
from typing import Any, Mapping, Sequence

from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

from slidegen.generate import ACCENT, GRAY, NAVY, RULE, add_rect, add_text, header
from slidegen.layout_fit import FitError, PayloadFitResult, select_payload_fit, stepped
from slidegen.textfit import text_width_in, wrap_body, wrap_natural

TEXT_X, TEXT_W = 1.42, 10.04
TOP_GAP, DETAIL_GAP, LINE_SPACING = 0.14, 0.04, 1.6


@dataclass(frozen=True)
class AgendaRow:
    title: tuple[str, ...]
    detail: tuple[str, ...]
    page: int | None
    height: float


@dataclass(frozen=True)
class AgendaPlan:
    title_size: float
    detail_size: float
    gap: float
    rows: tuple[AgendaRow, ...]


def plan_agenda(
    items: Sequence[Mapping[str, Any]], available: float
) -> PayloadFitResult[AgendaPlan]:
    candidates = []
    profiles = [("standard", 20.0, 13.0, 0.22), ("gap", 20.0, 13.0, 0.08)]
    profiles += [("element", size, max(11.0, size - 7), 0.08) for size in stepped(19.5, 15, 0.5)]
    for stage, title_size, detail_size, gap in profiles:
        rows = []
        for item in items:
            title = tuple(wrap_natural(item["title"], TEXT_W, title_size, "bold"))
            detail = (
                tuple(wrap_body(item["detail"], TEXT_W, detail_size)) if "detail" in item else ()
            )
            if any(
                text_width_in(line, title_size, "bold") > TEXT_W + 0.01 for line in title
            ) or any(text_width_in(line, detail_size) > TEXT_W + 0.01 for line in detail):
                break
            height = len(title) * title_size * LINE_SPACING / 72
            if detail:
                height += DETAIL_GAP + len(detail) * detail_size * LINE_SPACING / 72
            rows.append(AgendaRow(title, detail, item.get("page"), height))
        if len(rows) == len(items):
            plan = AgendaPlan(title_size, detail_size, gap, tuple(rows))
            candidates.append(
                (stage, plan, TOP_GAP + sum(row.height for row in rows) + gap * (len(rows) - 1))
            )
    if not candidates:
        raise FitError(
            "agenda: 最小設定でも章名が収まりません。文言を短くするか、目次を分割してください。"
        )
    return select_payload_fit(
        "agenda",
        available,
        candidates,
        guidance="章名・補足を短くするか、目次を複数ページへ分けてください。",
    )


def s_agenda(slide: Any, spec: Mapping[str, Any], page: int) -> None:
    area = header(slide, spec["kicker"], spec["title"], spec.get("lead"))
    plan = plan_agenda(spec["items"], area.height).values
    y = area.top + TOP_GAP

    def text(
        x: float,
        top: float,
        width: float,
        height: float,
        value: str,
        size: float,
        *,
        bold: bool = False,
        align: Any = PP_ALIGN.LEFT,
    ) -> Any:
        shape = add_text(
            slide,
            x,
            top,
            width,
            height,
            value,
            size,
            bold=bold,
            color=NAVY if bold else GRAY,
            spacing=Pt(size * LINE_SPACING),
            align=align,
            wrap=False,
        )
        for paragraph in shape.text_frame.paragraphs:
            paragraph.space_before = paragraph.space_after = Pt(0)
        return shape

    for i, row in enumerate(plan.rows):
        number = text(0.76, y, 0.48, 0.36, f"{i + 1:02d}", 13, bold=True)
        number.text_frame.paragraphs[0].runs[0].font.color.rgb = ACCENT
        title_h = len(row.title) * plan.title_size * LINE_SPACING / 72
        text(TEXT_X, y, TEXT_W, title_h, "\n".join(row.title), plan.title_size, bold=True)
        if row.detail:
            text(
                TEXT_X,
                y + title_h + DETAIL_GAP,
                TEXT_W,
                len(row.detail) * plan.detail_size * LINE_SPACING / 72,
                "\n".join(row.detail),
                plan.detail_size,
            )
        if row.page is not None:
            text(11.75, y, 0.82, 0.36, f"{row.page:02d}", 13, align=PP_ALIGN.RIGHT)
        y += row.height
        if i < len(plan.rows) - 1:
            add_rect(slide, TEXT_X, y + plan.gap / 2, 11.15, 0.008, RULE)
        y += plan.gap
