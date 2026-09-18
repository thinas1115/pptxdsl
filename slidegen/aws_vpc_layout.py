"""AWS VPC/AZ/Subnet構成を描く専用renderer。"""
from dataclasses import dataclass

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from slidegen import generate
from slidegen.asset_paths import ASSET_DIR, resolve_icon_path
from slidegen.diagrams import add_arrow, arrow_label
from slidegen.layout_fit import FitError, fit_text_or_raise, select_fit, stepped
from slidegen.textfit import line_height_in, text_width_in


SUBNET_LINE = RGBColor(0x9A, 0xA0, 0x9E)


@dataclass(frozen=True)
class Box:
    x: float
    y: float
    w: float
    h: float

    @property
    def cx(self):
        return self.x + self.w / 2

    @property
    def cy(self):
        return self.y + self.h / 2


@dataclass(frozen=True)
class NodeBox:
    x: float
    y: float
    w: float
    h: float
    icon_size: float

    @property
    def cx(self):
        return self.x + self.w / 2

    @property
    def cy(self):
        return self.y + self.h / 2

    @property
    def left(self):
        return self.cx - self.icon_size / 2 - 0.06

    @property
    def right(self):
        return self.cx + self.icon_size / 2 + 0.06

    @property
    def top(self):
        return self.cy - self.icon_size / 2 - 0.06

    @property
    def bottom(self):
        return self.cy + self.icon_size / 2 + 0.06


def _frame(slide, box, label, *, color, name, width=1.2):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(box.x), Inches(box.y),
        Inches(box.w), Inches(box.h))
    shape.name = name
    shape.fill.background()
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    shape.shadow.inherit = False
    _box_label(slide, box.x + 0.13, box.y + 0.07, box.w - 0.26, label,
               color=color, name=f"{name}:label")
    return shape


def _box_label(slide, x, y, w, text, *, color, name):
    size, lines = fit_text_or_raise(
        "aws_vpc_layout", "label", text, w, 0.24, 9.2,
        min_pt=7.5, weight="bold", spacing=1.05, role="compact")
    rendered = "\n".join(lines)
    label_w = min(w, text_width_in(rendered, size, "bold") + 0.10)
    label_h = min(0.24, line_height_in(size, 1.05) * len(lines) + 0.03)
    tb = generate.add_text(
        slide, x, y, label_w, label_h, rendered, size,
        bold=True, color=color, anchor=MSO_ANCHOR.MIDDLE, spacing=1.05)
    tb.name = name
    tb.fill.solid()
    tb.fill.fore_color.rgb = generate.CANVAS
    return tb


def _node(slide, node, box, *, icon_size, title_size, sub_size):
    path = resolve_icon_path(node["icon"])
    if not path.exists():
        raise FileNotFoundError(
            f"アイコン {node['icon']} が {ASSET_DIR} にありません。")
    slide.shapes.add_picture(str(path), Inches(box.cx - icon_size / 2),
                             Inches(box.cy - icon_size / 2),
                             Inches(icon_size), Inches(icon_size))
    title_y = box.cy + icon_size / 2 + 0.05
    _node_label(slide, box.cx, title_y, node["label"],
                max_w=box.w, size=title_size, min_pt=8.0,
                bold=True, color=generate.NAVY)
    if node.get("sub"):
        _node_label(slide, box.cx, title_y + 0.25, node["sub"],
                    max_w=box.w, size=sub_size, min_pt=7.0,
                    bold=False, color=generate.GRAY)


def _node_label(slide, cx, y, text, *, max_w, size, min_pt, bold, color):
    weight = "bold" if bold else "regular"
    slot_h = 0.22 if bold else 0.18
    actual_size, lines = fit_text_or_raise(
        "aws_vpc_layout", "resource", text, max_w, slot_h, size,
        min_pt=min_pt, weight=weight, spacing=1.05,
        role="natural" if bold else "compact")
    rendered = "\n".join(lines)
    label_w = min(max_w, text_width_in(rendered, actual_size, weight) + 0.10)
    label_h = min(slot_h, line_height_in(actual_size, 1.05) * len(lines) + 0.03)
    tb = generate.add_text(
        slide, cx - label_w / 2, y, label_w, label_h, rendered, actual_size,
        bold=bold, color=color, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE, spacing=1.05)
    tb.fill.solid()
    tb.fill.fore_color.rgb = generate.CANVAS
    return tb


def _resource_slot(subnet_box, index, count, *, icon_size, has_sub=False):
    top = subnet_box.y + 0.31
    required_h = 0.31 + icon_size + 0.05 + (0.43 if has_sub else 0.22) + 0.04
    if subnet_box.h < required_h:
        raise FitError(
            "aws_vpc_layout: subnet内のリソース表示領域が不足しています。"
            "subnet数またはリソース数を減らしてください。")
    slot_w = subnet_box.w / count
    cx = subnet_box.x + slot_w * index + slot_w / 2
    cy = top + icon_size / 2
    return NodeBox(cx - slot_w / 2 + 0.04, cy - icon_size / 2,
                   slot_w - 0.08, icon_size, icon_size)


def _layout(area, spec):
    azs = spec["azs"]
    max_subnets = max(len(az["subnets"]) for az in azs)
    note_h = 0.34 if spec.get("note") else 0.0
    available_h = area.height - note_h

    def required_height(values):
        has_sub = any(resource.get("sub") for az in azs for subnet in az["subnets"]
                      for resource in subnet.get("resources", []))
        resource_h = 0.31 + values["icon_size"] + 0.05 \
            + (0.43 if has_sub else 0.22) + 0.04
        # _resource_slot、VPC/AZ/subnetの実際のpaddingと同じ寸法で判定する。
        return (0.011 + 0.08 + 0.42 + 0.20 + 0.36 + 0.10 + max_subnets * resource_h
                + values["subnet_gap"] * (max_subnets - 1))

    def candidates():
        for gap in stepped(0.24, 0.14, 0.05):
            values = {
                "az_gap": gap,
                "subnet_gap": max(0.10, gap - 0.08),
                "icon_size": 0.48,
                "title_size": 9.0,
                "sub_size": 7.8,
            }
            used = required_height(values)
            yield ("standard" if gap == 0.24 else "gap", values, used)
        for icon_size in stepped(0.44, 0.36, 0.04):
            values = {
                "az_gap": 0.14,
                "subnet_gap": 0.10,
                "icon_size": icon_size,
                "title_size": 8.2,
                "sub_size": 7.2,
            }
            used = required_height(values)
            yield "element", values, used

    fit = select_fit(
        "aws_vpc_layout", available_h, candidates(),
        guidance="AZ数、subnet数、リソース数を減らすか、スライドを分割してください。")
    values = dict(fit.values)
    values["stage"] = fit.stage

    external_w = 1.35 if spec.get("external") else 0.0
    external_gap = 0.22 if spec.get("external") else 0.0
    vpc_x = generate.MARGIN + external_w + external_gap
    vpc_y = area.top + 0.04
    vpc_w = generate.BODY_W - external_w - external_gap
    vpc_h = available_h - 0.08
    if vpc_h < 3.50:
        raise FitError(
            "aws_vpc_layout: VPC図の高さが不足しています。"
            "leadを短くするか、subnet数を減らしてください。")

    vpc_box = Box(vpc_x, vpc_y, vpc_w, vpc_h)
    inner_pad_x = 0.30
    inner_top = 0.42
    inner_bottom = 0.20
    az_count = len(azs)
    az_w = (vpc_w - inner_pad_x * 2
            - values["az_gap"] * (az_count - 1)) / az_count
    az_h = vpc_h - inner_top - inner_bottom
    if az_w < 2.35 or az_h < 3.05:
        raise FitError(
            "aws_vpc_layout: AZ領域が不足しています。"
            "AZ数を減らすか、内容を複数スライドへ分割してください。")

    subnet_top_pad = 0.36
    subnet_h = (
        az_h - subnet_top_pad - values["subnet_gap"] * (max_subnets - 1)
        - 0.10
    ) / max_subnets
    if subnet_h < 0.82:
        raise FitError(
            "aws_vpc_layout: subnetの高さが不足しています。"
            "subnet数またはリソース数を減らしてください。")

    az_boxes = {}
    subnet_boxes = {}
    for az_index, az in enumerate(azs):
        az_x = vpc_x + inner_pad_x + az_index * (az_w + values["az_gap"])
        az_box = Box(az_x, vpc_y + inner_top, az_w, az_h)
        az_boxes[az["id"]] = az_box
        for subnet_index, subnet in enumerate(az["subnets"]):
            subnet_y = az_box.y + subnet_top_pad + subnet_index * (
                subnet_h + values["subnet_gap"])
            subnet_boxes[subnet["id"]] = Box(
                az_box.x + 0.16, subnet_y, az_box.w - 0.32, subnet_h)
    return vpc_box, az_boxes, subnet_boxes, values


def _flow_label(slide, route_pts, text, frames, nodes):
    """ラベルの実測外形がコンテナ境界やノードを隠さない位置を選ぶ。"""
    width = min(1.2, text_width_in(text, 8.2) + 0.10)
    height = line_height_in(8.2, 1.1) + 0.04
    side_offset = max(node.icon_size for node in nodes.values()) / 2 + 0.06 \
        + width / 2 + 0.08

    def clear(cx, cy):
        left, right = cx - width / 2, cx + width / 2
        top, bottom = cy - height / 2, cy + height / 2
        for frame in frames:
            if top < frame.y + frame.h and bottom > frame.y:
                if any(left - 0.02 < edge < right + 0.02
                       for edge in (frame.x, frame.x + frame.w)):
                    return False
            if left < frame.x + frame.w and right > frame.x:
                if any(top - 0.02 < edge < bottom + 0.02
                       for edge in (frame.y, frame.y + frame.h)):
                    return False
        for node in nodes.values():
            if (left < node.right and right > node.left
                    and top < node.bottom + 0.50 and bottom > node.top):
                return False
        return True

    # 長い区間から探す。折れ点ではなく、各区間の内部にラベルを置く。
    segments = sorted(zip(route_pts, route_pts[1:]),
                      key=lambda pair: abs(pair[1][0] - pair[0][0])
                      + abs(pair[1][1] - pair[0][1]), reverse=True)
    for start, end in segments:
        for fraction in (0.5, 0.25, 0.75, 0.4, 0.6, 0.1, 0.9):
            x = start[0] + (end[0] - start[0]) * fraction
            y = start[1] + (end[1] - start[1]) * fraction
            offsets = ((0, -0.12), (0, -0.22), (0, 0.22)) \
                if start[1] == end[1] else \
                ((side_offset, 0), (-side_offset, 0))
            for dx, dy in offsets:
                if clear(x + dx, y + dy):
                    label = arrow_label(slide, x + dx, y + dy, text, w=1.2, size=8.2)
                    label.name = "aws-vpc-layout:flow-label"
                    return label
    raise FitError(f"aws_vpc_layout: 通信ラベル「{text}」の表示領域が不足しています。"
                   "ラベルを短くするか、通信を複数スライドへ分割してください。")


def _draw_flow(slide, flow, nodes, frames):
    source = nodes[flow["from"]]
    target = nodes[flow["to"]]
    dash = flow.get("dash")
    both = bool(flow.get("both"))
    if abs(source.cy - target.cy) < 0.08:
        if source.cx <= target.cx:
            start, end = (source.right, source.cy), (target.left, target.cy)
        else:
            start, end = (source.left, source.cy), (target.right, target.cy)
        route_pts = [start, end]
    elif abs(source.cx - target.cx) < 0.20:
        if source.cy <= target.cy:
            start, end = (source.cx, source.bottom), (target.cx, target.top)
        else:
            start, end = (source.cx, source.top), (target.cx, target.bottom)
        route_pts = [start, end]
    else:
        start = (source.right if source.cx < target.cx else source.left,
                 source.cy)
        end = (target.left if source.cx < target.cx else target.right,
               target.cy)
        mid_x = (start[0] + end[0]) / 2
        route_pts = [start, (mid_x, start[1]), (mid_x, end[1]), end]
    if len(route_pts) == 2:
        add_arrow(slide, *route_pts[0], *route_pts[1], dash=dash, both=both)
    else:
        from slidegen.diagrams3 import route
        route(slide, route_pts, dash=dash, both=both, width=1.25)
    if flow.get("label"):
        _flow_label(slide, route_pts, flow["label"], frames, nodes)


def s_aws_vpc_layout(slide, spec, page):
    """AWS VPC/AZ/Subnetの階層と代表的な通信を描く。"""
    area = generate.header(slide, spec["kicker"], spec["title"], spec.get("lead"))
    vpc_box, az_boxes, subnet_boxes, fit = _layout(area, spec)
    vpc_label = spec["vpc"]["label"]
    if spec["vpc"].get("cidr"):
        vpc_label = f"{vpc_label} {spec['vpc']['cidr']}"
    _frame(slide, vpc_box, vpc_label, color=generate.NAVY,
           name="aws-vpc-layout:vpc", width=1.3)

    nodes = {}
    for az in spec["azs"]:
        _frame(slide, az_boxes[az["id"]], az["label"], color=generate.ACCENT,
               name=f"aws-vpc-layout:az:{az['id']}", width=1.15)
        for subnet in az["subnets"]:
            subnet_box = subnet_boxes[subnet["id"]]
            label = subnet["label"]
            if subnet.get("cidr"):
                label = f"{label} {subnet['cidr']}"
            _frame(slide, subnet_box, label, color=SUBNET_LINE,
                   name=f"aws-vpc-layout:subnet:{subnet['id']}", width=0.9)
            resources = subnet.get("resources", [])
            for index, resource in enumerate(resources):
                slot = _resource_slot(
                    subnet_box, index, len(resources), icon_size=fit["icon_size"],
                    has_sub=bool(resource.get("sub")))
                nodes[resource["id"]] = slot

    for external in spec.get("external", []):
        target_y = vpc_box.cy
        outgoing = next(
            (flow for flow in spec.get("flows", [])
             if flow.get("from") == external["id"] and flow.get("to") in nodes),
            None)
        if outgoing:
            target_y = nodes[outgoing["to"]].cy
        slot = NodeBox(
            generate.MARGIN + 0.08, target_y - fit["icon_size"] / 2,
            1.18, fit["icon_size"], fit["icon_size"])
        nodes[external["id"]] = slot

    for flow in spec.get("flows", []):
        _draw_flow(slide, flow, nodes,
                   [vpc_box, *az_boxes.values(), *subnet_boxes.values()])

    for az in spec["azs"]:
        for subnet in az["subnets"]:
            for resource in subnet.get("resources", []):
                _node(slide, resource, nodes[resource["id"]],
                      icon_size=fit["icon_size"],
                      title_size=fit["title_size"],
                      sub_size=fit["sub_size"])
    for external in spec.get("external", []):
        _node(slide, external, nodes[external["id"]],
              icon_size=fit["icon_size"],
              title_size=fit["title_size"],
              sub_size=fit["sub_size"])
    if spec.get("note"):
        generate.note_line(slide, spec["note"])
