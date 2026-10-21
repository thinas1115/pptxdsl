"""設定キー・期待値・確認コードの対応と、非折り返しの下限を守る。"""
from copy import deepcopy
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from pptx import Presentation

from slidegen.check_layout import check
from slidegen.generate_from_json import main
from slidegen.layout_fit import FitError
from slidegen.training_renderers import _fit_code_text, s_code_lab
from slidegen.validate_content import validate

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "examples/authoring/config-checks.json"


class ConfigChecksTests(unittest.TestCase):
    def test_example_keys_expected_values_and_saved_code(self):
        deck = json.loads(SOURCE.read_text(encoding="utf-8"))
        self.assertFalse(validate(deck))
        spec = deck["slides"][0]
        code = "\n".join(section["code"] for section in spec["sections"])
        result = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        for key, expected in (("replicas", "3"), ("health_path", "/health"), ("timeout_s", "5")):
            self.assertIn(f'cfg["{key}"]', code)
            self.assertTrue(any(f"{key}: {expected}" in item for item in spec["checks"]))
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "config.pptx"
            main(SOURCE, output)
            self.assertFalse(check(output))
            shapes = list(Presentation(output).slides[0].shapes)
            for section in spec["sections"]:
                shape = next(s for s in shapes if s.has_text_frame and s.text == section["code"])
                self.assertFalse(shape.text_frame.word_wrap)
                self.assertTrue(all(run.font.size.pt >= 9.5
                                    for p in shape.text_frame.paragraphs for run in p.runs))

    def test_code_floor_and_split_guidance(self):
        # 高さを狭めても下限で収容できる場合と、横幅・高さの限界を分けて確認。
        from slidegen.textfit import line_height_in
        self.assertEqual(_fit_code_text("x = 1", 7, line_height_in(9.5, 1.18)), 9.5)
        for code, width, height in (("x" * 300, 7, 3), ("x\n" * 16, 7, .1)):
            with self.assertRaisesRegex(FitError, "分割"):
                _fit_code_text(code, width, height)
        spec = deepcopy(json.loads(SOURCE.read_text(encoding="utf-8"))["slides"][0])
        spec["sections"][0]["code"] = "x" * 300
        prs = Presentation()
        with self.assertRaisesRegex(FitError, "分割"):
            s_code_lab(prs.slides.add_slide(prs.slide_layouts[6]), spec, 1)


if __name__ == "__main__":
    unittest.main()
