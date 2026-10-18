"""工程の補足距離・条件文字・経路非交差・過密停止を検証する。"""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE

from slidegen.check_layout import check
from slidegen.diagrams2 import _fit_process_flow, _rects_overlap, s_process
from slidegen.layout_fit import FitError
from tools.gallery.generate_process_cases import main, make_deck


class ProcessAnnotationTests(unittest.TestCase):
    def test_saved_annotations_are_close_and_do_not_cover_routes(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "process.pptx"
            main(output)
            self.assertFalse(check(output))
            for slide, spec in zip(Presentation(output).slides, make_deck()["slides"]):
                shapes = {s.name: s for s in slide.shapes}
                if "steps" in spec:
                    for i, step in enumerate(spec["steps"]):
                        key = f"process:attribute-label:{i}"
                        if not step.get("attribute") and not step.get("actor"):
                            self.assertNotIn(key, shapes)
                            continue
                        desc, label = shapes[f"process:desc:{i}"], shapes[key]
                        gap = (label.top - desc.top - desc.height) / 914400
                        self.assertAlmostEqual(gap, .20, places=4)
                        self.assertEqual(desc.left + desc.width // 2, label.left + label.width // 2)
                        self.assertGreaterEqual(label.text_frame.paragraphs[0].runs[0].font.size.pt, 9.5)
                else:
                    labels = [s for s in slide.shapes if s.name.startswith("process:condition:")]
                    self.assertEqual(len(labels), sum(bool(e.get("label")) for e in spec["flow"]["edges"]))
                    for label in labels:
                        self.assertIn(label.fill.type, (None, MSO_FILL_TYPE.BACKGROUND))
                        self.assertFalse(label.text_frame.word_wrap)
                        self.assertGreaterEqual(label.text_frame.paragraphs[0].runs[0].font.size.pt, 9.5)
                        rect = tuple(v / 914400 for v in (label.left, label.top,
                                                        label.left + label.width, label.top + label.height))
                        for line in slide.shapes:
                            if line.shape_type != MSO_SHAPE_TYPE.LINE:
                                continue
                            bounds = tuple(v / 914400 for v in (line.left, line.top,
                                                               line.left + line.width, line.top + line.height))
                            self.assertFalse(_rects_overlap(rect, bounds, .02), (label.text, bounds))

    def test_flow_profiles_and_minimum_stop(self):
        levels = [["a", "b", "c"]]
        self.assertEqual(_fit_process_flow(5, levels).stage, "standard")
        self.assertEqual(_fit_process_flow(4.1, levels).stage, "gap")
        self.assertEqual(_fit_process_flow(3.5, levels).stage, "element")
        with self.assertRaises(FitError):
            _fit_process_flow(3, levels)

    def test_oversized_description_attribute_and_condition_stop(self):
        for target in ("desc", "attribute", "condition"):
            spec = deepcopy(make_deck()["slides"][4 if target == "condition" else 0])
            if target == "condition":
                spec["flow"]["edges"][2]["label"] = "判断条件" * 100
            elif target == "attribute":
                spec["steps"][0]["attribute"]["label"] = "確認対象" * 100
            else:
                spec["steps"][0]["desc"] = "説明文" * 100
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            with self.assertRaises(FitError):
                s_process(slide, spec, 1)


if __name__ == "__main__":
    unittest.main()
