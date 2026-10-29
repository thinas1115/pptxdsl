"""プロトコル図の凡例・状態変化・非比例縮尺の説明を確認する。"""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation
from pptx.util import Inches, Pt

from slidegen import generate
from slidegen.check_layout import check
from slidegen.training_renderers import (
    _fit_protocol, _protocol_changed_stages, s_protocol_state_flow, s_protocol_anatomy,
)
from slidegen.layout_fit import FitError
from types import SimpleNamespace
from tests.fixtures.gallery.content_patterns import PATTERN_DECK


def specimen(type_):
    return deepcopy(next(spec for spec in PATTERN_DECK["slides"] if spec["type"] == type_))


def render(spec):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    renderer = s_protocol_anatomy if spec["type"] == "nw_frame_anatomy" else s_protocol_state_flow
    generate.render_slide(renderer, slide, spec, 1)
    return prs, slide


class ProtocolSemanticsTests(unittest.TestCase):
    def test_change_uses_stage_order_not_input_order_or_processing_detail(self):
        stages = [{"id": name} for name in ("a", "b", "c", "d")]
        states = [{"stage": "d", "label": "B"}, {"stage": "c", "label": "B"},
                  {"stage": "b", "label": "A", "appearance": "internal", "detail": "変換処理"},
                  {"stage": "a", "label": "A"}]
        self.assertEqual(_protocol_changed_stages(stages, states), {"c"})

    def test_encapsulation_add_remove_are_changes_even_when_value_is_unchanged(self):
        stages = [{"id": name} for name in ("a", "b", "c", "d")]
        states = [{"stage": name, "label": "A"} for name in ("a", "b", "c", "d")]
        for state in states[1:3]:
            state.update(appearance="encapsulated", encapsulation="HDR")
        self.assertEqual(_protocol_changed_stages(stages, states), {"b", "d"})

    def test_legend_uses_same_four_markers_and_only_used_appearances(self):
        spec = specimen("nw_protocol_flow")
        for state, appearance in zip(spec["flows"][0]["states"],
                                     ("plain", "encapsulated", "internal", "alert", "plain")):
            state["appearance"] = appearance
            if appearance == "encapsulated":
                state["encapsulation"] = "HDR"
        _, slide = render(spec)
        for appearance in ("plain", "encapsulated", "internal", "alert"):
            self.assertTrue(any(shape.name == f"protocol-state-flow:legend:{appearance}" for shape in slide.shapes))
        for shape in slide.shapes:
            if shape.name == "protocol-state-flow:marker:plain":
                self.assertEqual(shape.line.color.rgb, generate.NAVY)
            elif shape.name == "protocol-state-flow:marker:alert":
                self.assertEqual(shape.line.color.rgb, generate.CORAL)
        for flow in spec["flows"]:
            for state in flow["states"]:
                state["appearance"] = "plain"
                state.pop("encapsulation", None)
        _, slide = render(spec)
        self.assertFalse(any(shape.name == "protocol-state-flow:legend:alert" for shape in slide.shapes))

    def test_nat_change_underline_survives_save_and_unchanged_destination_is_not_marked(self):
        prs, _ = render(specimen("nw_protocol_flow"))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "nat.pptx"
            prs.save(path)
            self.assertFalse(check(path))
            shapes = Presentation(path).slides[0].shapes
            changed = [shape for shape in shapes if shape.name.startswith("protocol-state-flow:value:changed:")]
            self.assertEqual([shape.name for shape in changed], ["protocol-state-flow:value:changed:0:nat"])
            self.assertTrue(all(run.font.underline for shape in changed
                                for paragraph in shape.text_frame.paragraphs for run in paragraph.runs))

    def test_scale_note_without_annotations_and_stronger_annotation_lines(self):
        for with_annotations in (False, True):
            spec = specimen("nw_frame_anatomy")
            if not with_annotations:
                spec.pop("lead", None)
                spec["frames"][0].pop("annotations", None)
            prs, slide = render(spec)
            note = next(shape for shape in slide.shapes if shape.name == "protocol-anatomy:scale-note")
            self.assertEqual(note.text, "幅はbit数に比例しない概念図")
            for shape in slide.shapes:
                if shape.name.startswith("protocol-anatomy:annotation:"):
                    self.assertEqual(shape.line.color.rgb, generate.GRAY)
                    self.assertEqual(shape.line.width, Pt(1.1))
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "anatomy.pptx"
                prs.save(path)
                self.assertFalse(check(path))

    def test_maximum_cases_have_no_text_or_line_collisions(self):
        from tools.gallery.generate_protocol_cases import make_deck
        for spec in make_deck()["slides"]:
            prs, _ = render(spec)
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "maximum.pptx"
                prs.save(path)
                self.assertFalse(check(path), spec["title"])

    def test_anatomy_preserves_all_fit_stages_and_stops_below_minimum(self):
        from tools.gallery.generate_protocol_cases import make_deck
        frames = make_deck()["slides"][-1]["frames"]
        observed = set()
        for step in range(200):
            try:
                result = _fit_protocol(SimpleNamespace(height=6.5 - step * .01), frames)
            except FitError:
                continue
            observed.add(result.stage)
        self.assertEqual(observed, {"standard", "gap", "element", "font"})
        with self.assertRaises(FitError):
            _fit_protocol(SimpleNamespace(height=1), frames)


if __name__ == "__main__":
    unittest.main()
