"""選定例を通常の入力契約で生成し、判断・条件・左右の構造を保持する。"""
from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation

from slidegen.check_layout import check
from slidegen.generate_from_json import main
from slidegen.validate_content import validate


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "examples/authoring/type-selection.json"


class AuthoringExampleTests(unittest.TestCase):
    def test_regular_input_and_saved_pptx(self):
        deck = json.loads(SOURCE.read_text(encoding="utf-8"))
        self.assertFalse(validate(deck))  # 回帰サンプル許可で通常の入力検証を迂回しない。
        self.assertEqual([s["type"] for s in deck["slides"]],
                         ["paired_comparison", "two_column", "scope_boundary", "bullets"])
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "selection.pptx"
            main(SOURCE, output)
            self.assertFalse(check(output))
            prs = Presentation(output)
            for slide, spec in zip(prs.slides, deck["slides"]):
                text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
                self.assertIn(spec["title"], text)
                if "lead" in spec:
                    self.assertIn(spec["lead"], text)
            scope = deck["slides"][2]
            for phrase in ("既存窓口を残し", "個人情報を含まない"):
                self.assertIn(phrase, scope["lead"])
            self.assertNotEqual(len(deck["slides"][1]["left"]["bullets"]),
                                len(deck["slides"][1]["right"]["bullets"]))
            self.assertIn("本番切替は今回の判断に含めません", deck["slides"][3]["lead"])

    def test_retired_summary_and_condition_fields_stay_rejected(self):
        deck = json.loads(SOURCE.read_text(encoding="utf-8"))
        for index, field in ((2, "preconditions"), (3, "conclusion")):
            bad = deepcopy(deck)
            bad["slides"][index][field] = "廃止フィールドの拒否試験"
            self.assertTrue(validate(bad))
        bad = deepcopy(deck)
        bad["slides"][3]["type"] = "decision_summary"
        self.assertTrue(validate(bad))


if __name__ == "__main__":
    unittest.main()
