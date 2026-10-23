"""工程線の太さ・期間座標・ラベル可読性・明示停止を検証する。"""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation

from slidegen.check_layout import check
from slidegen.diagrams2 import PROGRAM_LINE_PT, s_program_roadmap
from slidegen.layout_fit import FitError
from slidegen.timeline_layout import fit_program_roadmap, pack_activities
from tools.gallery.generate_roadmap_cases import main, make_deck, overlong_label_case


class RoadmapReadabilityTests(unittest.TestCase):
    def test_saved_spans_preserve_periods_and_dominate_grid(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "roadmap.pptx"
            main(output)
            self.assertFalse(check(output))
            for slide, spec in zip(Presentation(output).slides, make_deck()["slides"]):
                shapes = {shape.name: shape for shape in slide.shapes}
                grids = [s for s in slide.shapes if s.name == "roadmap:grid"]
                self.assertTrue(grids)
                self.assertEqual(PROGRAM_LINE_PT, 1.4)
                self.assertGreater(PROGRAM_LINE_PT, grids[0].line.width.pt)
                for i, track in enumerate(spec["tracks"]):
                    placements, _ = pack_activities(track["activities"], spec["periods"])
                    for item in placements:
                        line = shapes[f"roadmap:span:{i}:{item.index}"]
                        self.assertEqual(line.line.width.pt, PROGRAM_LINE_PT)
                        width = (item.end - item.start) * 9.09 / len(spec["periods"]) - .07
                        self.assertAlmostEqual(line.width.inches, width, places=4)
                        label = shapes[f"roadmap:activity:{i}:{item.index}"]
                        self.assertEqual(label.text, item.activity["label"])
                        self.assertGreaterEqual(label.text_frame.paragraphs[0].runs[0].font.size.pt, 7.5)
                        self.assertAlmostEqual((line.left + line.width / 2) / 914400,
                                               (label.left + label.width / 2) / 914400, places=4)
                    if "milestone" in track:
                        label = shapes[f"roadmap:milestone-label:{i}"]
                        self.assertGreaterEqual(label.text_frame.paragraphs[0].runs[0].font.size.pt, 7)

    def test_profiles_keep_readability_floors_and_stop(self):
        cases = [(5.27, [1, 2, 1], "standard"), (4.4, [2, 2, 4, 1, 1], "gap"),
                 (4.8, [3, 3, 3, 3, 3], "element")]
        for available, counts, stage in cases:
            fit = fit_program_roadmap(available, counts)
            self.assertEqual(fit.stage, stage)
            self.assertGreaterEqual(fit.values["period_pt"], 8)
            self.assertGreaterEqual(fit.values["activity_pt"], 8)
            self.assertGreaterEqual(fit.values["track_pt"], 9.5)
        with self.assertRaises(FitError):
            fit_program_roadmap(5.27, [4] * 6)

    def test_overlong_labels_stop_without_tiny_text(self):
        for field in ("activity", "milestone", "goal", "name"):
            spec = deepcopy(make_deck()["slides"][2])
            track = spec["tracks"][0]
            if field == "activity":
                track["activities"][0]["label"] = "確認作業" * 100
            elif field == "milestone":
                track["milestone"]["label"] = "確認条件" * 100
            else:
                track[field] = "確認対象" * 100
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            with self.assertRaises(FitError):
                s_program_roadmap(slide, spec, 1)

    def test_original_width_rejects_withdrawn_long_label_case(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        with self.assertRaisesRegex(FitError, "最小フォント7.5ptでも幅1.41in"):
            s_program_roadmap(slide, overlong_label_case(), 1)


if __name__ == "__main__":
    unittest.main()
