"""保存PPTXの文字書式・余白・行送りに対する真値付き検査。"""
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from slidegen import generate as g
from slidegen.check_layout import check, rect_of, table_cells
from slidegen.font_config import font_family
from slidegen.text_metrics import TextMetricsError, layout_text_frame
from slidegen.text_roles import set_text_role
from slidegen.textfit import boundary_spacing_in, text_width_in, width_safety, wrapper_for_role


def canvas():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(g.SLIDE_W), Inches(g.SLIDE_H)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def restored(prs):
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    return Presentation(stream).slides[0]


class TextMetricsTests(unittest.TestCase):
    def test_saved_title_lead_body_compact_use_matching_wrapper(self):
        text = "資料の作成条件と実行後の確認方法を整理します。"
        for role, bold, size, spacing in (("natural", True, 20, 1.12), ("body", False, 14, 1.18),
                                           ("body", False, 12, 1.3), ("compact", False, 11, 1.15)):
            prs, slide = canvas()
            g.add_text(slide, 1, 1, 2.6, 3, text, size, bold=bold, spacing=spacing, role=role)
            shape = restored(prs).shapes[0]
            layout = layout_text_frame(shape.text_frame, rect_of(shape))
            expected = wrapper_for_role(role)(text, 2.6, size, "bold" if bold else "regular")
            self.assertEqual([line.text for line in layout.lines], expected)
            self.assertAlmostEqual(layout.height, len(expected) * size * spacing / 72, places=5)
            for line in layout.lines:
                self.assertAlmostEqual(line.rect[2] - line.rect[0],
                                       text_width_in(line.text, size, "bold" if bold else "regular"), places=5)

    def test_mixed_runs_use_all_sizes_and_weights(self):
        prs, slide = canvas()
        shape = g.add_text(slide, 1, 1, 8, 2, "small", 9, spacing=1.7, align=PP_ALIGN.CENTER)
        paragraph = shape.text_frame.paragraphs[0]
        run = paragraph.add_run()
        run.text = "BIG"
        g.set_run(run, 28, bold=True)
        shape = restored(prs).shapes[0]
        layout = layout_text_frame(shape.text_frame, rect_of(shape))
        self.assertEqual(len(layout.lines), 1)
        line = layout.lines[0]
        expected = text_width_in("small", 9) + text_width_in("BIG", 28, "bold")
        self.assertAlmostEqual(line.rect[2] - line.rect[0], expected, places=5)
        self.assertAlmostEqual(line.rect[0], 1 + (8 - expected) / 2, places=5)
        self.assertEqual(line.max_size, 28)
        self.assertAlmostEqual(layout.height, 28 * 1.7 / 72, places=5)

    def test_paragraph_alignment_fixed_spacing_before_and_after(self):
        prs, slide = canvas()
        shape = g.add_text(slide, 1, 1, 8, 3, "Left", 12, spacing=1.0)
        first = shape.text_frame.paragraphs[0]
        first.space_after = Pt(5)
        second = shape.text_frame.add_paragraph()
        second.alignment = PP_ALIGN.RIGHT
        second.line_spacing = Pt(18)
        second.space_before = Pt(3)
        second.space_after = Pt(2)
        run = second.add_run()
        run.text = "Right"
        g.set_run(run, 10)
        shape = restored(prs).shapes[0]
        layout = layout_text_frame(shape.text_frame, rect_of(shape))
        self.assertAlmostEqual(layout.height, (12 + 5 + 3 + 18 + 2) / 72, places=5)
        self.assertAlmostEqual(layout.lines[0].rect[0], 1, places=5)
        self.assertAlmostEqual(layout.lines[1].rect[2], 9, places=5)
        self.assertAlmostEqual(layout.lines[1].advance, 18 / 72, places=5)
        self.assertAlmostEqual(layout.lines[1].rect[1], 1 + (12 + 5 + 3 + (18 - 10) / 2) / 72, places=5)

    def test_margins_and_middle_bottom_anchor(self):
        for anchor in (MSO_ANCHOR.MIDDLE, MSO_ANCHOR.BOTTOM):
            prs, slide = canvas()
            shape = g.add_text(slide, 1, 1, 6, 3, "本文", 18, spacing=1.0, anchor=anchor)
            frame = shape.text_frame
            frame.margin_left, frame.margin_right = Inches(.2), Inches(.3)
            frame.margin_top, frame.margin_bottom = Inches(.4), Inches(.5)
            shape = restored(prs).shapes[0]
            layout = layout_text_frame(shape.text_frame, rect_of(shape))
            self.assertEqual(tuple(round(value, 4) for value in layout.inner), (1.2, 1.4, 6.7, 3.5))
            self.assertAlmostEqual(layout.lines[0].rect[0], 1.2, places=5)
            y = 1.4 + (2.1 - .25) / 2 if anchor == MSO_ANCHOR.MIDDLE else 3.5 - .25
            self.assertAlmostEqual(layout.lines[0].rect[1], y, places=5)

    def test_manual_break_empty_paragraph_and_no_wrap(self):
        prs, slide = canvas()
        shape = g.add_text(slide, 1, 1, 1, 4, "A", 12, wrap=False, spacing=1.2)
        first = shape.text_frame.paragraphs[0]
        first.text = "A\vB"
        for run in first.runs:
            g.set_run(run, 12)
        paragraph = shape.text_frame.add_paragraph()
        paragraph.font.size = Pt(22)
        paragraph = shape.text_frame.add_paragraph()
        paragraph.line_spacing = 1.0
        run = paragraph.add_run()
        run.text = "LONG_IDENTIFIER_WITHOUT_WRAP"
        g.set_run(run, 12)
        shape = restored(prs).shapes[0]
        layout = layout_text_frame(shape.text_frame, rect_of(shape))
        self.assertEqual([line.text for line in layout.lines], ["A", "B", "", "LONG_IDENTIFIER_WITHOUT_WRAP"])
        self.assertAlmostEqual(layout.height, (2 * 12 * 1.2 + 22 + 12) / 72, places=5)
        self.assertGreater(layout.lines[-1].rect[2], 2)

    def test_table_cell_margin_is_applied_only_once(self):
        prs, slide = canvas()
        table = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(5), Inches(1))
        g._cell(table.table.cell(0, 0), "セル本文", 12, center=True)
        shape = restored(prs).shapes[0]
        outer, inner, frame, _, _ = next(table_cells(shape))
        layout = layout_text_frame(frame, inner, box_is_inner=True, default_role="compact")
        self.assertEqual(layout.inner, inner)
        self.assertEqual(layout.role, "compact")
        self.assertAlmostEqual((layout.bounds[0] + layout.bounds[2]) / 2, 3.5, places=5)
        self.assertGreater(layout.bounds[0], outer[0] + .09)
        self.assertLess(layout.bounds[3], outer[3] - .04)

    def test_inherited_paragraph_style(self):
        prs, slide = canvas()
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        frame = shape.text_frame
        frame.margin_left = frame.margin_top = frame.margin_right = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.font.name = font_family()
        paragraph.font.size = Pt(19)
        paragraph.font.bold = True
        paragraph.line_spacing = 1.4
        paragraph.add_run().text = "Inherited"
        shape = restored(prs).shapes[0]
        line = layout_text_frame(shape.text_frame, rect_of(shape)).lines[0]
        self.assertEqual(line.max_size, 19)
        self.assertAlmostEqual(line.advance, 19 * 1.4 / 72, places=5)
        self.assertAlmostEqual(line.rect[2] - line.rect[0], text_width_in("Inherited", 19, "bold"), places=5)

    def test_checker_rejects_overflow_collision_and_unknown_font(self):
        with TemporaryDirectory() as temporary:
            path = Path(temporary) / "broken.pptx"
            prs, slide = canvas()
            g.add_text(slide, 1, 1, .4, .4, "LONG_IDENTIFIER", 18, wrap=False)
            prs.save(path)
            self.assertIn("TEXT-OOB", [kind for _, kind, _, _ in check(path)])
            prs, slide = canvas()
            shape = g.add_text(slide, 1, 1, 4, 1, "衝突\n衝突", 24)
            for paragraph in shape.text_frame.paragraphs:
                paragraph.line_spacing = Pt(1)
            prs.save(path)
            self.assertIn("T-T", [kind for _, kind, _, _ in check(path)])
            prs, slide = canvas()
            shape = g.add_text(slide, 1, 1, 4, 1, "Mismatch", 18)
            shape.text_frame.paragraphs[0].runs[0].font.name = "Unknown Font"
            prs.save(path)
            self.assertIn("TEXT-METRICS", [kind for _, kind, _, _ in check(path)])

    def test_unmeasurable_alignment_is_not_silently_treated_as_left(self):
        prs, slide = canvas()
        shape = g.add_text(slide, 1, 1, 4, 1, "本文", 18, align=PP_ALIGN.JUSTIFY)
        with self.assertRaises(TextMetricsError):
            layout_text_frame(shape.text_frame, rect_of(shape))

    def test_mixed_run_wrap_preserves_text_and_each_line_size(self):
        prs, slide = canvas()
        width = max(text_width_in("AB", 12), text_width_in("WORD", 24, "bold")) + .05
        shape = g.add_text(slide, 1, 1, width, 3, "AB ", 12, role="compact")
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = "WW WORD"
        g.set_run(run, 24, bold=True)
        layout = layout_text_frame(shape.text_frame, rect_of(shape))
        self.assertEqual([line.text for line in layout.lines], ["AB WW", "WORD"])
        for line in layout.lines:
            self.assertEqual(line.max_size, 24)
            self.assertLessEqual(line.rect[2] - line.rect[0], width)

    def test_checker_does_not_fill_empty_space_between_different_alignments(self):
        prs, slide = canvas()
        shape = g.add_text(slide, 1, 1, 8, 1, "LEFT", 12)
        paragraph = shape.text_frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.RIGHT
        paragraph.line_spacing = 1.3
        run = paragraph.add_run()
        run.text = "RIGHT"
        g.set_run(run, 12)
        g.add_text(slide, 1, 1 + 12 * 1.3 / 72, 2, .4, "CLEAR", 12)
        with TemporaryDirectory() as temporary:
            path = Path(temporary) / "normal.pptx"
            prs.save(path)
            self.assertFalse(check(path), check(path))

    def test_first_line_indent_does_not_shrink_later_lines(self):
        prs, slide = canvas()
        shape = g.add_text(slide, 1, 1, 2, 3, "確認項目の内容と実行条件を整理して保存します。", 12, role="compact")
        properties = shape.text_frame.paragraphs[0]._p.get_or_add_pPr()
        properties.set("marL", str(Inches(.2)))
        properties.set("indent", str(Inches(.4)))
        layout = layout_text_frame(shape.text_frame, rect_of(shape))
        self.assertGreater(len(layout.lines), 1)
        self.assertAlmostEqual(layout.lines[0].rect[0], 1.6, places=5)
        self.assertAlmostEqual(layout.lines[1].rect[0], 1.2, places=5)
        self.assertLessEqual(layout.lines[0].rect[2], 3.00001)
        self.assertLessEqual(layout.lines[1].rect[2], 3.00001)

    def test_japanese_latin_spacing_crosses_a_style_boundary(self):
        prs, slide = canvas()
        shape = g.add_text(slide, 1, 1, 6, 1, "日本語", 12, wrap=False)
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = "API"
        g.set_run(run, 12, bold=True)
        line = layout_text_frame(shape.text_frame, rect_of(shape)).lines[0]
        expected = (text_width_in("日本語", 12) + text_width_in("API", 12, "bold")
                    + boundary_spacing_in("語A", 12) * width_safety())
        self.assertAlmostEqual(line.rect[2] - line.rect[0], expected, places=5)

    def test_saved_autofit_scales_font_and_line_spacing(self):
        prs, slide = canvas()
        shape = g.add_text(slide, 1, 1, 4, 1, "本文", 20, spacing=1.4)
        body = shape.text_frame._txBody.bodyPr
        for child in list(body):
            if child.tag in {qn("a:noAutofit"), qn("a:normAutofit"), qn("a:spAutoFit")}:
                body.remove(child)
        autofit = body.makeelement(qn("a:normAutofit"), {"fontScale": "80000", "lnSpcReduction": "10000"})
        body.insert(0, autofit)
        shape = restored(prs).shapes[0]
        layout = layout_text_frame(shape.text_frame, rect_of(shape))
        self.assertAlmostEqual(layout.lines[0].max_size, 16)
        self.assertAlmostEqual(layout.height, 16 * 1.4 * .9 / 72, places=5)

    def test_invalid_saved_values_are_diagnostics_not_nan_rectangles(self):
        for attribute, value in (("marL", "nan"), ("indent", "inf")):
            prs, slide = canvas()
            shape = g.add_text(slide, 1, 1, 4, 1, "本文", 12)
            shape.text_frame.paragraphs[0]._p.get_or_add_pPr().set(attribute, value)
            with self.assertRaises(TextMetricsError):
                layout_text_frame(shape.text_frame, rect_of(shape))
        prs, slide = canvas()
        shape = g.add_text(slide, 1, 1, 4, 1, "本文", 12)
        paragraph = shape.text_frame.add_paragraph()
        paragraph._p.get_or_add_pPr().get_or_add_defRPr().set("sz", "nan")
        with self.assertRaises(TextMetricsError):
            layout_text_frame(shape.text_frame, rect_of(shape))

    def test_fitted_box_forwards_the_measurement_role(self):
        from slidegen.candidate_renderers import _text_in_box
        prs, slide = canvas()
        _text_in_box(slide, "test", "label", 1, 1, 4, 1, "短い属性値", 12, 10,
                     bold=True, role="compact")
        shape = restored(prs).shapes[0]
        self.assertEqual(layout_text_frame(shape.text_frame, rect_of(shape)).role, "compact")


if __name__ == "__main__":
    unittest.main()
