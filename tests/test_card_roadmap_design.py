"""番号付きカードの旧配置、工程表、収容段階と最小値停止を検証する。"""

from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.check_layout import check
from slidegen.layout_fit import FitError
from slidegen.layout_settings import ROADMAP
from slidegen.timeline_layout import plan_program_roadmap
from slidegen.diagrams2 import s_program_roadmap
from tools.gallery.generate_card_roadmap_cases import main, make_deck


class CardRoadmapDesignTests(unittest.TestCase):
    def test_fifteen_cases_pass_layout_and_styles_have_distinct_hierarchy(self):
        data = make_deck()
        self.assertEqual(len(data["slides"]), 15)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "cases.pptx"
            main(output)
            self.assertFalse(check(output))
            for slide, spec in zip(Presentation(output).slides, data["slides"]):
                if spec["type"] != "cards":
                    continue
                panels = [shape for shape in slide.shapes if shape.name.startswith("cards:editorial:panel:")]
                if spec["style"] == "editorial":
                    self.assertFalse(panels)
                    self.assertFalse(any(shape.name.startswith("cards:editorial:accent:")
                                         for shape in slide.shapes))
                    text_shapes = [shape for shape in slide.shapes if shape.has_text_frame]
                    for index in range(len(spec["cards"])):
                        number = next(shape for shape in text_shapes
                                      if shape.text == f"{index + 1:02d}")
                        self.assertEqual(number.text, f"{index + 1:02d}")
                        card = spec["cards"][index]
                        heading = next(shape for shape in text_shapes
                                       if shape.text.replace("\n", "") == card["heading"])
                        body = next(shape for shape in text_shapes
                                    if shape.text.replace("\n", "") == card["body"])
                        self.assertAlmostEqual(heading.left.inches - number.left.inches, .72)
                        self.assertAlmostEqual(body.left.inches, heading.left.inches)
                        self.assertAlmostEqual(body.top.inches - heading.top.inches, .58)
                        for shape, minimum in ((heading, 13), (body, 11)):
                            self.assertGreaterEqual(shape.top.inches, generate.BODY_TOP)
                            self.assertLessEqual((shape.top + shape.height) / 914400,
                                                 generate.BODY_BOTTOM)
                            self.assertGreaterEqual(shape.text_frame.paragraphs[0].runs[0].font.size.pt, minimum)
                else:
                    self.assertFalse(panels)
                    values = {card["value"] for card in spec["cards"]}
                    large_values = [shape for shape in slide.shapes if shape.has_text_frame and shape.text in values]
                    self.assertEqual(len(large_values), len(values))
                    self.assertTrue(all(shape.text_frame.paragraphs[0].runs[0].font.size.pt >= 22 for shape in large_values))

    def test_editorial_uses_gap_and_font_before_explicit_stop(self):
        spec = deepcopy(make_deck()["slides"][3])
        for card in spec["cards"]:
            card["body"] = "判断に必要な情報を確認し、対象と条件を明確にしてから次の作業へ進みます。"
        observed, stopped = set(), False
        original = generate.fit_vertical_stacks

        def capture(*args, **kwargs):
            fitted = original(*args, **kwargs)
            observed.add(fitted.stage)
            self.assertGreaterEqual(fitted.size, 11)
            return fitted

        for step in range(95):
            area = generate.ContentArea(1.58, 1.58 + 5.27 - step * .03)
            prs = Presentation()
            prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            with patch.object(generate, "header", return_value=area), patch.object(generate, "fit_vertical_stacks", side_effect=capture):
                try:
                    generate.s_cards(slide, spec, 1)
                except FitError:
                    stopped = True
        self.assertEqual(observed, {"standard", "gap", "font"})
        self.assertTrue(stopped)

    def test_delayed_theme_and_goal_stay_in_left_column(self):
        spec = {"periods": ["前期", "中期", "後期"], "tracks": [{
            "name": "試行", "goal": "結果の記録", "activities": [
                {"label": "受付", "start": 1, "end": 2},
                {"label": "照合", "start": 1.5, "end": 3},
            ]}]}
        plan = plan_program_roadmap(spec, generate.ContentArea(1.58, 6.85))
        self.assertAlmostEqual(plan.grid_x, ROADMAP.label_x + ROADMAP.track_width)
        first = plan.activities[0][0]
        self.assertGreater(first.x1, plan.grid_x)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        s_program_roadmap(slide, dict(spec, kicker="旧配置の回帰検証", title="工程表"), 1)
        name = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text == "試行")
        goal = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text == "結果の記録")
        self.assertLess((name.left + name.width) / 914400, plan.grid_x)
        self.assertLess((goal.left + goal.width) / 914400, plan.grid_x)
        self.assertLess(plan.bottom, 6.85)
        with self.assertRaises(FitError):
            plan_program_roadmap(spec, generate.ContentArea(1.58, 2.0))

    def test_short_first_bar_preserves_original_lane_and_milestone(self):
        spec = {"periods": [str(index) for index in range(12)], "tracks": [{
            "name": "短期作業", "goal": "結果の確認",
            "milestone": {"at": 6, "label": "完了"},
            "activities": [
                {"label": "半月作業", "start": 4.25, "end": 4.75},
                {"label": "後続作業", "start": 5, "end": 6},
            ]}]}
        original = deepcopy(spec)
        plan = plan_program_roadmap(spec, generate.ContentArea(1.58, 6.85))
        first, following = plan.activities[0]
        self.assertAlmostEqual(plan.grid_x, 3.5)
        self.assertAlmostEqual(first.lane_top, following.lane_top)
        self.assertGreater(plan.milestones[0].y, following.lane_top)
        self.assertLess(plan.bottom, 6.85)
        self.assertEqual(spec, original)
        with self.assertRaises(FitError):
            plan_program_roadmap(spec, generate.ContentArea(1.58, 2.0))


if __name__ == "__main__":
    unittest.main()
