"""入力形式・移行警告・登録・タイトル例外・パス公開方針の回帰。"""
from copy import deepcopy
import contextlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from pptx import Presentation

from slidegen.cover_footer import parse_cover_footer_config
from slidegen import generate
from slidegen.generate_from_json import RENDER, main
from slidegen.validate_content import VALIDATORS, _TYPE_KEYS, validate, validation_warnings


def deck():
    return {"meta": {"title": "入力契約"}, "slides": [{
        "type": "bullets", "kicker": "分類", "title": "検討項目",
        "bullets": [{"text": "説明内容"}],
    }]}


class InputContractTests(unittest.TestCase):
    def test_type_sets_match(self):
        self.assertEqual(set(RENDER), set(VALIDATORS))
        self.assertEqual(set(RENDER), set(_TYPE_KEYS))

    def test_legacy_bullets_warn_without_mutation(self):
        data = deck()
        self.assertFalse(validation_warnings(data))
        data["slides"][0]["bullets"] = [["説明内容", None]]
        original = deepcopy(data)
        self.assertFalse(validate(data))
        self.assertIn("2.0.0", validation_warnings(data)[0])
        self.assertEqual(data, original)
        for value in (["説明", "消える値"], ["説明"], [None, None]):
            data["slides"][0]["bullets"] = [value]
            self.assertTrue(validate(data))

    def test_migration_preserves_rendered_shapes(self):
        canonical = deck()["slides"][0]
        legacy = {**canonical, "bullets": [["説明内容", None]]}
        outputs = []
        for spec in (canonical, legacy):
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            generate.s_bullets(slide, spec, 1)
            outputs.append(slide.shapes._spTree.xml)
        self.assertEqual(*outputs)

    def test_nominal_review_is_narrow_and_bound_to_title(self):
        data = deck()
        slide = data["slides"][0]
        slide["title"] = "こたつ"
        self.assertTrue(validate(data))
        review = {"title": "こたつ", "reason": "暖房器具の名称であり動作を述べる文章ではないため"}
        slide["title_nominal_review"] = review
        self.assertFalse(validate(data))
        for title in ("こたつ。", "こたつ\n", "効果を改善", "この方式が必要", "資料を作る", "新方式は有効だ"):
            slide["title"] = title
            review["title"] = title
            self.assertTrue(validate(data), title)
        slide["title"] = "こたつ"
        for invalid in (True, {}, {**review, "title": "こたつ", "reason": "短い"},
                        {**review, "title": "別の見出し"}, {**review, "allow_all": True}):
            slide["title_nominal_review"] = invalid
            self.assertTrue(validate(data), invalid)
        slide["title"] = "要点"
        slide["title_nominal_review"] = {**review, "title": "要点"}
        self.assertTrue(validate(data), "不要な例外指定を拒否する")

    def test_url_contracts(self):
        for kind in ("table", "references"):
            data = deck()
            if kind == "table":
                slide = {"type": kind, "kicker": "出典", "title": "規格",
                         "columns": ["名称", "説明"], "rows": [["規格", "説明"]],
                         "note": "詳細", "note_link": {"label": "原典", "url": ""}}
                target = slide["note_link"]
            else:
                target = {"id": "1", "title": "原典", "url": ""}
                slide = {"type": kind, "kicker": "出典", "title": "規格", "references": [target]}
            data["slides"] = [slide]
            target["url"] = "https://example.com/path#section"
            self.assertFalse(validate(data))
            target["url"] = "http://example.com/archive"
            self.assertEqual(bool(validate(data)), kind == "table")
            for url in ("https://", "https:///path", "https://user:secret@example.com",
                        "https://example.com/a b", "https://[bad", "https://example.com:bad",
                        "https://example.com/a\\b", "https://example.com/\x00", "javascript:alert(1)"):
                target["url"] = url
                self.assertTrue(validate(data), url)

    def test_absolute_background_paths_rejected_portably(self):
        self.assertIsNotNone(parse_cover_footer_config({}).cover.background_image)
        for value in ("/private/cover.png", "Z:/private/cover.png", " Z:/private/cover.png ", "Z:cover.png",
                      "\\\\server\\private\\cover.png", "~/cover.png"):
            with self.assertRaises(ValueError) as caught:
                parse_cover_footer_config({"cover": {"background_image": value}})
            self.assertIn("相対パス", str(caught.exception))
            self.assertNotIn(value, str(caught.exception))

    def test_cli_warning_and_defensive_renderer_error(self):
        with TemporaryDirectory() as temporary:
            source = Path(temporary) / "input.json"
            data = deck()
            data["slides"][0]["bullets"] = [["説明内容", None]]
            source.write_text(json.dumps(data), encoding="utf-8")
            result = subprocess.run([sys.executable, "-m", "slidegen.validate_content", str(source)],
                                    capture_output=True, text=True, encoding="utf-8",
                                    env=dict(os.environ, PYTHONIOENCODING="utf-8"), check=False)
            self.assertEqual(result.returncode, 0)
            self.assertIn("警告", result.stderr)
            self.assertNotIn(temporary, result.stderr)
            with patch.dict(RENDER, {}, clear=True), contextlib.redirect_stderr(io.StringIO()):
                with self.assertRaises(SystemExit) as caught:
                    main(source, Path(temporary) / "output.pptx")
            self.assertIn("rendererがありません", str(caught.exception))
            self.assertNotIn(temporary, str(caught.exception))

    def test_non_string_type_is_user_error(self):
        for value in ([], {}, None, 1):
            data = deck()
            data["slides"][0]["type"] = value
            self.assertTrue(validate(data))


if __name__ == "__main__":
    unittest.main()
