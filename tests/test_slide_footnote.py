"""共通補足欄の全type描画、本文の予約領域、収容停止を検証する。"""
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.generate_from_json import RENDER
from slidegen.layout_fit import FitError
from slidegen.slide_footnote import fit_footnote
from slidegen.validate_content import validate
from tests.fixtures.gallery.content_patterns import PATTERN_DECK


def main():
    specimens = {spec['type']: deepcopy(spec) for spec in PATTERN_DECK['slides']}
    specimens['aws_vpc_layout']['azs'] = specimens['aws_vpc_layout']['azs'][:1]
    specimens['aws_vpc_layout']['azs'][0]['subnets'] = specimens['aws_vpc_layout']['azs'][0]['subnets'][:1]
    specimens['aws_vpc_layout'].pop('external', None)
    specimens['aws_vpc_layout'].pop('flows', None)
    assert set(specimens) == set(RENDER)
    generate.DECK = {'meta': {'title': '共通補足欄の検証', 'author': '作成担当'},
                     'slides': list(specimens.values())}
    prs = Presentation()
    prs.slide_width = Inches(generate.SLIDE_W)
    prs.slide_height = Inches(generate.SLIDE_H)
    for index, spec in enumerate(specimens.values(), 1):
        spec['footnote'] = {'label': '注意', 'text': 'この補足欄は全type共通です。'}
        assert not validate({'meta': generate.DECK['meta'], 'slides': [spec]}, allow_sample_content=True)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        generate.render_slide(RENDER[spec['type']], slide, spec, index)
        notes = {shape.name: shape for shape in slide.shapes if shape.name.startswith('slide-footnote:')}
        assert notes['slide-footnote:text'].text == spec['footnote']['text']
        assert notes['slide-footnote:label'].text == '注意'
        for shape in slide.shapes:
            if shape.name.startswith('slide-footnote:'):
                continue
            # 背景は除外。本文テキスト・線・アイコンは補足欄の上へ収まる。
            if shape.width > Inches(13) or shape.height > Inches(7):
                continue
            assert shape.top + shape.height <= notes['slide-footnote:rule'].top + Inches(.02), (spec['type'], shape.name)
        assert generate._CONTENT_BOTTOM.get() == generate.BODY_BOTTOM
        generate.footer(slide, index)
    output = Path('out/footnote-all-types.pptx')
    output.parent.mkdir(exist_ok=True)
    prs.save(output)

    spec = deepcopy(specimens['bullets'])
    spec.pop('footnote')
    positions = []
    for _ in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        generate.render_slide(RENDER['bullets'], slide, spec, 1)
        assert not any(sh.name.startswith('slide-footnote:') for sh in slide.shapes)
        positions.append([(sh.top, sh.height) for sh in slide.shapes])
    assert positions[0] == positions[1]
    spec['footnote'] = {'text': 'ラベルを省略した補足です。'}
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    generate.render_slide(RENDER['bullets'], slide, spec, 1)
    assert not any(sh.name == 'slide-footnote:label' for sh in slide.shapes)

    stages = set()
    for n in range(1, 100):
        try:
            layout = fit_footnote({'text': '補足の条件を確認します。' * n}, generate.BODY_W)
        except FitError:
            break
        stages.add(layout.stage)
    assert stages == {'standard', 'gap', 'font'}, stages
    try:
        fit_footnote({'text': '補足' * 1000}, generate.BODY_W)
    except FitError:
        pass
    else:
        raise AssertionError('過密な補足欄が停止していません')
    spec['footnote'] = {'text': '短い補足'}
    def failing(*args):
        raise FitError('本文が収まりません')
    try:
        generate.render_slide(failing, slide, spec, 1)
    except SystemExit:
        pass
    assert generate._CONTENT_BOTTOM.get() == generate.BODY_BOTTOM
    for value in ('誤った形式', {}, {'text': ''}, {'text': '補足', 'label': ''}, {'text': '補足', 'x': 1}):
        spec['footnote'] = value
        assert validate({'meta': {'title': '検証'}, 'slides': [spec]}, allow_sample_content=True)
    legacy = deepcopy(specimens['concept'])
    legacy.pop('footnote')
    legacy['misconception'] = '旧補足'
    errors = validate({'meta': {'title': '検証'}, 'slides': [legacy]}, allow_sample_content=True)
    assert any('misconception' in error and '未対応' in error for error in errors)
    print('共通補足欄: 全type・収容3段階・明示停止・旧フィールド拒否を確認')


if __name__ == '__main__':
    main()
