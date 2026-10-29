"""文章役割の保存・復元と他の拡張情報の保持を確認する。"""
from io import BytesIO
import unittest

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from slidegen.text_roles import set_text_role, text_role


class TextRoleTests(unittest.TestCase):
    def test_shape_and_cell_roles_survive_save(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        shape.name = "references:link[1]"
        shape.text = "表示文字列"
        table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(4), Inches(1))
        cell = table.table.cell(0, 0)
        cell.text = "セル本文"
        set_text_role(shape.text_frame, "natural")
        set_text_role(cell.text_frame, "compact")
        stream = BytesIO()
        prs.save(stream)
        stream.seek(0)
        restored = Presentation(stream).slides[0]
        self.assertEqual(text_role(restored.shapes[0].text_frame), "natural")
        self.assertEqual(restored.shapes[0].name, "references:link[1]")
        self.assertEqual(restored.shapes[0].text, "表示文字列")
        self.assertEqual(text_role(restored.shapes[1].table.cell(0, 0).text_frame), "compact")

    def test_replace_preserves_other_extensions(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        body = shape.text_frame._txBody.bodyPr
        extensions = body.makeelement(qn("a:extLst"), {})
        other = body.makeelement(qn("a:ext"), {"uri": "other-test-extension"})
        extensions.append(other)
        body.append(extensions)
        self.assertEqual(text_role(shape.text_frame), "body")
        set_text_role(shape.text_frame, "natural")
        set_text_role(shape.text_frame, "compact")
        self.assertEqual(text_role(shape.text_frame), "compact")
        self.assertEqual(len(extensions), 2)
        self.assertIs(extensions[0], other)
        with self.assertRaises(ValueError):
            set_text_role(shape.text_frame, "unknown")


if __name__ == "__main__":
    unittest.main()
