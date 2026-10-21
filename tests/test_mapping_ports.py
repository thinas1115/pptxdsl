"""対応線の端点分散・意味の保持・密度上限での明示停止を検証する。"""
from collections import defaultdict
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation
from pptx.util import Inches

from slidegen.candidate_renderers import ACCENT, GRAY, _mapping_ports, s_mapping
from slidegen.check_layout import check
from slidegen.layout_fit import FitError
from slidegen.validate_content import validate
from tools.gallery.generate_mapping_cases import main, make_deck


class MappingPortTests(unittest.TestCase):
    def test_ports_follow_other_item_order_and_reject_overcrowding(self):
        left, right = {"a": 1.0, "b": 2.0}, {"x": 1.0, "y": 2.0}
        links = [{"from": "a", "to": "y"}, {"from": "a", "to": "x"},
                 {"from": "b", "to": "x"}]
        ports = _mapping_ports(left, right, .58, links)
        self.assertLess(ports[1][0], ports[0][0])
        self.assertLess(ports[1][1], ports[2][1])
        for side, key, positions in ((0, "from", left), (1, "to", right)):
            for link, pair in zip(links, ports):
                self.assertGreaterEqual(pair[side], positions[link[key]] + .10 - 1e-6)
                self.assertLessEqual(pair[side], positions[link[key]] + .48 + 1e-6)
        with self.assertRaisesRegex(FitError, "分割"):
            _mapping_ports(left, right, .22, links)

    def test_saved_endpoints_are_distinct_and_keep_each_relation(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "mapping.pptx"
            main(output)
            self.assertFalse(check(output))
            for slide, spec in zip(Presentation(output).slides, make_deck()["slides"]):
                shapes = {s.name: s for s in slide.shapes}
                connected = defaultdict(list)
                for index, link in enumerate(spec["links"]):
                    arrow = shapes[f"mapping:link:{index}"]
                    self.assertEqual(arrow.line.color.rgb, ACCENT if link.get("emphasis") else GRAY)
                    self.assertIn('type="triangle"', arrow._element.xml)
                    for side, key in (("left", "from"), ("right", "to")):
                        port = shapes[f"mapping:port:{side}:{index}"]
                        item = shapes[f"mapping:item:{side}:{link[key]}"]
                        cy = port.top + port.height / 2
                        self.assertAlmostEqual((arrow.begin_y if side == "left" else arrow.end_y) / 914400,
                                               cy / 914400, places=5)
                        self.assertGreater(cy, item.top)
                        self.assertLess(cy, item.top + item.height)
                        self.assertAlmostEqual((port.left + port.width / 2) / 914400,
                                               (item.left + item.width if side == "left" else item.left) / 914400,
                                               places=5)
                        connected[(side, link[key])].append(cy / 914400)
                for endpoints in connected.values():
                    endpoints.sort()
                    self.assertTrue(all(b - a >= .055 - 1e-5 for a, b in zip(endpoints, endpoints[1:])))
                for side, key in (("left", "left_items"), ("right", "right_items")):
                    for item in spec[key]:
                        text = shapes[f"mapping:text:{side}:{item['id']}"]
                        self.assertEqual(text.text.replace("\n", ""), item["text"])
                        self.assertTrue(all(run.font.size.pt >= 9.5 for p in text.text_frame.paragraphs for run in p.runs))

    def test_oversized_names_and_link_limit_have_explicit_failure(self):
        spec = deepcopy(make_deck()["slides"][0])
        spec["left_items"][0]["text"] = "確認事項" * 150
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        with self.assertRaisesRegex(FitError, "分割"):
            s_mapping(prs.slides.add_slide(prs.slide_layouts[6]), spec, 1)
        deck = make_deck()
        deck["slides"] = [deepcopy(deck["slides"][4])]
        deck["slides"][0]["links"].append({"from": "l5", "to": "r2"})
        self.assertTrue(any("10" in error for error in validate(deck, allow_sample_content=True)))
        self.assertTrue(any("分割" in error for error in validate(deck, allow_sample_content=True)))


if __name__ == "__main__":
    unittest.main()
