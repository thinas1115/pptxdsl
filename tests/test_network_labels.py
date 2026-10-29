"""ネットワーク図のマスク実測幅・物理/論理線・凡例・収容を検証する。"""
from copy import deepcopy
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.check_layout import check
from slidegen.layout_fit import FitError
from slidegen.textfit import text_width_in
from slidegen.training_renderers import _fit_network, _network_legend, _node_label, s_network
from tools.gallery.generate_network_cases import make_deck


def presentation():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
    return prs


class NetworkLabelTests(unittest.TestCase):
    def test_masks_are_measured_and_match_the_lane_background(self):
        prs = presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _node_label(slide, 5, 2, 1.68, "PC A", 10.5, sub="営業", fill=generate.LIGHT)
        for shape in slide.shapes:
            self.assertEqual(shape.fill.fore_color.rgb, generate.LIGHT)
            run = shape.text_frame.paragraphs[0].runs[0]
            widest = max(text_width_in(line, run.font.size.pt, "bold" if run.font.bold else "regular")
                         for line in shape.text.split("\n"))
            self.assertAlmostEqual(shape.width / 914400, widest + .08, places=5)
            self.assertLess(shape.width, Inches(1.2))
            self.assertFalse(shape.text_frame.word_wrap)

    def test_full_legend_wraps_within_body_width(self):
        rows, height = _network_legend([{"kind": kind} for kind in
                                       ("access", "trunk", "routed", "control", "broadcast", "blocked")])
        self.assertGreater(len(rows), 1)
        self.assertGreater(height, .24)
        for row in rows:
            self.assertLessEqual(sum(item[2] for item in row) + .24 * (len(row) - 1), generate.BODY_W - .16)

    def test_trunk_has_physical_line_and_matching_legend(self):
        spec = next(spec for spec in make_deck()["slides"] if any(link["kind"] == "trunk" for link in spec["links"]))
        prs = presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        generate.render_slide(s_network, slide, spec, 1)
        physical = [shape for shape in slide.shapes if shape.name.startswith("network-link:physical:")]
        self.assertTrue(physical)
        self.assertTrue(all(shape.line.color.rgb == generate.GRAY for shape in physical))
        self.assertTrue(any(shape.name == "network-legend:physical" for shape in slide.shapes))
        self.assertTrue(any(shape.name == "network-legend:trunk" for shape in slide.shapes))

    def test_public_network_cases_have_no_layout_findings(self):
        for spec in make_deck()["slides"]:
            prs = presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            generate.render_slide(s_network, slide, spec, 1)
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "network.pptx"
                prs.save(path)
                self.assertFalse(check(path), spec["title"])

    def test_fit_stages_and_minimum_failure(self):
        stages = set()
        for index in range(400):
            try:
                fit = _fit_network(SimpleNamespace(height=7.5 - index * .01), 4, 2, .58)
            except FitError:
                continue
            stages.add(fit.stage)
        self.assertEqual(stages, {"standard", "gap", "element", "font"})
        with self.assertRaises(FitError):
            _fit_network(SimpleNamespace(height=1), 4, 2, .58)
        spec = deepcopy(make_deck()["slides"][0])
        spec["lanes"][0]["sub"] = "収容不能な説明" * 1000
        prs = presentation()
        with self.assertRaises(FitError):
            s_network(prs.slides.add_slide(prs.slide_layouts[6]), spec, 1)


if __name__ == "__main__":
    unittest.main()
