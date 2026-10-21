"""点・名称・引出線の近接と、密集時の明示停止を保存結果で検証する。"""
from math import hypot
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from pptx import Presentation

from slidegen import diagrams2 as d
from slidegen.check_layout import check
from slidegen.layout_fit import FitError
from tools.gallery.generate_matrix_cases import main, make_deck


class MatrixLabelTests(unittest.TestCase):
    def test_pure_layout_distance_and_collision(self):
        for spec in make_deck()["slides"]:
            _stage, _width, xy, labels = d.fit_matrix_labels(spec["points"], 1.68, 2.22, 10.5, 4.02)
            leaders = {}
            for i, (px, py) in enumerate(xy):
                self.assertAlmostEqual(px, 1.68 + 10.5 * spec["points"][i]["x"])
                self.assertAlmostEqual(py, 2.22 + 4.02 * (1 - spec["points"][i]["y"]))
                edge = d._matrix_nearest_edge((px, py), labels[i])
                self.assertLessEqual(hypot(edge[0] - px, edge[1] - py), d.MATRIX_LABEL_DISTANCE)
                for j, other in labels.items():
                    if j != i:
                        self.assertFalse(d._rects_overlap(labels[i], other, 0))
                for ox, oy in xy:
                    self.assertFalse(d._rects_overlap(labels[i], (ox - .17, oy - .17, ox + .17, oy + .17), 0))
                leader = d._matrix_leader_box((px, py), labels[i])
                if leader:
                    leaders[i] = leader
            for i, leader in leaders.items():
                for j, rect in labels.items():
                    if i != j:
                        self.assertFalse(d._rects_overlap(leader, rect, 0))
                for j, other in leaders.items():
                    if i != j:
                        self.assertFalse(d._rects_overlap(leader, other, 0))

    def test_saved_labels_and_leaders(self):
        total_leaders = 0
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "matrix.pptx"
            main(output)
            self.assertFalse(check(output))
            for slide, spec in zip(Presentation(output).slides, make_deck()["slides"]):
                for i, point in enumerate(spec["points"]):
                    dot = next(s for s in slide.shapes if s.name == f"matrix:point:{i}")
                    label = next(s for s in slide.shapes if s.name == f"matrix:label:{i}")
                    self.assertEqual(label.text, point["name"])
                    self.assertFalse(label.text_frame.word_wrap)
                    self.assertGreaterEqual(label.text_frame.paragraphs[0].runs[0].font.size.pt, 9)
                    self.assertAlmostEqual(dot.left.inches + dot.width.inches / 2, 1.68 + 10.5 * point["x"], places=4)
                    xy = (dot.left.inches + dot.width.inches / 2, dot.top.inches + dot.height.inches / 2)
                    rect = (label.left.inches, label.top.inches,
                            label.left.inches + label.width.inches, label.top.inches + label.height.inches)
                    edge = d._matrix_nearest_edge(xy, rect)
                    distance = hypot(edge[0] - xy[0], edge[1] - xy[1])
                    self.assertLessEqual(distance, d.MATRIX_LABEL_DISTANCE + .0001)
                    leaders = [s for s in slide.shapes if s.name == f"matrix:leader:{i}"]
                    self.assertEqual(len(leaders), int(distance > .26))
                    if leaders:
                        total_leaders += 1
                        self.assertAlmostEqual(leaders[0].line.width.pt, .7)
                        self.assertFalse(leaders[0]._element.xpath('.//a:headEnd | .//a:tailEnd'))
            self.assertGreater(total_leaders, 0)

    def test_fallback_order_keeps_a_real_minimum_layout(self):
        original = d._matrix_label_positions
        widths = []
        def constrained(*args, **kwargs):
            widths.append((kwargs["label_w"], kwargs["gap"]))
            if len(widths) < 3:
                raise FitError("先行配置が収まらない場合の契約検証")
            return original(*args, **kwargs)
        with patch.object(d, "_matrix_label_positions", side_effect=constrained):
            stage, width, _xy, _labels = d.fit_matrix_labels(make_deck()["slides"][3]["points"], 1.68, 2.22, 10.5, 4.02)
        self.assertEqual(widths, [(2., .05), (2., .015), (1.65, .015)])
        self.assertEqual((stage, width), ("element", 1.65))

    def test_impossible_dense_and_long_labels_stop(self):
        for points in (
            [{"name": "検証対象方式", "x": .5, "y": .5} for _ in range(8)],
            [{"name": "過密で読めない長い名称" * 20, "x": .5, "y": .5}],
        ):
            with self.assertRaises(FitError):
                d.fit_matrix_labels(points, 1.68, 2.22, 10.5, 4.02)


if __name__ == "__main__":
    unittest.main()
