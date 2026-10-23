"""配布ZIPの自己完結性と、cloneから切り離した実行を検証する。"""
import hashlib
import json
import os
from pathlib import Path
import subprocess
import shutil
import sys
import tempfile
import unittest
import zipfile

from PIL import Image
from pptx import Presentation

from tools.plugin.build import PROJECT_FILES, ROOT, build
from tests.test_pptxdsl_skill import _broken_local_links, _smoke_deck


def _hashes(root: Path, *, source_text: bool = False) -> dict[str, str]:
    def content(path: Path) -> bytes:
        data = path.read_bytes()
        if source_text and (path.suffix.lower() in {".md", ".json", ".svg", ".txt"} or path.name == "LICENSE"):
            return data.replace(b"\r\n", b"\n").replace(b"\r", b"\n")
        return data
    return {path.relative_to(root).as_posix(): hashlib.sha256(content(path)).hexdigest()
            for path in root.rglob("*") if path.is_file()}


def _plugin_smoke_deck() -> dict:
    deck = _smoke_deck()
    deck["slides"].extend([
        {"type": "section_divider", "kicker": "第2章", "title": "ネットワーク構成",
         "lead": "VPCとsubnetの境界"},
        {"type": "aws_vpc_layout", "kicker": "構成", "title": "VPC構成",
         "vpc": {"label": "VPC", "cidr": "10.0.0.0/16"},
         "external": [{"id": "user", "label": "利用者", "icon": "icons/aws/users.png"}],
         "azs": [{"id": "az-a", "label": "AZ-a", "subnets": [
             {"id": "private", "label": "private subnet", "resources": [
                 {"id": "service", "label": "Lambda", "icon": "icons/aws/lambda.png"}]}]}],
         "flows": [{"from": "user", "to": "service", "label": "HTTPS"}]},
    ])
    return deck


def _run(script: Path, cwd: Path, *args: str, success: bool = True):
    env = os.environ.copy()
    env.update(PYTHONIOENCODING="utf-8", PYTHONDONTWRITEBYTECODE="1", PYTHONPATH="")
    result = subprocess.run([sys.executable, "-I", str(script), *args], cwd=cwd,
                            env=env, capture_output=True, text=True, encoding="utf-8")
    assert (result.returncode == 0) == success, result.stdout + result.stderr
    return result


def _check_archive_and_runtime() -> None:
    with tempfile.TemporaryDirectory(prefix="pptxdsl-plugin-") as temporary:
        base = Path(temporary)
        plugin, archive = build(base / "build")
        _, second = build(base / "repeat")
        assert archive.read_bytes() == second.read_bytes(), "同じ入力から異なるZIPが生成されました"
        original = _hashes(plugin)
        try:
            build(base / "build")
        except ValueError:
            pass
        else:
            raise AssertionError("既存配布物の上書きを拒否しませんでした")
        assert original == _hashes(plugin)

        installed = base / "installed" / "pptxdsl"
        with zipfile.ZipFile(archive) as bundle:
            assert "plugin.json" in bundle.namelist()
            assert not any("split_layout" in name or "split_renderer" in name
                           or "split-layout" in name for name in bundle.namelist())
            assert not any(set(Path(name).parts) & {"tests", "tools", ".git", "out", "__pycache__"}
                           for name in bundle.namelist())
            bundle.extractall(installed)
        portable = json.loads((installed / "plugin.json").read_text(encoding="utf-8"))
        compatibility = json.loads((installed / ".codex-plugin/plugin.json").read_text(encoding="utf-8"))
        assert portable["name"] == compatibility["name"] == installed.name
        assert portable["version"] == compatibility["version"]
        assert portable["extensions"]["com.openai"]["interface"] == compatibility["interface"]
        assert "mcpServers" not in portable and "apps" not in portable
        skill = installed / "skills/pptxdsl"
        project = skill / "project"
        for markdown in installed.rglob("*.md"):
            assert not _broken_local_links(markdown), markdown.relative_to(installed)
        assert _hashes(ROOT / "slidegen/assets", source_text=True) == _hashes(project / "slidegen/assets")
        assert (project / "LICENSE").is_file()
        assert (project / "THIRD_PARTY_NOTICES.md").is_file()
        assert (project / "slidegen/data/sample_fingerprints.json").is_file()

        before = _hashes(installed)
        task = base / "task with spaces"
        task.mkdir()
        deck = _plugin_smoke_deck()
        (task / "content.json").write_text(json.dumps(deck, ensure_ascii=False), encoding="utf-8")
        runner = skill / "scripts/run.py"
        example = project / "examples/authoring/type-selection.json"
        assert example.read_text(encoding="utf-8") == (ROOT / "examples/authoring/type-selection.json").read_text(encoding="utf-8")
        _run(runner, task, "validate", str(example))
        _run(runner, task, "generate", str(example), "out/authoring.pptx")
        _run(runner, task, "check", "out/authoring.pptx")
        assert len(Presentation(task / "out/authoring.pptx").slides) == 4
        image_example = project / "examples/authoring/image-reading.json"
        _run(runner, task, "validate", str(image_example))
        _run(runner, task, "generate", str(image_example), "out/image-reading.pptx")
        _run(runner, task, "check", "out/image-reading.pptx")
        assert len(Presentation(task / "out/image-reading.pptx").slides) == 4
        _run(runner, task, "validate", "content.json")
        _run(runner, task, "generate", "content.json", "out/deck.pptx")
        _run(runner, task, "check", "out/deck.pptx")
        fonts = json.loads(_run(runner, task, "fonts").stdout)
        assert fonts["regular"]["family"] == fonts["bold"]["family"]
        assert fonts["regular"]["style"] == "Regular" and fonts["bold"]["style"] == "Bold"
        assert all("path" not in face for face in fonts.values())
        assert len(Presentation(task / "out/deck.pptx").slides) == len(deck["slides"])
        details = json.loads(_run(runner, task, "probe").stdout)
        # Windowsの一時ディレクトリは短縮名・長い名前のどちらでも同じ場所を指す。
        assert Path(details["powerpoint"]["render_script"]).resolve() == (project / "render.ps1").resolve()
        available = details["available_backends"]
        required = os.environ.get("PPTXDSL_PLUGIN_REQUIRE_RENDER") == "1"
        assert available or not required, "Pluginの実レンダリング手段がありません"
        if required:
            backend = "libreoffice" if "libreoffice" in available else "powerpoint"
            configured = os.environ.get("PPTXDSL_PLUGIN_SMOKE_OUT")
            png = (ROOT / configured).resolve() if configured else task / "out/png"
            _run(runner, task, "render", "out/deck.pptx", str(png), "--backend", backend)
            _run(runner, task, "sheet", str(png))
            assert len(list(png.glob("slide_*.png"))) == len(deck["slides"])
            with Image.open(png / "sheet.png") as image:
                assert image.size[0] > 0 and image.size[1] > 0
            print(f"Plugin実レンダリング: {backend}")
        else:
            print("SKIP: Plugin実レンダリング: PPTXDSL_PLUGIN_REQUIRE_RENDER未指定")
        (task / "invalid.json").write_text("{", encoding="utf-8")
        failure = _run(runner, task, "generate", "invalid.json", "out/invalid.pptx", success=False)
        assert not (task / "out/invalid.pptx").exists()
        assert "Traceback" not in failure.stderr
        _run(runner, task, "generate", "content.json", success=False)
        _run(runner, task, "render", "missing.pptx", "out/missing-png", success=False)
        assert not (task / "out/missing-png").exists()
        assert before == _hashes(installed), "インストール先に生成物が書き込まれました"


def _check_untracked_files_excluded() -> None:
    with tempfile.TemporaryDirectory(prefix="pptxdsl-plugin-source-") as temporary:
        source = Path(temporary) / "source"
        for relative in (*PROJECT_FILES, "plugins/pptxdsl/.codex-plugin/plugin.json",
                         "plugins/pptxdsl/.claude-plugin/plugin.json",
                         ".agents/skills/pptxdsl/SKILL.md", "slidegen/generate_from_json.py"):
            target = source / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(ROOT / relative, target)
        subprocess.run(["git", "init", "-q", str(source)], check=True)
        subprocess.run(["git", "-C", str(source), "add", "slidegen"], check=True)
        private = source / "slidegen/private.local"
        private.write_text("利用者のローカル入力", encoding="utf-8")
        _, archive = build(Path(temporary) / "package", source_root=source)
        with zipfile.ZipFile(archive) as bundle:
            assert not any(name.endswith("private.local") for name in bundle.namelist())
            assert "skills/pptxdsl/project/slidegen/generate_from_json.py" in bundle.namelist()


class PluginPackageTests(unittest.TestCase):
    def test_untracked_files_excluded(self):
        _check_untracked_files_excluded()

    def test_archive_and_runtime(self):
        _check_archive_and_runtime()


if __name__ == "__main__":
    unittest.main()
