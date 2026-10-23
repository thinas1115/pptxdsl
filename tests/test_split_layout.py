"""半幅の独立収容、共有部分の一回描画、拒否条件と保存前停止を検証する。"""

from copy import deepcopy
import json
import os
from pathlib import Path
import tempfile
import subprocess
import sys
import unittest
from unittest.mock import patch

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from slidegen import generate as g
from slidegen.check_layout import check
from slidegen.diagram_layout import Layout
from slidegen.generate_from_json import build_presentation
from slidegen.generation_context import current_context
from slidegen.layout_fit import FitError
from slidegen.schema_definition import SPLIT_PAIRS, split_block_schemas
from slidegen.split_layout import DrawingRegion, plan_split
from slidegen.split_renderer import ADAPTERS
from slidegen.validate_content import validate, validation_warnings
from tests.fixtures.gallery.split_patterns import SPLIT_SLIDES
from tools.gallery.generate_split_cases import main, make_deck


def deck(spec):
    return {"meta": {"title": "複合の試験", "footer": "共通のフッター"}, "slides": [spec]}


def render(spec):
    return build_presentation(deck(spec), allow_sample_content=True)


class SplitLayoutTests(unittest.TestCase):
    def test_plan_is_equal_width_disjoint_and_does_not_mutate_area(self):
        area = g.ContentArea(1.8, 6.5, True)
        before = deepcopy(area)
        plan = plan_split(area, shifted=True, has_note=True)
        self.assertEqual(area, before)
        self.assertEqual(plan.left.width, plan.right.width)
        self.assertAlmostEqual(plan.right.left - plan.left.right, plan.gutter)
        self.assertAlmostEqual(plan.left.bottom, 6.2)
        self.assertTrue(plan.left.shifted)
        self.assertEqual(plan, plan_split(area, shifted=True, has_note=True))
        for arguments in ((0, 2, 0, 3), (0, 3, 2, 2), (0, 2, float("inf"), 3)):
            with self.assertRaises(FitError):
                DrawingRegion(*arguments)

    def test_declared_pairs_and_body_adapters_match(self):
        self.assertEqual(set(ADAPTERS), set(split_block_schemas()))
        expected = set(SPLIT_PAIRS) | {(right, left) for left, right in SPLIT_PAIRS}
        data = make_deck()
        pairs = {(spec["left"]["type"], spec["right"]["type"]) for spec in data["slides"]}
        self.assertEqual(pairs, expected)
        self.assertFalse(validate(data, allow_sample_content=True))

    def test_thirteen_saved_pages_have_no_layout_findings(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "split.pptx"
            main(output)
            self.assertEqual(len(Presentation(output).slides), 13)
            self.assertFalse(check(output))

    def test_shared_header_lead_note_footnote_and_footer_are_drawn_once(self):
        spec = deepcopy(make_deck()["slides"][-1])
        spec["lead"] = "共通の説明文を一度だけ表示します。"
        saved = render(spec).slides[0]
        texts = [shape.text for shape in saved.shapes if shape.has_text_frame]
        for value in (spec["title"], spec["kicker"], spec["lead"], spec["note"],
                      spec["footnote"]["text"], "共通のフッター"):
            self.assertEqual(texts.count(value), 1, value)
        backgrounds = [shape for shape in saved.shapes if shape.left == 0 and shape.top == 0
                       and shape.width.inches > 13 and shape.height.inches > 7]
        self.assertEqual(len(backgrounds), 1)

    def test_unknown_renderer_recursion_shared_child_fields_and_coordinates_are_rejected(self):
        for field, value in (("type", "org"), ("type", "split"), ("width", 5),
                             ("title", "子のタイトル"), ("lead", "子の説明"),
                             ("note", "子の注記"), ("footnote", {"text": "子の脚注"})):
            spec = deepcopy(SPLIT_SLIDES[0])
            spec["left"][field] = value
            errors = validate(deck(spec), allow_sample_content=True)
            self.assertTrue(errors, (field, value))
            self.assertTrue(any("left" in error for error in errors), errors)
        spec = deepcopy(SPLIT_SLIDES[0])
        spec["right"] = deepcopy(spec["left"])
        self.assertTrue(validate(deck(spec), allow_sample_content=True), "diagram同士は未対応")

    def test_child_assets_and_references_still_require_semantic_validation(self):
        spec = deepcopy(SPLIT_SLIDES[1])
        spec["left"]["image"] = "images/missing-split-fixture.png"
        self.assertTrue(any("left.image" in error and "assets" in error
                            for error in validate(deck(spec), allow_sample_content=True)))
        spec = deepcopy(SPLIT_SLIDES[0])
        spec["left"]["diagram"]["edges"][0]["to"] = "missing"
        self.assertTrue(any("left.edges" in error and "見つかりません" in error
                            for error in validate(deck(spec), allow_sample_content=True)))

    def test_one_side_fitting_does_not_resize_other_side(self):
        original = deepcopy(SPLIT_SLIDES[1])
        crowded = deepcopy(original)
        for item in crowded["right"]["bullets"]:
            item["text"] *= 3
        fitted = []
        original_fit = g.fit_vertical_stacks

        def observe(*args, **kwargs):
            result = original_fit(*args, **kwargs)
            fitted.append(result)
            return result

        with patch.object(g, "fit_vertical_stacks", side_effect=observe):
            baseline, result = render(original), render(crowded)
        def images(prs):
            return [(shape.left, shape.top, shape.width, shape.height,
                     shape.crop_left, shape.crop_right, shape.image.sha1)
                    for shape in prs.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(images(baseline), images(result))
        self.assertLess(fitted[-1].size, fitted[0].size)
        self.assertGreaterEqual(fitted[-1].size, 12)

    def test_left_right_and_both_crowded_inputs_stop_without_mutating_input_or_context(self):
        for side in ("left", "right", "both"):
            spec = deepcopy(SPLIT_SLIDES[0])
            if side in ("left", "both"):
                diagram = spec["left"]["diagram"]
                diagram["rows"] = [f"r{index}" for index in range(20)]
                diagram["containers"] = []
                diagram["nodes"] = {f"n{index}": {"col": "入力", "row": f"r{index}",
                                               "title": "確認先", "icon": "icons/fluent/app.png"}
                                    for index in range(20)}
                diagram["edges"] = [{"from": "n0", "to": "n1"}]
            if side in ("right", "both"):
                spec["right"]["cards"][0]["body"] *= 1000
            before, context = deepcopy(spec), current_context()
            self.assertFalse(validate(deck(spec), allow_sample_content=True))
            with self.assertRaisesRegex(FitError, r"split\.(left|right)"):
                render(spec)
            self.assertEqual(spec, before)
            self.assertIs(current_context(), context)

    def test_region_overflow_is_rejected_not_clipped_or_shifted(self):
        def escaped(slide, _spec, region):
            g.add_rect(slide, region.left - .2, region.top, .4, .4, g.NAVY)
        with patch.dict(ADAPTERS, {"bullets": escaped}):
            with self.assertRaisesRegex(FitError, r"split.right.*領域"):
                render(deepcopy(SPLIT_SLIDES[1]))

    def test_legacy_child_bullets_keep_deprecation_warning(self):
        spec = deepcopy(SPLIT_SLIDES[1])
        spec["right"]["bullets"] = [["項目の確認", None]]
        self.assertFalse(validate(deck(spec), allow_sample_content=True))
        self.assertTrue(any("right.bullets" in message and "2.0.0" in message
                            for message in validation_warnings(deck(spec))))

    def test_half_diagram_uses_spacing_and_icon_stages_before_stop(self):
        diagram = deepcopy(SPLIT_SLIDES[0]["left"]["diagram"])
        observed, stopped = set(), False
        for step in range(80):
            region = DrawingRegion(.55, 1.66, 5.88, 6.85 - step * .05)
            try:
                observed.add(Layout(diagram, content_area=region).fit_stage)
            except FitError:
                stopped = True
        self.assertEqual(observed, {"standard", "gap", "icon"})
        self.assertTrue(stopped)

    def test_overcrowded_cli_leaves_existing_output_untouched_and_has_safe_errors(self):
        spec = {"type": "split", "kicker": "収容", "title": "密度の確認",
                "left": {"type": "image", "image": "images/pptxdsl-repository.png"},
                "right": {"type": "bullets", "bullets": [{"text": "独立した本文領域への収容条件を検査します。" * 1000}]}}
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "input.json", Path(directory) / "existing.pptx"
            source.write_text(json.dumps(deck(spec), ensure_ascii=False), encoding="utf-8")
            output.write_bytes(b"existing-artifact")
            result = subprocess.run([sys.executable, "-m", "slidegen.generate_from_json", str(source), str(output)],
                                    cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True,
                                    encoding="utf-8", env=dict(os.environ, PYTHONIOENCODING="utf-8"), check=False)
            messages = result.stdout + result.stderr
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("split.right", messages)
            self.assertNotIn("Traceback", messages)
            self.assertNotIn(directory, messages)
            self.assertEqual(output.read_bytes(), b"existing-artifact")


if __name__ == "__main__":
    unittest.main()
