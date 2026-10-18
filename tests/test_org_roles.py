"""体制図の名称・役割・担当区分の可読性と明示停止を検証する。"""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation

from slidegen.check_layout import check, contrast_ratio
from slidegen.generate import ContentArea, NAVY, TEXT, LIGHT, SURFACE
from slidegen.layout_fit import FitError
from slidegen.org_layout import OrgLayout, fit_org_layout, s_org
from tools.gallery.generate_org_cases import main, make_deck


class OrgRoleTests(unittest.TestCase):
    def test_names_only_keep_hierarchy_and_connection_geometry(self):
        specs = make_deck()["slides"]
        names = OrgLayout(specs[0]["org"], ContentArea())
        roles = OrgLayout(specs[1]["org"], ContentArea())
        self.assertEqual(names.boxes, roles.boxes)
        self.assertEqual(names.routes, roles.routes)
        self.assertEqual(roles.fit.values["sub_pt"], 11.5)
        self.assertGreaterEqual(contrast_ratio(tuple(TEXT), tuple(SURFACE)), 4.5)
        self.assertGreaterEqual(contrast_ratio(tuple(LIGHT), tuple(NAVY)), 4.5)

    def test_saved_roles_members_and_minimum_sizes(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "org.pptx"
            main(output)
            self.assertFalse(check(output))
            for slide, spec in zip(Presentation(output).slides, make_deck()["slides"]):
                shapes = {s.name: s for s in slide.shapes}
                for key, node in spec["org"]["nodes"].items():
                    box = shapes[f"org:box:{key}"]
                    for kind, value, minimum in (("name", node["name"], 10.5), ("sub", node.get("sub"), 9.5),
                                                  ("members", " / ".join(node.get("members", [])), 9.5)):
                        if not value:
                            self.assertNotIn(f"org:{kind}:{key}", shapes)
                            continue
                        text = shapes[f"org:{kind}:{key}"]
                        self.assertEqual(text.text.replace("\n", ""), value.replace("\n", ""))
                        for paragraph in text.text_frame.paragraphs:
                            for run in paragraph.runs:
                                self.assertGreaterEqual(run.font.size.pt, minimum)
                        self.assertGreaterEqual(text.top, box.top)
                        self.assertLessEqual(text.top + text.height, box.top + box.height)

    def test_compression_keeps_role_floor_and_then_stops(self):
        org = make_deck()["slides"][5]["org"]
        self.assertEqual(fit_org_layout(org, 6).stage, "standard")
        self.assertEqual(fit_org_layout(org, 5).stage, "gap")
        self.assertEqual(fit_org_layout(org, 4).stage, "element")
        minimum = fit_org_layout(org, 3.5)
        self.assertEqual(minimum.stage, "element")
        self.assertGreaterEqual(minimum.values["sub_pt"], 9.5)
        self.assertGreaterEqual(minimum.values["members_pt"], 9.5)
        with self.assertRaises(FitError):
            fit_org_layout(org, 3)

    def test_oversized_role_or_members_stop_without_tiny_text(self):
        for field in ("sub", "members"):
            spec = deepcopy(make_deck()["slides"][4])
            spec["org"]["nodes"]["owner"][field] = ("説明文" * 100 if field == "sub" else ["担当者" * 100])
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            with self.assertRaises(FitError):
                s_org(slide, spec, 1)


if __name__ == "__main__":
    unittest.main()
