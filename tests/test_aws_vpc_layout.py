"""AWS階層図の段階的収容と、ラベルによる枠線の隠れを検証する。"""
from copy import deepcopy

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.aws_vpc_layout import _layout, _resource_slot, s_aws_vpc_layout
from slidegen.layout_fit import FitError
from slidegen.validate_content import validate
from tests.test_ai_content_contract import _deck
from tests.fixtures.gallery.content_patterns import PATTERN_DECK


def main():
    spec = {"type": "aws_vpc_layout", "kicker": "構成", "title": "VPC構成",
            "vpc": {"label": "VPC"}, "azs": [{"id": "az", "label": "AZ-a",
            "subnets": [{"id": f"sub-{i}", "label": "private subnet",
                         "resources": [{"id": f"node-{i}", "label": "Lambda",
                                        "icon": "icons/aws/lambda.png"}]}
                        for i in range(3)]}]}
    for height, stage in ((4.8, "standard"), (4.7, "gap"), (4.5, "element")):
        _, _, subnets, fit = _layout(generate.ContentArea(2.0, 2.0 + height, True), spec)
        assert fit["stage"] == stage, fit
        for subnet in subnets.values():
            _resource_slot(subnet, 0, 1, icon_size=fit["icon_size"])
    try:
        _layout(generate.ContentArea(2.0, 6.2, True), spec)
    except FitError as error:
        assert "不足" in str(error)
    else:
        raise AssertionError("最小アイコンでも収まらない図を拒否しませんでした")

    with_sub = deepcopy(spec)
    for subnet in with_sub["azs"][0]["subnets"]:
        subnet["resources"][0]["sub"] = "関数"
    _, _, subnets, fit = _layout(generate.ContentArea(1.5, 6.9), with_sub)
    for subnet in subnets.values():
        _resource_slot(subnet, 0, 1, icon_size=fit["icon_size"], has_sub=True)

    assert not validate(_deck(spec))
    invalids = []
    for field, value in (("vpc", {"label": "VPC", "x": 1}), ("azs", []),
                         ("external", [{}] * 3), ("flows", [{}] * 11)):
        invalid = deepcopy(spec)
        invalid[field] = value
        invalids.append(invalid)
    for change in ("duplicate", "subnet-count", "resource-count", "missing-icon"):
        invalid = deepcopy(spec)
        subnet = invalid["azs"][0]["subnets"][0]
        if change == "duplicate":
            subnet["resources"][0]["id"] = invalid["azs"][0]["id"]
        elif change == "subnet-count":
            invalid["azs"][0]["subnets"].append(deepcopy(subnet))
        elif change == "resource-count":
            subnet["resources"] *= 3
        else:
            subnet["resources"][0]["icon"] = "icons/aws/missing.png"
        invalids.append(invalid)
    for flow in ({"from": "node-0", "to": "missing"},
                 {"from": "node-0", "to": "node-0"},
                 {"from": "node-0", "to": "node-1", "dash": "solid"},
                 {"from": "node-0", "to": "node-1", "both": "yes"}):
        invalid = deepcopy(spec)
        invalid["flows"] = [flow]
        invalids.append(invalid)
    for invalid in invalids:
        assert validate(_deck(invalid)), invalid

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    gallery = next(item for item in PATTERN_DECK["slides"] if item["type"] == "aws_vpc_layout")
    s_aws_vpc_layout(slide, gallery, 1)
    for shape in slide.shapes:
        if shape.name.startswith("aws-vpc-layout:") and shape.name.endswith(":label"):
            assert shape.text_frame.word_wrap is False
    frames = [shape for shape in slide.shapes if shape.name.startswith("aws-vpc-layout:")
              and ":label" not in shape.name and shape.name != "aws-vpc-layout:flow-label"]
    labels = [shape for shape in slide.shapes if shape.name == "aws-vpc-layout:flow-label"]
    assert len(labels) == len([flow for flow in gallery["flows"] if flow.get("label")])
    pad = Inches(0.01)
    for label in labels:
        left, right = label.left, label.left + label.width
        top, bottom = label.top, label.top + label.height
        for frame in frames:
            if top < frame.top + frame.height and bottom > frame.top:
                assert not any(left - pad < edge < right + pad
                               for edge in (frame.left, frame.left + frame.width)), label.text
            if left < frame.left + frame.width and right > frame.left:
                assert not any(top - pad < edge < bottom + pad
                               for edge in (frame.top, frame.top + frame.height)), label.text
    print("OK: AWS VPC図の標準・余白圧縮・縮小・明示停止、通信ラベルの枠線保護")


if __name__ == "__main__":
    main()
