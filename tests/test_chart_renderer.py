"""グラフの符号・ゼロ・ラベル・可読性・段階的収容の回帰。"""
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate as g
from slidegen.chart_renderer import MIN_TEXT_PT, plan_chart
from slidegen.check_layout import check
from slidegen.layout_fit import FitError
from slidegen.validate_content import validate


KINDS = ("bar", "column", "line", "stacked_bar", "stacked_column")


def chart(kind="bar", values=None):
    return {"kind": kind, "categories": ["対象A", "対象B", "対象C"],
            "series": [["実績", values if values is not None else [-25, 0, 50]]]}


def render(data):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(g.SLIDE_W), Inches(g.SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    g.s_chart(slide, {"type": "chart", "kicker": "数値検証", "title": "実績比較", "chart": data}, 1)
    return prs, slide


class ChartRendererTests(unittest.TestCase):
    def test_every_kind_negative_mixed_and_zero(self):
        with TemporaryDirectory() as temporary:
            for kind in KINDS:
                for values in ([-25, -10, -5], [-25, 0, 50], [0, 0, 0]):
                    data = chart(kind, values)
                    data["show_values"] = True
                    prs, slide = render(data)
                    labels = [shape.text for shape in slide.shapes if shape.name == "chart:data-value"]
                    self.assertCountEqual(labels, [str(value) for value in values])
                    bars = [shape for shape in slide.shapes if shape.name == "chart:data-bar"]
                    self.assertEqual(len(bars), 0 if kind == "line" else sum(value != 0 for value in values))
                    self.assertTrue(all(shape.width > 0 and shape.height > 0 for shape in bars))
                    self.assertTrue(all(shape.text_frame.word_wrap is False for shape in slide.shapes
                                        if shape.has_text_frame and shape.text and shape.name.startswith("chart:")))
                    path = Path(temporary) / f"{kind}.pptx"
                    prs.save(path)
                    self.assertFalse(check(path), (kind, values, check(path)))

    def test_show_values_and_format_for_all_kinds(self):
        for kind in KINDS:
            for visible in (True, False):
                data = chart(kind, [0.25, 0.5, 0.75])
                data.update(show_values=visible, number_format="0.0%")
                _, slide = render(data)
                labels = [shape.text for shape in slide.shapes if shape.name == "chart:data-value"]
                self.assertCountEqual(labels, ["25.0%", "50.0%", "75.0%"] if visible else [])

    def test_signed_stacks_and_zero_shapes(self):
        for kind in ("stacked_bar", "stacked_column"):
            data = {"kind": kind, "categories": ["対象"],
                    "series": [["A", [30]], ["B", [-25]], ["C", [0]], ["D", [10]]], "show_values": False}
            plan = plan_chart(data, g.ContentArea())
            axis = next(line for line in plan.lines if line[-1] == "zero-axis")
            bars = [bar for bar in plan.bars if bar[-1] == "data-bar"]
            self.assertEqual(len(bars), 3)
            if kind == "stacked_bar":
                self.assertGreaterEqual(bars[0][0], axis[0] - 1e-6)
                self.assertLess(bars[1][0], axis[0])
                self.assertAlmostEqual(bars[2][0], bars[0][0] + bars[0][2])
            else:
                self.assertLess(bars[0][1], axis[1])
                self.assertGreaterEqual(bars[1][1], axis[1] - 1e-6)
                self.assertAlmostEqual(bars[2][1] + bars[2][3], bars[0][1])

    def test_normal_density_range_and_minimum(self):
        for kind in KINDS:
            for count in range(1, 13):
                data = {"kind": kind, "categories": [f"項目{i}" for i in range(count)],
                        "series": [["実績", [10] * count]], "show_values": False}
                plan = plan_chart(data, g.ContentArea())
                self.assertGreaterEqual(plan.size, MIN_TEXT_PT)
            for series_count in range(1, 5):
                data = {"kind": kind, "categories": ["項目"], "series": [[f"系列{i}", [10]] for i in range(series_count)],
                        "show_values": False}
                self.assertGreaterEqual(plan_chart(data, g.ContentArea()).size, MIN_TEXT_PT)

    def test_overdense_and_long_names_stop_explicitly(self):
        data = chart()
        data["categories"] = ["非常に長いカテゴリ名" * 30] * 3
        with self.assertRaises(FitError):
            plan_chart(data, g.ContentArea())
        data = chart()
        data["series"][0][0] = "長い系列名" * 50
        data["show_legend"] = True
        with self.assertRaises(FitError):
            plan_chart(data, g.ContentArea())
        data = {"kind": "bar", "categories": [f"項目{i}" for i in range(12)],
                "series": [[f"系列{i}", [100] * 12] for i in range(4)]}
        with self.assertRaisesRegex(FitError, "10.5pt"):
            plan_chart(data, g.ContentArea())

    def test_normal_gap_element_and_minimum_stages(self):
        for count, bottom, stage, size in ((4, 6.85, "normal", 12), (5, 6.85, "gap", 12),
                                           (7, 6.85, "element", 11), (7, 6.75, "element", 10.5)):
            data = {"kind": "bar", "categories": [f"項目{i}" for i in range(count)],
                    "series": [["実績", [10] * count], ["計画", [20] * count]], "show_values": False}
            plan = plan_chart(data, g.ContentArea(bottom=bottom))
            self.assertEqual((plan.stage, plan.size), (stage, size))

    def test_long_names_are_preserved_and_wrapped(self):
        names = ["東日本営業本部の法人営業担当における月次実績", "西日本営業本部の公共営業担当における月次実績"]
        data = {"kind": "bar", "categories": ["既存サービスの拡張案件", "新規サービスの導入案件"],
                "series": [[name, [10, 20]] for name in names], "show_values": False}
        plan = plan_chart(data, g.ContentArea())
        labels = [item[0] for item in plan.texts if item[-1] == "legend"]
        self.assertEqual([text.replace("\n", "") for text in labels], names)
        self.assertTrue(any("\n" in label for label in labels))

    def test_four_series_values_include_zero_in_stacks(self):
        with TemporaryDirectory() as temporary:
            for kind in KINDS:
                for values in ([30, -25, 0, 10], [0, 0, 0, 0]):
                    data = {"kind": kind, "categories": ["項目"],
                            "series": [[f"系列{i}", [value]] for i, value in enumerate(values)], "show_values": True}
                    prs, slide = render(data)
                    self.assertCountEqual([shape.text for shape in slide.shapes if shape.name == "chart:data-value"],
                                          [f"系列{i}: 0" if kind.startswith("stacked_") and value == 0 else str(value)
                                           for i, value in enumerate(values)])
                    path = Path(temporary) / f"{kind}.pptx"
                    prs.save(path)
                    self.assertFalse(check(path), (kind, values, check(path)))

    def test_zero_only_stack_labels_follow_series_order(self):
        for kind in ("stacked_bar", "stacked_column"):
            data = {"kind": kind, "categories": ["対象"], "show_values": True,
                    "series": [[f"系列{i}", [0]] for i in range(4)]}
            plan = plan_chart(data, g.ContentArea())
            labels = [item for item in plan.texts if item[-1] == "data-value"]
            self.assertEqual([item[0] for item in labels], [f"系列{i}: 0" for i in range(4)])
            positions = [item[2 if kind == "stacked_bar" else 1] for item in labels]
            self.assertEqual(positions, sorted(positions))
            self.assertEqual(len(set(positions)), 4)

    def test_validator_rejects_unsupported_format_nonfinite(self):
        data = {"meta": {"title": "実績"}, "slides": [{"type": "chart", "kicker": "分類", "title": "実績", "chart": chart()}]}
        for value in (float("nan"), float("inf"), True):
            invalid = deepcopy(data)
            invalid["slides"][0]["chart"]["series"][0][1][0] = value
            self.assertTrue(validate(invalid))
        data["slides"][0]["chart"]["number_format"] = "未対応"
        self.assertTrue(validate(data))
        for kind in ([], {}, None):
            invalid = deepcopy(data)
            invalid["slides"][0]["chart"]["kind"] = kind
            self.assertTrue(validate(invalid))


if __name__ == "__main__":
    unittest.main()
