"""Subnetラベルの文字コントラストと階層境界・下端の保持を検証する。"""
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.aws_vpc_layout import SUBNET_LINE
from slidegen.check_layout import check, contrast_ratio
from tools.gallery.generate_aws_vpc_cases import main, make_deck


class SubnetLabelTests(unittest.TestCase):
    def test_saved_labels_and_boundary_are_independent(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "subnets.pptx"
            main(path)
            self.assertFalse(check(path))
            for slide, spec in zip(Presentation(path).slides, make_deck()["slides"]):
                shapes = {shape.name: shape for shape in slide.shapes}
                vpc = shapes["aws-vpc-layout:vpc"]
                self.assertLessEqual((vpc.top + vpc.height) / Inches(1), generate.BODY_BOTTOM)
                for az in spec["azs"]:
                    for subnet in az["subnets"]:
                        name = f"aws-vpc-layout:subnet:{subnet['id']}"
                        frame, label = shapes[name], shapes[name + ":label"]
                        self.assertEqual(frame.line.color.rgb, SUBNET_LINE)
                        self.assertEqual(frame.line.width.pt, 0.9)
                        self.assertIn(subnet["label"], label.text)
                        self.assertIn(subnet["cidr"], label.text)
                        self.assertFalse(label.text_frame.word_wrap)
                        for paragraph in label.text_frame.paragraphs:
                            for run in paragraph.runs:
                                self.assertEqual(run.font.color.rgb, generate.NAVY)
                                self.assertGreaterEqual(contrast_ratio(run.font.color.rgb, generate.CANVAS), 4.5)
                        self.assertLess(label.top + label.height, frame.top + frame.height)


if __name__ == "__main__":
    unittest.main()
