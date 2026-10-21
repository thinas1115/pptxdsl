"""選択肢の列位置、モード間対応、縦並び、収容下限を確認する。"""
from copy import deepcopy
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest

from pptx import Presentation

from slidegen.check_layout import check
from slidegen.generate import ACCENT, GRAY
from slidegen.layout_fit import FitError
from slidegen.training_renderers import _fit_questions, s_knowledge_check
from tools.gallery.generate_knowledge_cases import main, make_deck


class KnowledgeChoiceTests(unittest.TestCase):
    def test_saved_columns_and_answers(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "knowledge.pptx"
            main(output)
            self.assertFalse(check(output))
            slides = list(Presentation(output).slides)
            for page, (slide, spec) in enumerate(zip(slides, make_deck()["slides"])):
                shapes = {s.name: s for s in slide.shapes}
                positions = {}
                for q, question in enumerate(spec["questions"]):
                    for i, text in enumerate(question["options"]):
                        shape = shapes[f"knowledge:option:{q}:{i}"]
                        self.assertEqual(shape.text.replace("\n", ""), f"{chr(65 + i)}. {text}")
                        self.assertFalse(shape.text_frame.word_wrap)
                        run = shape.text_frame.paragraphs[0].runs[0]
                        self.assertGreaterEqual(run.font.size.pt, 10)
                        correct = spec["mode"] == "answers" and i == question["answer"]
                        self.assertEqual(bool(run.font.bold), correct)
                        self.assertEqual(run.font.color.rgb, ACCENT if correct else GRAY)
                        if page not in (4, 5):
                            positions.setdefault(i, shape.left)
                            self.assertEqual(shape.left, positions[i])
                        if spec["mode"] == "answers":
                            previous = next(s for s in slides[page - 1].shapes if s.name == shape.name)
                            self.assertEqual(shape.left, previous.left)
                    key = f"knowledge:explanation:{q}"
                    self.assertEqual(key in shapes, spec["mode"] == "answers")
                    if key in shapes:
                        self.assertEqual(shapes[key].text.replace("\n", ""), question["explanation"])
                if page in (4, 5):
                    options = [shapes[f"knowledge:option:0:{i}"] for i in range(3)]
                    self.assertIn("\n", options[0].text)
                    self.assertEqual(len({s.left for s in options}), 1)
                    for a, b in zip(options, options[1:]):
                        self.assertGreater(b.top, a.top + a.height)

    def test_standard_gap_font_and_stop(self):
        spec = make_deck()["slides"][1]
        standard, _ = _fit_questions(SimpleNamespace(height=6), spec["questions"], spec["mode"])
        observed = set()
        for height in (6, standard.used + .36 - .05,
                       standard.used + .36 - .30, .5):
            try:
                fit, _ = _fit_questions(SimpleNamespace(height=height), spec["questions"], spec["mode"])
                observed.add(fit.stage)
                self.assertGreaterEqual(fit.values["size"], 11.5)
            except FitError:
                observed.add("stop")
        self.assertEqual(observed, {"standard", "gap", "font", "stop"})

    def test_overcrowded_fields_stop(self):
        for field in ("question", "options", "explanation"):
            spec = deepcopy(make_deck()["slides"][1])
            if field == "options":
                spec["questions"][0]["options"][0] = "確認条件\n" * 80
            else:
                spec["questions"][0][field] = "確認条件\n" * 80
            prs = Presentation()
            with self.assertRaises(FitError):
                s_knowledge_check(prs.slides.add_slide(prs.slide_layouts[6]), spec, 1)


if __name__ == "__main__":
    unittest.main()
