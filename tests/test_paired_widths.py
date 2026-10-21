"""評価軸に応じた欄幅と行の対応・本文色・明示停止を検証する。"""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation

from slidegen.candidate_renderers import _paired_column_widths, s_paired_comparison
from slidegen.check_layout import check
from slidegen.generate import BODY_BOTTOM, BODY_W, TEXT
from slidegen.layout_fit import FitError
from tools.gallery.generate_paired_cases import main, make_deck


class PairedWidthTests(unittest.TestCase):
    def test_short_criteria_return_width_to_equal_body_columns(self):
        specs = make_deck()["slides"]
        left, center, right = _paired_column_widths(specs[0]["rows"])
        self.assertAlmostEqual(left, right)
        self.assertGreater(left, 4.95)
        self.assertLess(center, 1.68)
        self.assertAlmostEqual(left + center + right + .40, BODY_W)
        long = _paired_column_widths(specs[1]["rows"])
        self.assertGreater(long[1], center)
        self.assertLessEqual(long[1], 2.0)
        self.assertAlmostEqual(long[0], long[2])

    def test_saved_content_rows_colors_and_sizes(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "paired.pptx"
            main(output)
            self.assertFalse(check(output))
            for slide, spec in zip(Presentation(output).slides, make_deck()["slides"]):
                shapes = {s.name: s for s in slide.shapes}
                for index, row in enumerate(spec["rows"]):
                    boxes = [shapes[f"paired:{key}:{index}"] for key in ("left", "criterion", "right")]
                    left, criterion, right = boxes
                    self.assertEqual(left.width, right.width)
                    self.assertEqual(left.top, right.top)
                    self.assertEqual(left.height, right.height)
                    self.assertLessEqual((left.top + left.height) / 914400, BODY_BOTTOM - .05)
                    self.assertLess(left.left + left.width, criterion.left)
                    self.assertLess(criterion.left + criterion.width, right.left)
                    for key, box in zip(("left", "criterion", "right"), boxes):
                        self.assertEqual(box.text.replace("\n", ""), row[key])
                        for paragraph in box.text_frame.paragraphs:
                            for run in paragraph.runs:
                                self.assertGreaterEqual(run.font.size.pt, 8.5 if key == "criterion" else 9.5)
                                if key != "criterion":
                                    self.assertEqual(run.font.color.rgb, TEXT)

    def test_overlong_body_or_criterion_stops(self):
        for key in ("left", "criterion", "right"):
            spec = deepcopy(make_deck()["slides"][0])
            spec["rows"][0][key] = "長い確認事項" * 100
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            with self.assertRaises(FitError):
                s_paired_comparison(slide, spec, 1)


if __name__ == "__main__":
    unittest.main()
