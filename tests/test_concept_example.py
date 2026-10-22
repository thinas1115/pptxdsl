"""定義と補足の分担例の情報保持・schema・保存PPTXを検証する。"""
import json
from pathlib import Path
import re
import tempfile
import unittest

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.check_layout import check
from slidegen.validate_content import validate
from tests.fixtures.gallery.content_patterns import PATTERN_DECK
from tests.fixtures.gallery.content_lead_patterns import LEAD_PATTERN_DECK
from slidegen.training_renderers import s_concept


class ConceptExampleTests(unittest.TestCase):
    def test_gallery_and_document_example_preserve_definition_boundaries(self):
        root = Path(__file__).resolve().parents[1]
        guide = (root / "docs/type-selection-guide.md").read_text(encoding="utf-8")
        section = guide.split("### `concept`の定義と補足の分担", 1)[1]
        example = json.loads(re.search(r"```json\s*(.*?)```", section, re.S).group(1))
        cases = [example] + [next(spec for spec in deck["slides"] if spec["type"] == "concept")
                              for deck in (PATTERN_DECK, LEAD_PATTERN_DECK)]
        for spec in cases:
            self.assertLessEqual(len(spec["definition"]), 40)
            self.assertNotIn("Recovery Time Objective", spec["definition"])
            points = {point["label"]: point["text"] for point in spec["points"]}
            self.assertIn("Recovery Time Objective", points["名称"])
            self.assertEqual(set(points), {"名称", "起点", "終点", "設計への反映"})
            self.assertIn("事前に合意する目標値", spec["footnote"]["text"])
            self.assertFalse(validate({"meta": {"title": "分担例"}, "slides": [spec]}, allow_sample_content=True))
            prs = Presentation()
            prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            generate.render_slide(s_concept, slide, spec, 1)
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "concept.pptx"
                prs.save(path)
                self.assertFalse(check(path))
                text = "\n".join(shape.text for shape in Presentation(path).slides[0].shapes
                                 if shape.has_text_frame)
                self.assertIn("Recovery Time Objective", text)
                self.assertIn("誤解しやすい点", text)


if __name__ == "__main__":
    unittest.main()
