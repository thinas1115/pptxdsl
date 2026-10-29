"""フェーズ名とライフライン・要求・応答を分離する契約。"""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from slidegen import generate
from slidegen.generate_from_json import build_presentation
from slidegen.candidate_renderers import _fit_rows, s_sequence
from slidegen.check_layout import check
from slidegen.layout_fit import FitError
from slidegen.quality_markers import SEQUENCE_LIFELINE_PREFIX, SEQUENCE_MESSAGE_LABEL_PREFIX
from slidegen.validate_content import validate
from tools.gallery.generate_sequence_cases import main, make_deck


def render(spec):
    deck = {"meta": {"title": "シーケンス検証"}, "slides": [spec]}
    prs = build_presentation(deck, allow_sample_content=True)
    slide = prs.slides[0]
    return prs, slide


class SequencePhasesTest(unittest.TestCase):
    def test_whole_diagram_is_centered_with_phase_rail_and_rightmost_self_loop(self):
        for count in range(2, 7):
            for with_phases in (False, True):
                for self_loop in (False, True):
                    with self.subTest(count=count, phases=with_phases, self_loop=self_loop):
                        participants = [{"id": f"p{i}", "label": f"参加者{i + 1}"}
                                        for i in range(count)]
                        spec = {"type": "message_sequence", "kicker": "中央配置の検証",
                                "title": "送受信の流れ", "participants": participants,
                                "messages": [
                                    {"id": "m0", "from": "p0", "to": f"p{count - 1}", "label": "要求"},
                                    {"id": "m1", "from": f"p{count - 1}",
                                     "to": f"p{count - 1}" if self_loop else "p0", "label": "確認"}]}
                        if with_phases:
                            spec["phases"] = [{"label": "確認工程", "from": "m0", "to": "m1"}]
                        prs, slide = render(spec)
                        # ヘッダー・凡例・フッターを除いた図全体の描画範囲で判定する。
                        shapes = [shape for shape in slide.shapes
                                  if shape.name.startswith("sequence-")]
                        left = min(shape.left for shape in shapes)
                        right = max(shape.left + shape.width for shape in shapes)
                        self.assertAlmostEqual((left + right) / 2, prs.slide_width / 2, delta=2)
                        self.assertGreaterEqual(left, Inches(generate.MARGIN))
                        self.assertLessEqual(right, Inches(generate.MARGIN + generate.BODY_W))
                        lifelines = [shape for shape in shapes
                                     if shape.name.startswith(SEQUENCE_LIFELINE_PREFIX)]
                        span = max(shape.left for shape in lifelines) - min(shape.left for shape in lifelines)
                        half = 4.18 if count <= 3 else min(4.45, 2.80 + .55 * (count - 2))
                        self.assertAlmostEqual(span / 914400, 2 * half, places=5)

    def test_phase_rail_does_not_cover_lifelines_or_message_labels(self):
        for spec in make_deck()["slides"]:
            with self.subTest(case=spec["kicker"]):
                prs, _ = render(spec)
                with tempfile.TemporaryDirectory() as tmp:
                    saved = Path(tmp) / "case.pptx"
                    prs.save(saved)
                    slide = Presentation(saved).slides[0]
                lines = [sh for sh in slide.shapes if sh.name.startswith(SEQUENCE_LIFELINE_PREFIX)]
                leftmost = min(sh.left for sh in lines)
                labels = [sh for sh in slide.shapes if sh.name.startswith("sequence-phase-label:")]
                self.assertEqual(len(labels), len(spec.get("phases", [])))
                for label in labels:
                    self.assertLessEqual(label.left + label.width, leftmost - Inches(0.17))
                    self.assertGreaterEqual(label.left, Inches(generate.MARGIN))
                    self.assertFalse(label.text_frame.word_wrap)
                    self.assertGreaterEqual(label.text_frame.paragraphs[0].runs[0].font.size, Pt(9.5))
                    for message in slide.shapes:
                        if message.name.startswith(SEQUENCE_MESSAGE_LABEL_PREFIX):
                            self.assertLess(label.left + label.width, message.left)

    def test_request_return_async_and_lifeline_have_distinct_styles(self):
        spec = deepcopy(make_deck()["slides"][1])
        _, slide = render(spec)
        messages = {item["id"]: item for item in spec["messages"]}
        for shape in slide.shapes:
            if shape.name.startswith(SEQUENCE_LIFELINE_PREFIX):
                self.assertEqual(shape.line.color.rgb, generate.RULE)
                self.assertEqual(shape.line.width, Pt(0.6))
                self.assertIsNone(shape.line._get_or_add_ln().find(qn("a:prstDash")))
            elif shape.name.startswith("sequence-message-line:"):
                kind = messages[shape.name.split(":", 1)[1]].get("kind", "request")
                self.assertEqual(shape.line.color.rgb, {
                    "request": generate.NAVY, "return": generate.GRAY, "async": generate.ACCENT}[kind])
                dash = shape.line._get_or_add_ln().find(qn("a:prstDash"))
                self.assertEqual(dash.get("val") if dash is not None else None,
                                 "dash" if kind == "return" else None)
                self.assertEqual(shape.line.width, Pt(1.25))

    def test_dashed_arrows_keep_schema_order_for_powerpoint_color(self):
        spec = deepcopy(make_deck()["slides"][1])
        # 自己応答3辺も同じ契約にする。
        spec["messages"][2]["to"] = spec["messages"][2]["from"]
        spec["messages"][2]["kind"] = "return"
        _, slide = render(spec)
        dashed = []
        for shape in slide.shapes:
            if not hasattr(shape, "line"):
                continue
            line = shape.line._get_or_add_ln()
            dash = line.find(qn("a:prstDash"))
            if dash is not None:
                self.assertLess(list(line).index(line.find(qn("a:solidFill"))), list(line).index(dash))
                dashed.append(shape)
        self.assertGreaterEqual(len(dashed), 4)

    def test_public_and_dense_cases_pass_saved_layout_check(self):
        self.assertEqual(validate(make_deck(), allow_sample_content=True), [])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sequence.pptx"
            main(path)
            self.assertEqual(len(Presentation(path).slides), 8)
            self.assertEqual(check(path), [])

    def test_phase_label_does_not_silently_shrink_or_truncate(self):
        spec = deepcopy(make_deck()["slides"][1])
        spec["phases"][0]["label"] = "確認工程" * 100
        with self.assertRaisesRegex(FitError, "phases"):
            render(spec)

    def test_row_fit_keeps_compression_and_explicit_stop(self):
        stages = set()
        for height in (7.0, 6.0, 5.0, 4.0):
            result = _fit_rows("message_sequence", height, 12, row_h=0.42, min_row_h=0.29,
                               gap=0.12, min_gap=0.02, font=10.5, min_font=8.0)
            stages.add(result.stage)
        self.assertTrue({"standard", "gap", "element"} <= stages)
        with self.assertRaises(FitError):
            _fit_rows("message_sequence", 3.0, 12, row_h=0.42, min_row_h=0.29,
                      gap=0.12, min_gap=0.02, font=10.5, min_font=8.0)


if __name__ == "__main__":
    unittest.main()
