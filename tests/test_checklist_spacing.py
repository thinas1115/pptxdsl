"""保存PPTXで全項目の先頭行合わせ・均等gap・状態非依存を確認する。"""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from slidegen import generate
from slidegen.check_layout import check
from slidegen.layout_fit import FitError
from tools.gallery.generate_checklist_cases import main, make_deck


def positions(slide, spec):
    texts = {shape.text: shape for shape in slide.shapes if shape.has_text_frame}
    bodies = [texts[item["text"]] for item in spec["bullets"]]
    boxes = sorted([shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                    and abs(shape.left.inches - 0.99) < 0.001
                    and abs(shape.width.inches - 0.24) < 0.001
                    and abs(shape.height.inches - 0.24) < 0.001], key=lambda s: s.top)
    return bodies, boxes


class ChecklistSpacingTests(unittest.TestCase):
    def test_all_saved_rows_have_constant_gap_and_first_line_alignment(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "checklist.pptx"
            main(output)
            self.assertFalse(check(output))
            for slide, spec in zip(Presentation(output).slides, make_deck()["slides"]):
                bodies, boxes = positions(slide, spec)
                self.assertEqual(len(boxes), len(bodies))
                gaps = [(b.top - a.top - a.height) / Inches(1) for a, b in zip(bodies, bodies[1:])]
                self.assertLess(max(gaps) - min(gaps), 0.001)
                self.assertGreaterEqual(min(gaps), 0.259)
                for body, box in zip(bodies, boxes):
                    size = body.text_frame.paragraphs[0].runs[0].font.size.pt
                    expected = body.top.inches + size / 108 + 0.012
                    self.assertAlmostEqual(box.top.inches + box.height.inches / 2, expected, delta=0.001)
                    self.assertLessEqual(body.top.inches + body.height.inches, generate.BODY_BOTTOM)
                if len(bodies) == 2:
                    self.assertGreater(generate.BODY_BOTTOM - bodies[-1].top.inches - bodies[-1].height.inches, 1.5)

    def test_toggling_checked_state_does_not_change_geometry(self):
        for original in make_deck()["slides"]:
            geometry = []
            for invert in (False, True):
                spec = deepcopy(original)
                for item in spec["bullets"]:
                    item["checked"] = bool(item.get("checked")) != invert
                prs = Presentation()
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                generate.render_slide(generate.s_bullets, slide, spec, 1)
                bodies, boxes = positions(slide, spec)
                geometry.append([(s.left, s.top, s.width, s.height) for s in [*bodies, *boxes]])
            self.assertEqual(*geometry)

    def test_overcrowded_checklist_stops(self):
        spec = deepcopy(make_deck()["slides"][2])
        for item in spec["bullets"]:
            item["text"] *= 20
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        with self.assertRaises(FitError):
            generate.s_bullets(slide, spec, 1)


if __name__ == "__main__":
    unittest.main()
