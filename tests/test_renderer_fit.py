"""全rendererが過密入力を明示停止することを検証する。"""
from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.aws_vpc_layout import s_aws_vpc_layout
from slidegen.org_layout import s_org
from slidegen.diagrams2 import s_matrix, s_process, s_program_roadmap
from slidegen.layout_fit import FitError, select_fit


LONG = "提出品質を維持できないほど長い説明文です。" * 10


def _slide():
    prs = Presentation()
    prs.slide_width = Inches(generate.SLIDE_W)
    prs.slide_height = Inches(generate.SLIDE_H)
    return prs.slides.add_slide(prs.slide_layouts[6])


def _must_fail(renderer, spec, expected):
    try:
        renderer(_slide(), spec, 1)
    except FitError as e:
        assert expected in str(e), str(e)
    else:
        raise AssertionError(f"{spec['type']}が過密入力を拒否しませんでした")


def _base(type_):
    return {"type": type_, "kicker": "TEST", "title": "収容検証"}


def main():
    # section_divider: leadなし・leadありの標準配置と、実測した開始位置を確認する。
    section = _base("section_divider")
    section.update(kicker="第2章", title="導入計画")
    section_slide = _slide()
    generate.s_section_divider(section_slide, section, 1)
    section_texts = [shape for shape in section_slide.shapes
                     if getattr(shape, "has_text_frame", False)
                     and shape.text]
    assert [shape.text for shape in section_texts[:2]] == ["第2章", "導入計画"]
    assert section_texts[0].top / Inches(1) == generate.SECTION_DIVIDER_TOP
    assert section_texts[1].top / Inches(1) > section_texts[0].top / Inches(1)

    section_with_lead = dict(section, lead="試験導入から展開判断まで")
    lead_slide = _slide()
    generate.s_section_divider(lead_slide, section_with_lead, 1)
    lead_text = next(shape for shape in lead_slide.shapes
                     if getattr(shape, "has_text_frame", False)
                     and shape.text == section_with_lead["lead"])
    assert lead_text.top / Inches(1) > section_texts[1].top / Inches(1)
    assert (lead_text.top + lead_text.height) / Inches(1) \
        <= generate.SECTION_DIVIDER_BOTTOM + 0.01

    wrapped_kicker = dict(section_with_lead,
                          kicker="第2章 導入計画に関する長い章ラベルを自然に改行して表示する")
    wrapped_kicker_slide = _slide()
    generate.s_section_divider(wrapped_kicker_slide, wrapped_kicker, 1)
    wrapped_kicker_shape = next(shape for shape in wrapped_kicker_slide.shapes
                                if getattr(shape, "has_text_frame", False)
                                and shape.text.replace("\n", "")
                                == wrapped_kicker["kicker"])
    assert wrapped_kicker_shape.top / Inches(1) == generate.SECTION_DIVIDER_TOP
    assert (wrapped_kicker_shape.top + wrapped_kicker_shape.height) / Inches(1) \
        <= generate.SECTION_DIVIDER_BOTTOM + 0.01

    # 余白圧縮と文字縮小が標準の後に発動することを確認する。
    gap_spec = dict(section_with_lead,
                    lead="試験導入から展開判断まで。" * 21)
    gap_fit = select_fit(
        "section_divider",
        generate.SECTION_DIVIDER_BOTTOM - generate.SECTION_DIVIDER_TOP,
        generate._section_divider_candidate_values(gap_spec), guidance="x")
    assert gap_fit.stage == "gap", gap_fit
    font_spec = dict(section_with_lead,
                     lead="試験導入から展開判断まで。" * 23)
    font_fit = select_fit(
        "section_divider",
        generate.SECTION_DIVIDER_BOTTOM - generate.SECTION_DIVIDER_TOP,
        generate._section_divider_candidate_values(font_spec), guidance="x")
    assert font_fit.stage == "font", font_fit

    overfull = dict(section_with_lead,
                    lead="ABCDEFGHIJ" * 100)
    _must_fail(generate.s_section_divider, overfull, "不足")

    spec = _base("bullets")
    spec["bullets"] = [[LONG, None] for _ in range(6)]
    _must_fail(generate.s_bullets, spec, "不足")
    try:
        generate.render_slide(generate.s_bullets, _slide(), spec, 1)
    except SystemExit as e:
        assert "slides[0] (type=bullets)" in str(e), str(e)
        assert "箇条書きを減らす" in str(e), str(e)
    else:
        raise AssertionError("rendererエラーにスライド位置が付与されませんでした")

    checklist = _base("bullets")
    checklist.update(
        style="checklist",
        bullets=[{"text": LONG, "checked": i % 2 == 0} for i in range(6)],
    )
    _must_fail(generate.s_bullets, checklist, "不足")

    card_body = "判断に必要な事実と示唆を簡潔に整理する。"
    normal_cards = dict(
        _base("cards"), style="editorial",
        cards=[{"heading": f"項目{i + 1}", "body": card_body} for i in range(6)],
    )
    generate.s_cards(_slide(), normal_cards, 1)

    gap_cards = dict(
        normal_cards,
        cards=[{"heading": f"項目{i + 1}", "body": card_body * 3}
               for i in range(6)],
    )
    gap_slide = _slide()
    generate.s_cards(gap_slide, gap_cards, 1)
    gap_body = next(shape for shape in gap_slide.shapes
                    if getattr(shape, "has_text_frame", False)
                    and shape.text == card_body * 3)
    assert gap_body.text_frame.paragraphs[0].runs[0].font.size.pt == 14

    shrink_cards = dict(
        normal_cards,
        cards=[{"heading": f"項目{i + 1}", "body": card_body * 4}
               for i in range(6)],
    )
    shrink_slide = _slide()
    generate.s_cards(shrink_slide, shrink_cards, 1)
    shrink_body = next(shape for shape in shrink_slide.shapes
                       if getattr(shape, "has_text_frame", False)
                       and shape.text == card_body * 4)
    assert shrink_body.text_frame.paragraphs[0].runs[0].font.size.pt < 14

    overfull_cards = dict(
        normal_cards,
        cards=[{"heading": f"項目{i + 1}", "body": LONG} for i in range(6)],
    )
    _must_fail(generate.s_cards, overfull_cards, "カード本文")

    spec = _base("table")
    spec.update(columns=["項目", "説明"],
                rows=[[f"行{i}", LONG] for i in range(8)])
    _must_fail(generate.s_table, spec, "表の行を減らす")

    panel = {"heading": "比較", "bullets": [LONG for _ in range(6)]}
    spec = _base("two_column")
    spec.update(left=panel, right=panel)
    _must_fail(generate.s_twocol, spec, "左右の箇条書き")

    spec = _base("chart")
    spec["chart"] = {
        "categories": [str(i) for i in range(13)],
        "series": [["実績", [1] * 13]],
    }
    _must_fail(generate.s_chart, spec, "カテゴリ1〜12件")

    spec = _base("process")
    spec.update(steps=[{"name": "工程", "desc": "説明", "actor": "担当"}
                       for _ in range(7)], emph=[])
    _must_fail(s_process, spec, "工程は3〜6件")

    optional_actor = _base("process")
    optional_actor.update(steps=[
        {"name": "工程A", "desc": "担当表示なし"},
        {"name": "工程B", "desc": "担当表示あり", "actor": "担当区分"},
        {"name": "工程C", "desc": "担当表示なし"},
    ])
    s_process(_slide(), optional_actor, 1)

    generic_attribute = _base("process")
    generic_attribute.update(steps=[
        {"name": "入力", "desc": "要件を整理",
         "attribute": {"label": "INPUT", "value": "資料要件"}},
        {"name": "生成", "desc": "資料を生成",
         "attribute": {"label": "OUTPUT", "value": "PPTX"}},
        {"name": "確認", "desc": "品質を確認"},
    ])
    s_process(_slide(), generic_attribute, 1)

    spec = _base("program_roadmap")
    spec.update(periods=["1月", "2月", "3月", "4月"],
                tracks=[{"name": "テーマ", "activities": [
                    {"label": f"作業{j}", "start": 0, "end": 4}
                    for j in range(4)
                ]} for _ in range(6)])
    _must_fail(s_program_roadmap, spec, "最小設定")

    spec = _base("matrix")
    spec.update(x_axis="X", y_axis="Y", target_label="対象",
                points=[{"name": str(i), "x": 0.5, "y": 0.5}
                        for i in range(9)])
    _must_fail(s_matrix, spec, "点は1〜8件")

    spec = _base("org")
    spec["org"] = {
        "nodes": {
            f"level{i}": {"name": f"第{i + 1}階層", "sub": "担当範囲",
                          "members": ["担当A", "担当B"]}
            for i in range(6)
        },
        "levels": [[f"level{i}"] for i in range(6)],
        "edges": [{"from": f"level{i}", "to": f"level{i + 1}"}
                  for i in range(5)],
    }
    _must_fail(s_org, spec, "最小設定")

    dense_aws = _base("aws_vpc_layout")
    dense_aws.update(
        lead="長いリード文" * 20,
        vpc={"label": "VPC", "cidr": "10.0.0.0/16"},
        azs=[
            {"id": f"az{i}", "label": f"AZ-{i}", "subnets": [
                {"id": f"subnet{i}_{j}", "label": "very long private subnet label",
                 "resources": [
                     {"id": f"node{i}_{j}_a", "label": "Application Node A",
                      "icon": "icons/fluent/server.png"},
                     {"id": f"node{i}_{j}_b", "label": "Application Node B",
                      "icon": "icons/fluent/server.png"},
                 ]}
                for j in range(4)
            ]}
            for i in range(3)
        ],
    )
    _must_fail(s_aws_vpc_layout, dense_aws, "subnet")

    print("renderer fit tests passed")


if __name__ == "__main__":
    main()
