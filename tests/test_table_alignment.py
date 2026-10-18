"""表の長文・数値・識別子を保存PPTXで検証する。"""
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation
from pptx.enum.text import PP_ALIGN

from slidegen import generate
from slidegen.check_layout import check
from slidegen.layout_fit import FitError
from tools.gallery.generate_table_cases import main, make_deck


class TableAlignmentTests(unittest.TestCase):
    def test_numeric_notation_and_literal_boundary(self):
        for text in ("0", "85", "1,200", "12,345.50", "-8", "−2.5", "+12.0%", "0.75"):
            self.assertEqual(generate._table_column_alignments(["項目", "数値"], [["1", text]]),
                             [PP_ALIGN.LEFT, PP_ALIGN.RIGHT], text)
        for text in ("001", "00.5", "1,20", "12件", "未測定", "—", "1\n2", " 12", "1e3", "1.2.3"):
            self.assertEqual(generate._table_column_alignments(["項目", "本文"], [["1", text]]),
                             [PP_ALIGN.LEFT, PP_ALIGN.LEFT], text)
        self.assertEqual(generate._table_column_alignments(["項目", "本文"], [["A", "12"], ["B", "未測定"]]),
                         [PP_ALIGN.LEFT, PP_ALIGN.LEFT])

    def test_saved_content_width_and_alignment(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "table.pptx"
            main(path)
            self.assertFalse(check(path))
            for slide, spec in zip(Presentation(path).slides, make_deck()["slides"]):
                table = next(s.table for s in slide.shapes if s.has_table)
                alignments = generate._table_column_alignments(spec["columns"], spec["rows"])
                self.assertAlmostEqual(sum(c.width.inches for c in table.columns), generate.BODY_W, places=4)
                for row_index, values in enumerate([spec["columns"], *spec["rows"]]):
                    for col_index, value in enumerate(values):
                        cell = table.cell(row_index, col_index)
                        self.assertEqual(cell.text, value)
                        for paragraph in cell.text_frame.paragraphs:
                            self.assertEqual(paragraph.alignment, alignments[col_index])
                            self.assertGreaterEqual(paragraph.runs[0].font.size.pt, 10.5)
                if spec["title"] == "受付手順の確認事項":
                    self.assertGreater(table.columns[2].width, table.columns[1].width)

    def test_existing_text_table_keeps_left_alignment_and_widths(self):
        spec = make_deck()["slides"][-1]
        self.assertEqual(generate._table_column_alignments(spec["columns"], spec["rows"]),
                         [PP_ALIGN.LEFT] * len(spec["columns"]))
        self.assertAlmostEqual(sum(generate._auto_table_widths(spec["columns"], spec["rows"])), generate.BODY_W)

    def test_overcrowded_table_stops(self):
        spec = make_deck()["slides"][0]
        spec["rows"] = [["検証", "1", "説明が非常に長く、別ページへの分割が必要です。" * 80] for _ in range(8)]
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        with self.assertRaises(FitError):
            generate.s_table(slide, spec, 1)

    def test_six_rows_with_lead_require_page_split(self):
        spec = make_deck()["slides"][4]
        spec["lead"] = "説明を置く分だけ、表の利用可能な高さが減ります。"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        with self.assertRaises(FitError):
            generate.s_table(slide, spec, 1)


if __name__ == "__main__":
    unittest.main()
