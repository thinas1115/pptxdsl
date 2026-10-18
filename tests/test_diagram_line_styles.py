"""主経路・補助経路・境界の強弱と、破線の保存色・XML順序を検証する。"""
from pathlib import Path
import tempfile
import unittest

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from slidegen import generate
from slidegen.diagram_layout import (AUXILIARY_PATH_PT, MAIN_PATH_PT,
                                     PATH_LABEL_MIN_PT, Layout, _path_label_center, render_diagram)
from slidegen.diagrams import (BOUNDARY_LINE_PT, LINE, add_arrow, arrow_label,
                              container, set_line_dash)
from slidegen.diagrams3 import plain_line
from slidegen.layout_fit import FitError
from slidegen.training_renderers import _plain_connector
from tests.fixtures.gallery.diagram_specs import AWS_MULTIAZ_EXAMPLE
from tools.quality.line_styles import ROWS, verify


class DiagramLineStyleTests(unittest.TestCase):
    def test_dash_color_and_schema_order_after_save(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        color = RGBColor(0x66, 0x6E, 0x70)
        shapes = [add_arrow(slide, 1, 1, 2, 1, color=color, dash="dash", both=True),
                  plain_line(slide, 1, 2, 2, 2, color=color, dash="dash"),
                  _plain_connector(slide, 1, 3, 2, 3, color, dash="dash"),
                  _plain_connector(slide, 1, 4, 2, 4, color, arrow=True, dash="dash")]
        container(slide, 3, 1, 2, 3, "境界", color=color, dash="dash")
        shapes.append(slide.shapes[-2])  # containerが枠の後にラベルを作る。
        for index, shape in enumerate(shapes):
            shape.name = f"dash-test:{index}"
            set_line_dash(shape, "dash")  # 同じ指定を繰り返しても重複しない。
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "dash.pptx"
            prs.save(path)
            for shape in Presentation(path).slides[0].shapes:
                if not shape.name.startswith("dash-test:"):
                    continue
                self.assertEqual(shape.line.color.rgb, color)
                children = [child.tag for child in shape.line._get_or_add_ln()]
                self.assertEqual(children.count(qn("a:prstDash")), 1)
                self.assertLess(children.index(qn("a:solidFill")), children.index(qn("a:prstDash")))
                for tag in ("a:headEnd", "a:tailEnd"):
                    if qn(tag) in children:
                        self.assertLess(children.index(qn("a:prstDash")), children.index(qn(tag)))
                if qn("a:headEnd") in children and qn("a:tailEnd") in children:
                    self.assertLess(children.index(qn("a:headEnd")), children.index(qn("a:tailEnd")))

    def test_dense_diagram_hierarchy_and_label_minimum(self):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        render_diagram(slide, AWS_MULTIAZ_EXAMPLE)
        lines = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.LINE]
        self.assertTrue(any(s.line.width.pt == MAIN_PATH_PT and s.line.color.rgb == generate.NAVY for s in lines))
        self.assertTrue(any(s.line.width.pt == AUXILIARY_PATH_PT and s.line.color.rgb == LINE for s in lines))
        self.assertLess(BOUNDARY_LINE_PT, AUXILIARY_PATH_PT)
        self.assertLess(AUXILIARY_PATH_PT, MAIN_PATH_PT)
        labels = {edge["label"] for edge in AWS_MULTIAZ_EXAMPLE["edges"] if edge.get("label")}
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text in labels:
                for paragraph in shape.text_frame.paragraphs:
                    self.assertGreaterEqual(paragraph.runs[0].font.size.pt, PATH_LABEL_MIN_PT)
            if shape.name == "diagram:path-label":
                l, t = shape.left.inches, shape.top.inches
                r, b = l + shape.width.inches, t + shape.height.inches
                for left, top, right, bottom in Layout(AWS_MULTIAZ_EXAMPLE).cont_rect.values():
                    if t < bottom and b > top:
                        self.assertFalse(any(l < edge < r for edge in (left, right)), shape.text)
                    if l < right and r > left:
                        self.assertFalse(any(t < edge < b for edge in (top, bottom)), shape.text)

    def test_label_clearance_without_shapes_and_explicit_stop(self):
        x, y = _path_label_center(1.01, 2, 0.5, 0.2, [(1, 1, 5, 4)])
        self.assertGreaterEqual(x - 0.25, 1.02)
        self.assertEqual(y, 2)
        with self.assertRaises(FitError):
            _path_label_center(1, 2, 10, 1, [(1, 1, 5, 4)])

    def test_too_long_path_label_stops_instead_of_tiny_text(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        with self.assertRaises(FitError):
            arrow_label(slide, 2, 2, "経路の説明が長すぎるため分割が必要です" * 8,
                        w=1.9, size=10, min_size=9)

    def test_independent_png_check_rejects_wrong_blue(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image = Image.new("RGB", (1600, 900), "white")
            draw = ImageDraw.Draw(image)
            for y, color in ROWS:
                draw.line((700, round(y * 120), 1000, round(y * 120)), fill=tuple(color), width=3)
            image.save(root / "slide_01.png")
            result = verify(root, "other", "人工画素の単体試験")
            self.assertEqual(result["visual_review"], "not_performed")
            draw.line((700, round(ROWS[0][0] * 120), 1000, round(ROWS[0][0] * 120)),
                      fill=(79, 129, 189), width=3)
            image.save(root / "slide_01.png")
            with self.assertRaises(ValueError):
                verify(root, "other", "意図しない青の拒否試験")


if __name__ == "__main__":
    unittest.main()
