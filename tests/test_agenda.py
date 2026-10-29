"""目次の公開契約・段階的収容・保存PPTXを検証する。"""
from copy import deepcopy
import os
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation

from slidegen.agenda_renderer import plan_agenda
from slidegen.check_layout import check
from slidegen.generate_from_json import build_presentation
from slidegen.layout_fit import FitError
from slidegen.validate_content import validate


def spec(count=4, detail=True):
    names = ("目的と範囲", "利用条件", "比較項目", "実施手順", "役割分担", "確認方法", "質疑応答")
    return {"type": "agenda", "kicker": "目次の検証", "title": "説明会の構成",
            "items": [{"title": name, **({"detail": "この章の確認事項です。"} if detail else {})}
                      for name in names[:count]]}


def deck(slides):
    return {"meta": {"title": "目次の検証"}, "slides": slides}


class AgendaTests(unittest.TestCase):
    def test_contract_boundaries(self):
        for count in (1, 7):
            self.assertEqual(validate(deck([spec(count)])), [])
        for items in ([], [{"title": "章名"}] * 8, [{"title": "章名", "page": 0}],
                      [{"title": "章名", "page": True}], [{"title": "章名", "page": 1.5}],
                      [{"title": "章名", "detail": ""}], [{"title": "章名", "x": 1}]):
            with self.subTest(items=items):
                invalid = spec()
                invalid["items"] = items
                self.assertTrue(validate(deck([invalid])))

    def test_fit_stages_and_overfull_stop(self):
        items = spec()["items"]
        standard = plan_agenda(items, 5.27)
        self.assertEqual(standard.stage, "standard")
        gap = plan_agenda(items, standard.used - 0.2)
        self.assertEqual(gap.stage, "gap")
        compact = plan_agenda(items, gap.used - 0.2)
        self.assertEqual(compact.stage, "element")
        self.assertGreaterEqual(compact.values.title_size, 15)
        self.assertGreaterEqual(compact.values.detail_size, 11)
        minimum_height = 0.14 + 4 * (15 * 1.6 / 72 + 0.04 + 11 * 1.6 / 72) + 3 * 0.08
        minimum = plan_agenda(items, minimum_height)
        self.assertEqual(minimum.values.title_size, 15)
        self.assertEqual(minimum.values.detail_size, 11)
        with self.assertRaises(FitError):
            plan_agenda(items, minimum_height - 0.02)
        with self.assertRaises(FitError):
            plan_agenda(items, 0.5)
        with self.assertRaises(FitError):
            plan_agenda([{"title": "非常に長い章名" * 100}], 5.27)
        with self.assertRaises(FitError):
            plan_agenda([{"title": "X" * 200}], 5.27)

    def test_minimum_and_maximum_lead_footnote_and_page_labels(self):
        slides = [spec(1, False), spec(), spec(7), spec(7, False), spec(4)]
        slides[1]["items"][0]["page"] = 2
        slides[-1].update(lead="項目の順番に沿って説明します。", footnote={"text": "会議進行の検証用です。"})
        slides.append(deepcopy(slides[-1]))
        slides[-1]["items"][0]["title"] = "章内の目的と確認事項および検討対象と関連部門で合意する範囲の整理" * 2
        self.assertGreater(len(plan_agenda(slides[-1]["items"], 4.3).values.rows[0].title), 1)
        source = deck(slides)
        expected = deepcopy(source)
        prs = build_presentation(source)
        self.assertEqual(source, expected)
        self.assertIn("02", "\n".join(s.text for s in prs.slides[1].shapes if s.has_text_frame))
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "agenda.pptx"
            prs.save(output)
            self.assertEqual(check(output), [])
            saved = Presentation(output)
            self.assertEqual(len(saved.slides), 6)
        if os.environ.get("PPTXDSL_AGENDA_QA_OUT"):
            output = Path(os.environ["PPTXDSL_AGENDA_QA_OUT"])
            output.parent.mkdir(parents=True, exist_ok=True)
            prs.save(output)


if __name__ == "__main__":
    unittest.main()
