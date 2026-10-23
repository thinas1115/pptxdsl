"""公開ギャラリーの架空数値表示を、通常・lead版と保存PPTXで確認する。"""
from copy import deepcopy
import tempfile
from pathlib import Path
import unittest

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.check_layout import check
from slidegen.validate_content import validate
from tests.fixtures.gallery.content_patterns import PATTERN_DECK
from tests.fixtures.gallery.content_lead_patterns import LEAD_PATTERN_DECK


class GalleryDisclosureTests(unittest.TestCase):
    def test_metrics_disclosure_is_visible_on_both_variants(self):
        for deck in (PATTERN_DECK, LEAD_PATTERN_DECK):
            metrics = [spec for spec in deck["slides"]
                       if spec["type"] == "cards" and spec.get("style") == "metrics"]
            self.assertTrue(metrics)
            for source in metrics:
                spec = deepcopy(source)
                self.assertEqual(spec["footnote"]["label"], "架空データ")
                self.assertIn("実績や測定結果ではありません", spec["footnote"]["text"])
                self.assertEqual([item["value"] for item in spec["cards"]], ["-72%", "8分", "64%"])
                self.assertFalse(validate({"meta": deck["meta"], "slides": [spec]}, allow_sample_content=True))
                prs = Presentation()
                prs.slide_width = Inches(generate.SLIDE_W)
                prs.slide_height = Inches(generate.SLIDE_H)
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                generate.render_slide(generate.s_cards, slide, spec, 1)
                with tempfile.TemporaryDirectory() as directory:
                    output = Path(directory) / "disclosure.pptx"
                    prs.save(output)
                    saved = Presentation(output).slides[0]
                    shapes = {shape.name: shape for shape in saved.shapes}
                    self.assertEqual(shapes["slide-footnote:label"].text, "架空データ")
                    self.assertEqual(shapes["slide-footnote:text"].text, spec["footnote"]["text"])
                    self.assertFalse(check(output))


if __name__ == "__main__":
    unittest.main()
