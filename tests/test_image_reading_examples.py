"""全体・注目領域とcontain/coverを、通常入力から保存したPPTXで検証する。"""
import json
from pathlib import Path
import tempfile
import unittest

from PIL import Image
from pptx import Presentation

from slidegen.asset_paths import resolve_image_path
from slidegen.check_layout import check
from slidegen.generate_from_json import main
from slidegen.image_slide import PICTURE_NAME
from slidegen.validate_content import validate


SOURCE = Path(__file__).resolve().parents[1] / "examples/authoring/image-reading.json"


class ImageReadingTests(unittest.TestCase):
    def test_regular_input_preserves_images_and_visible_range(self):
        deck = json.loads(SOURCE.read_text(encoding="utf-8"))
        self.assertFalse(validate(deck))
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "images.pptx"
            main(SOURCE, output)
            self.assertFalse(check(output))
            pictures = []
            for slide, spec in zip(Presentation(output).slides, deck["slides"]):
                pic = next(s for s in slide.shapes if s.name == PICTURE_NAME)
                pictures.append(pic)
                self.assertEqual(pic.image.blob, resolve_image_path(spec["image"]).read_bytes())
                self.assertEqual(pic._element.nvPicPr.cNvPr.get("descr"), spec["alt"])
                if spec["fit"] == "contain":
                    self.assertEqual((pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom), (0, 0, 0, 0))
            self.assertEqual(pictures[2].image.blob, pictures[3].image.blob)
            self.assertGreater(pictures[3].crop_top, 0)
            self.assertEqual(pictures[3].crop_top, pictures[3].crop_bottom)
            self.assertEqual((pictures[3].crop_left, pictures[3].crop_right), (0, 0))
            # 同じ文字が全体図より大きく見える倍率を保証する。読みやすさは実PNGでも確認する。
            self.assertGreater((pictures[1].width / 1220) / (pictures[0].width / 1600), 1.4)

    def test_fictional_overview_and_focus_assets_have_expected_dimensions(self):
        deck = json.loads(SOURCE.read_text(encoding="utf-8"))
        for spec, size in zip(deck["slides"], ((1600, 900), (1220, 530))):
            with Image.open(resolve_image_path(spec["image"])) as image:
                self.assertEqual(image.size, size)
        self.assertIn("細かな操作文字は次ページ", deck["slides"][0]["lead"])
        self.assertIn("緑枠①", deck["slides"][1]["lead"])
        self.assertIn("端の文字や操作ボタン", deck["slides"][3]["lead"])


if __name__ == "__main__":
    unittest.main()
