"""Mac等の手動書出しを原本保護・自然順序・全件性で検証する。"""
from pathlib import Path
import os
import subprocess
import sys
import tempfile
import unittest

from PIL import Image
from pptx import Presentation

from tools.quality.import_slide_pngs import import_pngs


class ImportSlidePngTests(unittest.TestCase):
    def make_source(self, root, count=2):
        source = root / "export"
        source.mkdir()
        prs = Presentation()
        for number in range(1, count + 1):
            prs.slides.add_slide(prs.slide_layouts[6])
            Image.new("RGB", (1600, 900), (number, 30, 50)).save(source / f"スライド{number}.PNG")
        pptx = root / "deck.pptx"
        prs.save(pptx)
        return pptx, source, root / "png"

    def test_localized_names_preserve_originals_and_order_2_before_10(self):
        with tempfile.TemporaryDirectory() as temporary:
            pptx, source, output = self.make_source(Path(temporary), 10)
            originals = {p.name: p.read_bytes() for p in source.iterdir()}
            self.assertEqual(import_pngs(pptx, source, output), 10)
            self.assertEqual(len(list(output.glob("slide_*.png"))), 10)
            for number in range(1, 11):
                self.assertEqual((output / f"slide_{number:02d}.png").read_bytes(), originals[f"スライド{number}.PNG"])
            self.assertEqual({p.name: p.read_bytes() for p in source.iterdir()}, originals)

    def test_incomplete_duplicate_wrong_size_and_ambiguous_are_rejected(self):
        for case in ("incomplete", "duplicate", "wrong_size", "ambiguous"):
            with self.subTest(case=case), tempfile.TemporaryDirectory() as temporary:
                pptx, source, output = self.make_source(Path(temporary))
                if case == "incomplete":
                    (source / "スライド2.PNG").unlink()
                elif case == "duplicate":
                    (source / "Slide1.png").write_bytes((source / "スライド1.PNG").read_bytes())
                elif case == "wrong_size":
                    Image.new("RGB", (800, 450)).save(source / "スライド2.PNG")
                else:
                    (source / "スライド2.PNG").rename(source / "unknown.png")
                with self.assertRaises(ValueError):
                    import_pngs(pptx, source, output)
                self.assertFalse(output.exists())

    def test_existing_output_is_not_modified(self):
        with tempfile.TemporaryDirectory() as temporary:
            pptx, source, output = self.make_source(Path(temporary))
            output.mkdir()
            preserved = output / "keep.txt"
            preserved.write_text("保持", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "既に存在"):
                import_pngs(pptx, source, output)
            self.assertEqual(preserved.read_text(encoding="utf-8"), "保持")

    def test_cli_invalid_files_do_not_expose_private_paths_or_traceback(self):
        for case in ("missing_pptx", "invalid_pptx", "invalid_png"):
            with self.subTest(case=case), tempfile.TemporaryDirectory() as temporary:
                pptx, source, output = self.make_source(Path(temporary))
                if case == "missing_pptx":
                    pptx = Path(temporary) / "missing.pptx"
                elif case == "invalid_pptx":
                    pptx.write_bytes(b"not a presentation")
                else:
                    (source / "スライド1.PNG").write_bytes(b"not a PNG")
                result = subprocess.run(
                    [sys.executable, "-m", "tools.quality.import_slide_pngs", str(pptx), str(source), str(output)],
                    capture_output=True, text=True, encoding="utf-8",
                    env=dict(os.environ, PYTHONIOENCODING="utf-8"))
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("NG:", result.stderr)
                self.assertNotIn("Traceback", result.stderr)
                self.assertNotIn(temporary, result.stderr + result.stdout)
                self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
