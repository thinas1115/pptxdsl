"""作業名の2行優先・サイズ下限・レーン境界と工程帯の視認性を検証する。"""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation

from slidegen.candidate_renderers import CANVAS, SWIMLANE_RULE, _swimlane_label, s_swimlane
from slidegen.check_layout import check, contrast_ratio
from slidegen.layout_fit import FitError
from slidegen.quality_markers import SURFACE_ON_CANVAS_PREFIX
from tools.gallery.generate_swimlane_cases import main, make_deck


class SwimlaneLabelTests(unittest.TestCase):
    def test_saved_two_line_names_sizes_and_surface_edges(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "swimlane.pptx"
            main(output)
            self.assertFalse(check(output))
            for slide, spec in zip(Presentation(output).slides, make_deck()["slides"]):
                shapes = {s.name: s for s in slide.shapes}
                for step in spec["steps"]:
                    label = shapes[f"swimlane:step-label:{step['id']}"]
                    box = shapes[f"swimlane:step:{step['id']}"]
                    self.assertIn(step["name"], label.text.replace("\n", ""))
                    self.assertLessEqual(len(label.text.splitlines()), 2)
                    self.assertFalse(label.text_frame.word_wrap)
                    self.assertTrue(all(r.font.size.pt >= 9.5 for p in label.text_frame.paragraphs for r in p.runs))
                    self.assertGreaterEqual(label.left, box.left)
                    self.assertLessEqual(label.left + label.width, box.left + box.width)
                if "provider" in {s["id"] for s in spec["steps"]}:
                    label = shapes["swimlane:step-label:provider"]
                    self.assertTrue(all(r.font.size.pt == 11.5 for p in label.text_frame.paragraphs for r in p.runs))
                stages = [s for s in slide.shapes if s.name.startswith(SURFACE_ON_CANVAS_PREFIX)]
                self.assertEqual(len(stages), len(spec["stages"]))
                self.assertGreaterEqual(contrast_ratio(tuple(SWIMLANE_RULE), tuple(CANVAS)), 2.5)
                self.assertTrue(all(s.line.color.rgb == SWIMLANE_RULE for s in stages))
                self.assertEqual(sum(s.name.startswith("swimlane:lane-boundary:") for s in slide.shapes), len(spec["lanes"]))
                backgrounds = [i for i, s in enumerate(slide.shapes) if s.name.startswith("swimlane:lane-background:")]
                boundaries = [i for i, s in enumerate(slide.shapes) if s.name.startswith("swimlane:lane-boundary:")]
                self.assertGreater(min(boundaries), max(backgrounds))
                nodes = [s for s in slide.shapes if s.name.startswith("swimlane:step:")]
                for index, edge in enumerate(spec["edges"]):
                    if edge.get("kind") != "feedback":
                        continue
                    for segment in slide.shapes:
                        if not segment.name.startswith(f"swimlane:edge:{index}:"):
                            continue
                        x1, x2 = sorted((segment.begin_x, segment.end_x))
                        y1, y2 = sorted((segment.begin_y, segment.end_y))
                        for node in nodes:
                            left, top = node.left + 10, node.top + 10
                            right, bottom = node.left + node.width - 10, node.top + node.height - 10
                            crosses = (x1 == x2 and left < x1 < right and max(y1, top) < min(y2, bottom)
                                       or y1 == y2 and top < y1 < bottom and max(x1, left) < min(x2, right))
                            self.assertFalse(crosses, (edge, segment.name, node.name))

    def test_two_lines_precede_font_shrinking(self):
        prs = Presentation()
        label = _swimlane_label(prs.slides.add_slide(prs.slide_layouts[6]),
                                "name", 1, 1, 1.43, .50,
                                "機器・回線調査／復旧対応", 11.5, 9.5)
        self.assertEqual(len(label.text.splitlines()), 2)
        self.assertTrue(all(r.font.size.pt == 11.5 for p in label.text_frame.paragraphs for r in p.runs))

    def test_extreme_step_lane_and_stage_labels_stop(self):
        for field, key in (("steps", "name"), ("lanes", "label"), ("stages", "label")):
            spec = deepcopy(make_deck()["slides"][0])
            spec[field][0][key] = "非常に長い確認事項" * 60
            prs = Presentation()
            with self.assertRaisesRegex(FitError, "分割"):
                s_swimlane(prs.slides.add_slide(prs.slide_layouts[6]), spec, 1)


if __name__ == "__main__":
    unittest.main()
