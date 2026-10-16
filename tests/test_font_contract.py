"""測定faceとPPTX family・styleの契約を通常/異常設定で検証する。"""
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from pptx import Presentation
from pptx.oxml.ns import qn

from slidegen import font_config as config, generate, textfit


class FontContractTests(unittest.TestCase):
    def tearDown(self):
        config._resolve.cache_clear()
        config._code_face.cache_clear()
        config.load_font.cache_clear()

    def test_pptx_uses_measured_family_for_all_scripts(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for bold in (False, True):
            box = generate.add_text(slide, 1, 1, 10, 1, "日本語ABC123", 16, bold=bold)
            run = box.text_frame.paragraphs[0].runs[0]
            family = config.font_faces()["bold" if bold else "regular"].family
            self.assertEqual(run.font.name, family)
            self.assertEqual(run.font.bold, bold)
            for tag in ("a:ea", "a:cs"):
                self.assertEqual(run._r.rPr.find(qn(tag)).get("typeface"), family)
        self.assertIs(textfit._font("medium", 16), textfit._font("regular", 16))
        report = json.dumps(config.describe())
        for face in config.font_faces().values():
            self.assertNotIn(face.path, report)

    def test_unsupported_family_and_wrong_weight_are_rejected(self):
        fake = unittest.mock.Mock()
        fake.get_variation_axes.side_effect = OSError("not variable")
        with patch.object(config.ImageFont, "truetype", return_value=fake):
            fake.getname.return_value = ("Unknown Sans", "Regular")
            with self.assertRaisesRegex(config.FontConfigError, "未対応.*family"):
                config._face("font.ttf", 0, "regular")
            fake.getname.return_value = ("Noto Sans JP", "Regular")
            with self.assertRaisesRegex(config.FontConfigError, "Bold.*style"):
                config._face("font.ttf", 0, "bold")

    def test_fontconfig_does_not_silently_select_another_family(self):
        normal = config.FontFace("normal.ttf", 3, "Yu Gothic", "Regular")
        heavy = config.FontFace("bold.ttf", 3, "Yu Gothic", "Bold")
        with patch.dict(os.environ, {key: "" for key in config._ENV_KEYS}), \
                patch.object(config, "_fontconfig_match", return_value=("font.ttc", 3)), \
                patch.object(config, "_face", side_effect=[normal, heavy, normal, heavy]):
            with self.assertRaisesRegex(config.FontConfigError, "検出できません"):
                config.font_faces()

    def test_index_is_read_from_fontconfig(self):
        completed = subprocess.CompletedProcess([], 0, "font.ttc\t3\n")
        with patch.object(config.shutil, "which", return_value="fc-match"), \
                patch.object(config.subprocess, "run", return_value=completed):
            self.assertEqual(config._fontconfig_match("Noto Sans CJK JP", "Bold"), ("font.ttc", 3))

    def test_config_changes_do_not_reuse_cached_font(self):
        with patch.dict(os.environ, {key: "" for key in config._ENV_KEYS}), \
                tempfile.TemporaryDirectory() as directory:
            regular = Path(directory) / "normal.ttf"
            bold = Path(directory) / "bold.ttf"
            regular.touch()
            bold.touch()
            os.environ.update(PPTXDSL_FONT_REGULAR=str(regular), PPTXDSL_FONT_BOLD=str(bold))
            with patch.object(config, "_face", side_effect=lambda p, i, w: config.FontFace(str(p), i, "Noto Sans JP", w.title())):
                first = config.font_faces()["regular"]
                os.environ["PPTXDSL_FONT_REGULAR_INDEX"] = "2"
                second = config.font_faces()["regular"]
                self.assertEqual((first.index, second.index), (0, 2))
                os.environ["PPTXDSL_FONT_REGULAR_INDEX"] = "-1"
                with self.assertRaisesRegex(config.FontConfigError, "整数"):
                    config.font_faces()

    def test_cli_reports_missing_font_without_private_path(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "content.json"
            source.write_text(json.dumps({"meta": {"title": "フォント検証"}, "slides": [
                {"type": "title", "title": "書体設定", "subtitle": "実行環境の確認"}]}), encoding="utf-8")
            env = dict(os.environ, PYTHONIOENCODING="utf-8", PPTXDSL_FONT_REGULAR=str(Path(directory) / "private.ttf"))
            result = subprocess.run([sys.executable, "-m", "slidegen.generate_from_json", str(source), str(Path(directory) / "out.pptx")],
                                    env=env, capture_output=True, text=True, encoding="utf-8")
            output = result.stdout + result.stderr
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("フォント設定", output)
            self.assertIn("docs/font-contract.md", output)
            self.assertNotIn(directory, output)
            self.assertNotIn("Traceback", output)

    def test_code_uses_the_measured_monospace_family(self):
        from slidegen.training_renderers import _apply_code_font, _fit_code_text
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = generate.add_text(slide, 1, 1, 10, 1, "show vlan brief", 13)
        _apply_code_font(box)
        self.assertEqual(box.text_frame.paragraphs[0].runs[0].font.name, config.font_family("code"))
        self.assertEqual(textfit._font("code", 13).getname()[0], config.font_family("code"))
        if config.font_family("code") == "Consolas":
            with self.assertRaisesRegex(ValueError, "ASCII"):
                _fit_code_text("日本語コメント", 10, 2)

    def test_noto_mixed_script_spacing_scales_with_boundaries(self):
        measured = unittest.mock.Mock()
        measured.getlength.return_value = 96
        face = config.FontFace("font.ttf", 0, "Noto Sans JP", "Regular")
        with patch.object(textfit, "_font", return_value=measured), patch.object(textfit, "font_face", return_value=face):
            plain = textfit.text_width_in("ABCDEFG", 16)
            alternating = textfit.text_width_in("A日B本C語D", 16)
            self.assertAlmostEqual(alternating - plain, 6 * 16 / 72 / 4 * textfit.SAFETY)
            self.assertEqual(textfit.text_width_in("ABC 日本語", 16), plain)

    @unittest.skipUnless(os.environ.get("PPTXDSL_TEST_VARIABLE_FONT"), "可変フォントは専用CIで実ファイルを指定する")
    def test_real_variable_font_selects_400_and_700(self):
        with patch.dict(os.environ, {key: "" for key in config._ENV_KEYS}):
            os.environ["PPTXDSL_FONT_REGULAR"] = os.environ["PPTXDSL_TEST_VARIABLE_FONT"]
            faces = config.font_faces()
            self.assertEqual(faces["regular"].axes, (400,))
            self.assertEqual(faces["bold"].axes, (700,))
            self.assertNotEqual(textfit.text_width_in("WWW iii 012", 24), textfit.text_width_in("WWW iii 012", 24, "bold"))
            os.environ["PPTXDSL_FONT_MEDIUM"] = os.environ["PPTXDSL_FONT_REGULAR"]
            self.assertEqual(config.font_faces()["medium"], faces["regular"])


if __name__ == "__main__":
    unittest.main()
