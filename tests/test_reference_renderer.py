import io
import zipfile

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.layout_fit import FitError
from slidegen.reference_renderer import s_references
from slidegen.validate_content import validate


def _deck(slide):
    return {"meta": {"title": "参照テスト"}, "slides": [slide]}


def _slide(entries):
    return {
        "type": "references",
        "kicker": "参考資料",
        "title": "公式情報源",
        "references": entries,
    }


def _entry(index=1):
    return {
        "id": f"[{index}]",
        "title": f"公式ドキュメント {index}",
        "url": f"https://example.com/docs/{index}",
        "scope": "仕様確認",
    }


def test_references_requires_full_http_url():
    missing = _slide([{"id": "[1]", "title": "情報源"}])
    errors = validate(_deck(missing), allow_sample_content=True)
    assert any("references[0].url は必須" in error for error in errors)

    relative = _slide([{"id": "[1]", "title": "情報源", "url": "docs/source"}])
    errors = validate(_deck(relative), allow_sample_content=True)
    assert any("httpまたはhttpsの完全なURL" in error for error in errors)


def test_references_renders_visible_clickable_urls():
    spec = _slide([_entry(1), _entry(2)])
    assert validate(_deck(spec), allow_sample_content=True) == []
    generate.DECK = _deck(spec)
    prs = Presentation()
    prs.slide_width = Inches(generate.SLIDE_W)
    prs.slide_height = Inches(generate.SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    generate.render_slide(s_references, slide, spec, 1)
    text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
    assert "https://example.com/docs/1" in text
    stream = io.BytesIO()
    prs.save(stream)
    stream.seek(0)
    with zipfile.ZipFile(stream) as archive:
        rels = archive.read("ppt/slides/_rels/slide1.xml.rels").decode("utf-8")
    assert 'Target="https://example.com/docs/1"' in rels
    assert 'TargetMode="External"' in rels


def test_references_rejects_more_than_five_and_stops_on_overflow():
    errors = validate(_deck(_slide([_entry(i) for i in range(1, 7)])),
                      allow_sample_content=True)
    assert any("1〜5件" in error for error in errors)

    dense = _slide([
        dict(_entry(i), title="非常に長い参考資料名" * 30)
        for i in range(1, 6)
    ])
    prs = Presentation()
    prs.slide_width = Inches(generate.SLIDE_W)
    prs.slide_height = Inches(generate.SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    try:
        s_references(slide, dense, 1)
    except FitError:
        pass
    else:
        raise AssertionError("過密な参考資料がFitErrorになりませんでした")


def test_concept_requires_icon():
    spec = {
        "type": "concept", "kicker": "定義", "title": "用語とは",
        "term": "用語", "definition": "意味を説明します。",
        "points": [
            {"label": "観点1", "text": "説明1"},
            {"label": "観点2", "text": "説明2"},
        ],
    }
    errors = validate(_deck(spec), allow_sample_content=True)
    assert any("icon は必須" in error for error in errors)


if __name__ == "__main__":
    test_references_requires_full_http_url()
    test_references_renders_visible_clickable_urls()
    test_references_rejects_more_than_five_and_stops_on_overflow()
    test_concept_requires_icon()
    print("OK: reference renderer tests")
