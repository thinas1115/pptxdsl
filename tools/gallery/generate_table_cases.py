"""長短混在・数値・識別子・疎・8列の表を実描画で比較する。"""
from copy import deepcopy
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.validate_content import validate
from tests.fixtures.gallery.content_patterns import PATTERN_DECK


def make_deck():
    mixed = {"type": "table", "kicker": "長短の混在", "title": "受付手順の確認事項",
             "columns": ["区分", "件数", "確認事項"], "rows": [
                 ["試行前", "3", "通知先の登録と、担当者が不在の場合の引き継ぎ先を確認します。"],
                 ["試行中", "25", "検証データだけを使います。"],
                 ["試行後", "120", "受付から通知までの記録を照合し、重複した受付、通知漏れ、担当者の手戻りを記録して本番導入前の手順へ反映します。"]]}
    numeric = {"type": "table", "kicker": "数値列", "title": "検証用の数値比較",
               "columns": ["項目", "受付件数", "変動率", "処理時間"], "rows": [
                   ["1", "1,200", "-2.5%", "0.75"], ["2", "85", "+12.0%", "12.50"],
                   ["3", "0", "0.0%", "125.25"], ["4", "12,345", "−8.5%", "1.00"]]}
    literal = {"type": "table", "kicker": "数値以外の混在", "title": "測定値と識別子の区別",
               "columns": ["行名", "識別子", "測定値"], "rows": [
                   ["対象A", "001", "12件"], ["対象B", "002", "未測定"], ["対象C", "003", "—"]]}
    dense = {"type": "table", "kicker": "8列の比較", "title": "複数の検証項目",
             "lead": "数値表記だけの列と、説明を含む列の揃え方を分けて確認します。",
             "columns": ["区分", "受付", "完了", "差分", "変動率", "秒数", "版番号", "備考"],
             "rows": [[f"検証{i}", str(i * 100), str(i * 99), str(-i), f"{i}.5%", f"{i}.25", f"00{i}", "照合済み"]
                      for i in range(1, 6)]}
    dense_without_lead = deepcopy(dense)
    dense_without_lead.pop("lead")
    dense_without_lead["title"] = "8列と6行の比較"
    dense_without_lead["rows"].append(["検証6", "600", "594", "-6", "6.5%", "6.25", "006", "照合済み"])
    sparse = deepcopy(numeric)
    sparse["title"] = "少数行の数値比較"
    sparse["rows"] = sparse["rows"][:2]
    existing = deepcopy(next(s for s in PATTERN_DECK["slides"] if s["type"] == "table"))
    return {"meta": {"title": "表の列幅と揃えの回帰検証", "footer": "架空の数値・作業による表示検証"},
            "slides": [mixed, numeric, literal, dense, dense_without_lead, sparse, existing]}


def main(output):
    deck = make_deck()
    errors = validate(deck, allow_sample_content=True)
    if errors:
        raise ValueError("\n".join(errors))
    generate.configure_cover_footer()
    generate.DECK = deck
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
    for page, spec in enumerate(deck["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        generate.render_slide(generate.s_table, slide, spec, page)
        generate.footer(slide, page)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"表の回帰デッキ: {output}（{len(deck['slides'])}ページ）")


if __name__ == "__main__":
    main(sys.argv[1])
