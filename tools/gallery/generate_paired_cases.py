"""短い評価軸・長い評価軸・本文量の異なる対比較を検証する。"""
from copy import deepcopy
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.candidate_renderers import s_paired_comparison
from slidegen.validate_content import validate
from tests.fixtures.gallery.content_patterns import PATTERN_DECK
from tests.fixtures.gallery.content_lead_patterns import LEAD_PATTERN_DECK


def make_deck():
    base = {"type": "paired_comparison", "kicker": "欄幅の比較", "title": "短い評価軸と長い本文",
            "left_label": "方式A", "right_label": "方式B", "rows": [
                {"criterion": "受付", "left": "受付記録を担当窓口で確認し、問い合わせの内容と対応状況をまとめる。",
                 "right": "共通窓口で受付内容を分類し、担当区分と確認結果を記録する。"},
                {"criterion": "記録", "left": "個別の記録様式を使って、連絡先と回答内容を保存する。",
                 "right": "共通の記録様式を使って、受付から回答までの経過を保存する。"},
                {"criterion": "確認", "left": "担当窓口で確認する。", "right": "共通窓口で確認する。"}]}
    long = deepcopy(base)
    long["title"] = "長い評価軸と本文の対応"
    long["lead"] = "中央欄は文言に応じて広げ、左右本文には同じ幅を確保します。優劣を示す検証ではありません。"
    for row, text in zip(long["rows"], ("問い合わせ受付の範囲", "連絡履歴の保管方法", "担当区分の確認手順")):
        row["criterion"] = text
    sparse = deepcopy(base)
    sparse["title"] = "2つの評価軸"
    sparse["rows"] = sparse["rows"][:2]
    dense = deepcopy(base)
    dense["title"] = "6つの評価軸"
    dense["lead"] = "評価軸が増えても、同じ行で左右の本文を読み比べられるか確認します。"
    dense["rows"] = [{"criterion": f"確認軸{i + 1}", "left": "担当窓口で受付内容を確認する。",
                       "right": "共通窓口で受付内容を確認する。"} for i in range(6)]
    specs = [base, long, sparse, dense]
    for deck in (PATTERN_DECK, LEAD_PATTERN_DECK):
        specs.append(deepcopy(next(s for s in deck["slides"] if s["type"] == "paired_comparison")))
    return {"meta": {"title": "対比較の欄幅", "footer": "架空の方式による表示検証"}, "slides": specs}


def main(output):
    deck = make_deck()
    errors = validate(deck, allow_sample_content=True)
    if errors:
        raise ValueError("\n".join(errors))
    generate.configure_cover_footer()
    generate.DECK = deck
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
    for page, spec in enumerate(deck["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        generate.render_slide(s_paired_comparison, slide, spec, page)
        generate.footer(slide, page)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"対比較の回帰デッキ: {output}（{len(deck['slides'])}ページ）")


if __name__ == "__main__":
    main(sys.argv[1])
