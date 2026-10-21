"""スイムレーンの長い作業名・最大構成・工程帯を検証する。"""
from copy import deepcopy
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.candidate_renderers import s_swimlane
from slidegen.validate_content import validate
from tests.fixtures.gallery.candidate_review_cases import REVIEW_DECK
from tests.fixtures.gallery.content_patterns import PATTERN_DECK
from tests.fixtures.gallery.content_lead_patterns import LEAD_PATTERN_DECK


def make_deck():
    normal = deepcopy(next(s for s in PATTERN_DECK["slides"] if s["type"] == "swimlane_flow"))
    normal["title"] = "障害対応の担当と工程"
    normal.pop("lead", None)
    specs = [normal]
    for variant in ("疎", "長文", "上限"):
        spec = deepcopy(next(s for s in REVIEW_DECK["slides"]
                             if s["type"] == "swimlane_flow" and variant in s["kicker"]))
        spec["title"] = {"疎": "2部門の担当フロー", "長文": "長い作業名の担当フロー",
                         "上限": "6部門と6工程の担当フロー"}[variant]
        specs.append(spec)
    same_cell = deepcopy(normal)
    same_cell["title"] = "同じ工程に並ぶ2つの作業"
    same_cell["steps"][2]["name"] = "監視ログ確認"
    same_cell["steps"][5]["name"] = "調査依頼登録"
    specs.append(same_cell)
    specs.append(deepcopy(next(s for s in LEAD_PATTERN_DECK["slides"] if s["type"] == "swimlane_flow")))
    long_lane = deepcopy(normal)
    long_lane["title"] = "長い担当名と作業名"
    long_lane["lanes"][2]["label"] = "ネットワーク運用担当"
    long_lane["lanes"][3]["label"] = "回線事業者の調査窓口"
    long_lane["lead"] = "担当名にも必要な幅を確保し、作業名を2行以内で読み取れるか確認します。"
    specs.append(long_lane)
    return {"meta": {"title": "担当フローの文字", "footer": "架空の業務による表示検証"}, "slides": specs}


def main(output):
    deck = make_deck()
    errors = validate(deck, allow_sample_content=True)
    if errors:
        raise ValueError("\n".join(errors))
    generate.configure_cover_footer()
    generate.DECK = deck
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
    for page, spec in enumerate(deck["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        generate.render_slide(s_swimlane, slide, spec, page)
        generate.footer(slide, page)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"担当フローの回帰デッキ: {output}（{len(deck['slides'])}ページ）")


if __name__ == "__main__":
    main(sys.argv[1])
