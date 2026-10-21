"""完了状態・折り返し・疎密によるチェックリストの配置差を検証する。"""
from copy import deepcopy
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.validate_content import validate
from tests.fixtures.gallery.content_patterns import PATTERN_DECK
from tests.fixtures.gallery.content_lead_patterns import LEAD_PATTERN_DECK


def make_deck():
    short = ["担当者への連絡", "手順書の確認", "記録先の設定", "窓口の案内", "通知経路の試験", "試行開始の判断"]
    long = "試行で扱う検証用データに機密情報が含まれないことを確認し、運用担当と受け渡し方法を合意します。"
    specs = []
    for count, title in ((2, "少数項目の準備状況"), (4, "長短が混在する準備項目"), (6, "最大件数の準備項目")):
        items = [{"text": f"{index + 1}：" + (long if index % 3 == 1 else short[index]),
                  "checked": index % 2 == 0} for index in range(count)]
        spec = {"type": "bullets", "style": "checklist", "kicker": "チェック欄の配置検証",
                "title": title, "bullets": items}
        if count == 6:
            spec["lead"] = "完了状態にかかわらず、項目間隔と本文の先頭行に合わせたチェック欄を保ちます。"
        specs.append(spec)
    for deck in (PATTERN_DECK, LEAD_PATTERN_DECK):
        specs.append(deepcopy(next(s for s in deck["slides"]
                                   if s["type"] == "bullets" and s.get("style") == "checklist")))
    return {"meta": {"title": "チェックリスト配置の回帰検証", "footer": "架空の作業による表示検証"},
            "slides": specs}


def main(output):
    deck = make_deck()
    errors = validate(deck, allow_sample_content=True)
    if errors:
        raise ValueError("\n".join(errors))
    generate.configure_cover_footer()
    generate.DECK = deck
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
    for page, spec in enumerate(deck["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        generate.render_slide(generate.s_bullets, slide, spec, page)
        generate.footer(slide, page)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"チェックリスト回帰デッキ: {output}（{len(deck['slides'])}ページ）")


if __name__ == "__main__":
    main(sys.argv[1])
