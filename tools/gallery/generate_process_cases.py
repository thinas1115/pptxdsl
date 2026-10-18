"""工程補足と分岐条件の対応を実描画で検証する。"""
from copy import deepcopy
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.diagrams2 import s_process
from slidegen.validate_content import validate
from tests.fixtures.gallery.content_patterns import PATTERN_DECK
from tests.fixtures.gallery.content_lead_patterns import LEAD_PATTERN_DECK


def make_deck():
    serial = deepcopy(next(s for s in PATTERN_DECK["slides"] if s["type"] == "process"))
    serial["title"] = "工程説明と補足の対応"
    short = deepcopy(serial)
    short["title"] = "短い説明と担当区分"
    short["emph"] = [2]
    short["steps"] = [{"name": f"工程{i + 1}", "desc": "記録を確認。", "actor": f"担当{i + 1}"}
                      for i in range(3)]
    dense = deepcopy(serial)
    dense["title"] = "最大6工程の補足"
    dense["lead"] = "各工程の説明直後に補足を置き、対応を近い位置で読めるか検証します。"
    dense["steps"] = [{"name": f"工程{i + 1}", "desc": "入力記録と結果を照合し、確認事項を残す。",
                       "attribute": {"label": "確認対象", "value": f"検証記録{i + 1}"}}
                      for i in range(6)]
    plain = deepcopy(short)
    plain["title"] = "補足なしの工程説明"
    for step in plain["steps"]:
        step.pop("actor")
    flow = deepcopy(next(s for s in PATTERN_DECK["slides"] if s.get("flow")))
    flow["title"] = "承認条件と差戻し経路"
    lead = deepcopy(next(s for s in LEAD_PATTERN_DECK["slides"] if s.get("flow")))
    lead["title"] = "説明付きの承認経路"
    up = deepcopy(flow)
    up["title"] = "上向きの分岐条件"
    up["flow"]["levels"] = [["request", "check"], ["review"], ["execute", "revise"]]
    up["flow"]["edges"] = [{"from": "request", "to": "review"},
                              {"from": "check", "to": "review", "label": "確認済"},
                              {"from": "review", "to": "execute", "label": "承認"},
                              {"from": "review", "to": "revise", "label": "要修正"}]
    return {"meta": {"title": "工程補足の回帰検証", "footer": "架空の工程による表示検証"},
            "slides": [serial, short, dense, plain, flow, lead, up]}


def main(output):
    deck = make_deck()
    errors = validate(deck, allow_sample_content=True)
    if errors:
        raise ValueError("\n".join(errors))
    generate.configure_cover_footer()
    generate.DECK = deck
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
    for page, spec in enumerate(deck["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        generate.render_slide(s_process, slide, spec, page)
        generate.footer(slide, page)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"工程補足の回帰デッキ: {output}（{len(deck['slides'])}ページ）")


if __name__ == "__main__":
    main(sys.argv[1])
