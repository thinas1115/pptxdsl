"""基本rendererの回帰サンプルデッキを生成する。"""
import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from tests.fixtures.gallery.content import DECK


def main(out_path, cover_footer_config=None):
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        generate.configure_cover_footer(cover_footer_config)
    except ValueError as e:
        raise SystemExit(f"NG: 表紙・フッター設定: {e}") from e
    generate.DECK = DECK
    prs = Presentation()
    prs.slide_width = Inches(generate.SLIDE_W)
    prs.slide_height = Inches(generate.SLIDE_H)
    blank = prs.slide_layouts[6]
    for idx, spec in enumerate(DECK["slides"], 1):
        slide = prs.slides.add_slide(blank)
        generate.render_slide(generate.RENDER[spec["type"]], slide, spec, idx)
        if spec["type"] != "title":
            generate.footer(slide, idx)
    prs.save(out_path)
    print(f"saved: {out_path} ({len(DECK['slides'])} slides)")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("out_path", nargs="?",
                        default=Path(__file__).resolve().parents[2] / "out" / "sample_basic.pptx")
    parser.add_argument("--cover-footer-config", metavar="PATH",
                        help="表紙・フッター設定JSON")
    args = parser.parse_args()
    main(args.out_path, args.cover_footer_config)
