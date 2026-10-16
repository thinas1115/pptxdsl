"""公開研修のネットワーク図だけを集め、背景・配線・凡例を回帰検証する。"""
import json
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.training_renderers import s_network
from slidegen.validate_content import validate


def make_deck():
    root = Path(__file__).resolve().parents[2]
    source = json.loads((root / "examples/training/vlan/content.json").read_text(encoding="utf-8"))
    return {"meta": {"title": "ネットワーク図の回帰検証", "footer": "配線・背景・凡例の確認"},
            "slides": [spec for spec in source["slides"] if spec["type"] == "nw_topology"]}


def main(output):
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    deck = make_deck()
    errors = validate(deck, allow_sample_content=True)
    if errors:
        raise ValueError("\n".join(errors))
    generate.configure_cover_footer()
    generate.DECK = deck
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
    for index, spec in enumerate(deck["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        generate.render_slide(s_network, slide, spec, index)
        generate.footer(slide, index)
    prs.save(output)
    print(f"ネットワーク回帰デッキ: {output}（{len(deck['slides'])}ページ）")


if __name__ == "__main__":
    main(sys.argv[1])
