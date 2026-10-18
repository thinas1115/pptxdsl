"""プロトコル図の意味表示・最大件数を実描画する回帰デッキ。"""
from copy import deepcopy
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.generate_from_json import RENDER
from slidegen.validate_content import validate
from tests.fixtures.gallery.content_patterns import PATTERN_DECK


def make_deck():
    def sample(type_):
        return deepcopy(next(spec for spec in PATTERN_DECK["slides"] if spec["type"] == type_))

    nat = sample("nw_protocol_flow")
    nat["title"] = "送信元IPの変化と宛先IPの維持"
    minimum = deepcopy(nat)
    minimum.update(title="3段階・1系列の状態表示")
    minimum.pop("lead", None)
    minimum["stages"] = minimum["stages"][:3]
    minimum["flows"] = minimum["flows"][:1]
    minimum["flows"][0]["states"] = minimum["flows"][0]["states"][:3]
    maximum = deepcopy(nat)
    maximum.update(title="6段階・3系列と全マーカーの凡例")
    maximum.pop("lead", None)
    maximum["stages"].append({"id": "archive", "label": "記録先",
                               "icon": "icons/fluent/database.png"})
    maximum["flows"] = []
    for row in range(3):
        states = []
        for stage, appearance in zip(maximum["stages"],
                                     ("plain", "encapsulated", "internal", "alert", "plain", "plain")):
            state = {"stage": stage["id"], "label": "同じ値", "appearance": appearance}
            if appearance == "encapsulated":
                state["encapsulation"] = "HDR"
            if appearance == "alert":
                state.update(label="不一致", detail="確認対象")
            states.append(state)
        maximum["flows"].append({"label": f"属性{row + 1}", "states": states})

    anatomy = sample("nw_frame_anatomy")
    anatomy.update(title="非比例縮尺とフィールド注釈")
    empty = deepcopy(anatomy)
    empty.update(title="注釈なしの非比例縮尺表示")
    empty.pop("lead", None)
    empty["frames"][0].pop("annotations", None)
    dense = deepcopy(anatomy)
    dense.update(title="3フレームと12件の注釈")
    dense.pop("lead", None)
    template = dense["frames"][0]
    template["annotations"] = [{"field": field["id"], "text": "対応フィールドの説明"}
                                for field in template["fields"][:4]]
    dense["frames"] = [deepcopy(template) for _ in range(3)]
    for index, frame in enumerate(dense["frames"]):
        frame["label"] = f"フレーム{index + 1}"
    return {"meta": {"title": "プロトコル図の回帰検証", "footer": "凡例・値の変化・概念縮尺の検証"},
            "slides": [nat, minimum, maximum, anatomy, empty, dense]}


def main(output):
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    deck = make_deck()
    errors = validate(deck, allow_sample_content=True)
    if errors:
        raise ValueError("\n".join(errors))
    generate.configure_cover_footer()
    generate.DECK = deck
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
    for index, spec in enumerate(deck["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        generate.render_slide(RENDER[spec["type"]], slide, spec, index)
        generate.footer(slide, index)
    prs.save(output)
    print(f"プロトコル回帰デッキ: {output}（{len(deck['slides'])}ページ）")


if __name__ == "__main__":
    main(sys.argv[1])
