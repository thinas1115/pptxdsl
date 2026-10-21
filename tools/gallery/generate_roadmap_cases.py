"""12か月の多段工程と少期間・長名の可読性を実描画で検証する。"""
from copy import deepcopy
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.diagrams2 import s_program_roadmap
from slidegen.validate_content import validate
from tests.fixtures.gallery.content_patterns import PATTERN_DECK
from tests.fixtures.gallery.content_lead_patterns import LEAD_PATTERN_DECK
from tests.fixtures.gallery.content_stress_patterns import PROGRAM_ROADMAP_STRESS


def make_deck():
    normal = [deepcopy(s) for s in PATTERN_DECK["slides"] if s["type"] == "program_roadmap"]
    normal[0]["title"] = "12か月の多段工程"
    normal[1]["title"] = "複数テーマの年間工程"
    short = {"type": "program_roadmap", "kicker": "少期間", "title": "3期間の実施工程",
             "periods": ["第1期", "第2期", "第3期"], "tracks": [
                 {"name": "運用手順の確認", "goal": "受付から記録までを照合", "activities": [
                     {"label": "受付手順の確認", "start": 0, "end": 1},
                     {"label": "記録内容の照合", "start": 1, "end": 2},
                     {"label": "結果の取りまとめ", "start": 2, "end": 3}],
                  "milestone": {"at": 3, "label": "確認完了"}},
                 {"name": "連絡先の整理", "activities": [
                     {"label": "連絡先の確認と共有", "start": 0, "end": 2.5, "emph": True}]}]}
    long = deepcopy(short)
    long["title"] = "長い工程名と完了条件"
    long["tracks"] = [{"name": "問い合わせ受付の試行", "goal": "手順と記録様式を整備", "activities": [
        {"label": "問い合わせ受付手順の整備と試行", "start": 0, "end": 3}],
        "milestone": {"at": 3, "label": "試行結果の確認完了"}}]
    dense = deepcopy(PROGRAM_ROADMAP_STRESS)
    dense["title"] = "5テーマと15作業の期間"
    lead = [deepcopy(s) for s in LEAD_PATTERN_DECK["slides"] if s["type"] == "program_roadmap"]
    for index, spec in enumerate(lead, 1):
        spec["title"] = f"説明付きの年間工程{index}"
    return {"meta": {"title": "工程表の回帰検証", "footer": "架空の工程による表示検証"},
            "slides": normal + [short, long, dense] + lead}


def main(output):
    deck = make_deck()
    errors = validate(deck, allow_sample_content=True)
    if errors:
        raise ValueError("\n".join(errors))
    generate.configure_cover_footer()
    generate.DECK = deck
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
    for page, spec in enumerate(deck["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        generate.render_slide(s_program_roadmap, slide, spec, page)
        generate.footer(slide, page)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"工程表の回帰デッキ: {output}（{len(deck['slides'])}ページ）")


if __name__ == "__main__":
    main(sys.argv[1])
