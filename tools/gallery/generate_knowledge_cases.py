"""選択肢の整列・縦並びと、同じ内容の問題/解答ギャラリー。"""
from copy import deepcopy
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.training_renderers import s_knowledge_check
from slidegen.validate_content import validate
from tests.fixtures.gallery.content_patterns import PATTERN_DECK


def make_deck():
    normal = deepcopy(next(s for s in PATTERN_DECK["slides"] if s["type"] == "knowledge_check"))
    normal["title"] = "選択肢の列位置と正答の対応"
    four = deepcopy(normal)
    four["title"] = "2〜4選択肢の共通列"
    for i, question in enumerate(four["questions"]):
        question["question"] = ["投入後の実測値を期待値と比べる作業はどれか。",
                                "後から確認経緯をたどれるようにする作業はどれか。",
                                "判断権限を持つ人から実施の許可を得る作業はどれか。"][i]
        question["options"] = ["初期確認", "状態の照合", "記録の保存", "承認"][:i + 2]
        question["answer"] = i + 1
        question["explanation"] = ["状態の照合で、期待値と実測値が一致するか確認します。",
                                   "記録を保存し、確認した内容と判断の経緯を残します。",
                                   "実施前に判断権限を持つ人から承認を得ます。"][i]
    long = deepcopy(normal)
    long["title"] = "長い選択肢の縦並び"
    long["questions"] = [deepcopy(normal["questions"][0])]
    long["questions"][0]["options"] = [
        "利用者への影響と必要な検知条件を整理してから通知先と判断基準を定める。実際の運用で必要な対応時間と確認方法も合わせて整理し、記録を残す",
        "通知メールの見た目だけを決めて全件を同じ宛先へ送る",
        "担当者の座席順だけで通知の優先順位を決める"]
    lead = deepcopy(normal)
    lead["title"] = "説明付きの問題と解答"
    lead["lead"] = "問題と解答で選択肢の順序と列を保ち、正答と解説の対応を確認します。"
    slides = []
    for source in (normal, four, long, lead):
        for mode in ("questions", "answers"):
            spec = deepcopy(source)
            spec["mode"] = mode
            spec["kicker"] = "問題" if mode == "questions" else "正答・解説"
            slides.append(spec)
    return {"meta": {"title": "問題・解答ギャラリー", "footer": "架空の設問による表示検証"},
            "slides": slides}


def main(output):
    deck = make_deck()
    errors = validate(deck, allow_sample_content=True)
    if errors:
        raise ValueError("\n".join(errors))
    generate.configure_cover_footer()
    generate.DECK = deck
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
    for page, spec in enumerate(deck["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        generate.render_slide(s_knowledge_check, slide, spec, page)
        generate.footer(slide, page)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"問題・解答ギャラリー: {output}（{len(deck['slides'])}ページ）")


if __name__ == "__main__":
    main(sys.argv[1])
