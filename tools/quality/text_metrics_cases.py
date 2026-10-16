"""文字矩形の真値付きPPTXと、同一入力のOS間クロス検査記録を作る。"""
import argparse
from dataclasses import asdict
import hashlib
import json
from pathlib import Path
import platform
import urllib.request

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from slidegen import generate as g
from slidegen.check_layout import check, rect_of, table_cells
from slidegen.font_config import describe
from slidegen.text_metrics import layout_text_frame


FONT_URL = "https://raw.githubusercontent.com/google/fonts/66a36c8c94b1a5d992ee4e7f392fccfe4945767c/ofl/notosansjp/NotoSansJP%5Bwght%5D.ttf"
FONT_SHA256 = "c2f3b4d463500a2ddcd3849cded1fceeb9fd6d1c32e6cbecd568453ba50fc68f"


def prepare_font(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        with urllib.request.urlopen(FONT_URL, timeout=60) as response:
            data = response.read()
        if hashlib.sha256(data).hexdigest() != FONT_SHA256:
            raise ValueError("固定フォントのSHA256が一致しません")
        path.write_bytes(data)
    if hashlib.sha256(path.read_bytes()).hexdigest() != FONT_SHA256:
        raise ValueError("固定フォントのSHA256が一致しません")


def build(output):
    output.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(g.SLIDE_W), Inches(g.SLIDE_H)

    def page(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        g.add_rect(slide, 0, 0, g.SLIDE_W, g.SLIDE_H, g.CANVAS)
        g.add_text(slide, .72, .27, 11, .4, title, 14, bold=True)
        return slide

    slide = page("役割別の折り返し")
    for name, y, size, role, bold, spacing in (
        ("見出し", 1, 20, "natural", True, 1.12), ("lead", 3, 14, "body", False, 1.18),
        ("本文", 5, 12, "body", False, 1.3)):
        shape = g.add_text(slide, 1, y, 4, 1.5, "処理結果の確認方法と日本語APIの実行条件", size,
                           role=role, bold=bold, spacing=spacing)
        shape.name = "text-case:" + name
    slide = page("runごとの書式と段落配置")
    shape = g.add_text(slide, 1, 1, 9, 3, "small ", 9, spacing=1.7, align=PP_ALIGN.CENTER)
    shape.name = "text-case:mixed"
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "BIG"
    g.set_run(run, 28, bold=True)
    p = shape.text_frame.add_paragraph()
    p.alignment, p.line_spacing, p.space_before, p.space_after = PP_ALIGN.RIGHT, Pt(18), Pt(3), Pt(2)
    run = p.add_run()
    run.text = "段落ごとの右寄せ"
    g.set_run(run, 10)
    slide = page("表セルと余白・縦位置")
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(10), Inches(2))
    for ri, row in enumerate(table.table.rows):
        for ci, cell in enumerate(row.cells):
            g._cell(cell, f"セル{ri + 1}・{ci + 1}の内容", 12, center=ci == 1)
    for x, anchor in ((1, MSO_ANCHOR.MIDDLE), (7, MSO_ANCHOR.BOTTOM)):
        shape = g.add_text(slide, x, 4, 4, 2, "余白を除いた本文領域", 14, anchor=anchor)
        shape.text_frame.margin_left = Inches(.2)
        shape.text_frame.margin_right = Inches(.3)
        shape.text_frame.margin_top = Inches(.2)
        shape.text_frame.margin_bottom = Inches(.3)
    slide = page("手動改行と空段落")
    shape = g.add_text(slide, 1, 1, 6, 4, "A", 12, wrap=False, spacing=1.2)
    p = shape.text_frame.paragraphs[0]
    p.text = "A\vB"
    for run in p.runs:
        g.set_run(run, 12)
    shape.text_frame.add_paragraph().font.size = Pt(22)
    p = shape.text_frame.add_paragraph()
    p.line_spacing = 1.0
    run = p.add_run()
    run.text = "空段落の後の本文"
    g.set_run(run, 12)
    slide = page("既知のNG入力：はみ出しと行同士の衝突")
    shape = g.add_text(slide, 1, 1, .4, .4, "LONG_IDENTIFIER", 18, wrap=False)
    shape.name = "text-case:overflow"
    shape = g.add_text(slide, 1, 4, 4, 1, "衝突\n衝突", 24)
    shape.name = "text-case:collision"
    for p in shape.text_frame.paragraphs:
        p.line_spacing = Pt(1)
    path = output / "cases.pptx"
    prs.save(path)
    result = report(path)
    if any(finding[0] != 5 for finding in result["findings"]):
        raise ValueError("正常ページで誤検知しています")
    if not {"TEXT-OOB", "T-T"} <= {item[1] for item in result["findings"]}:
        raise ValueError("既知のNGを拒否できません")
    _write(output / "report.json", result)


def report(path):
    prs = Presentation(path)
    frames = []
    for si, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_table:
                sources = [(f"{shape.name}[{ri},{ci}]", frame, inner, True, "compact")
                           for _, inner, frame, ri, ci in table_cells(shape) if frame.text.strip()]
            elif shape.has_text_frame and shape.text_frame.text.strip():
                sources = [(shape.name, shape.text_frame, rect_of(shape), False, "body")]
            else:
                sources = []
            for name, frame, box, inner, default in sources:
                layout = layout_text_frame(frame, box, box_is_inner=inner, default_role=default)
                frames.append({"slide": si, "name": name, **asdict(layout)})
    return {"deck": path.name, "fonts": describe(), "frames": frames,
            "findings": sorted(set(check(path))),
            "environment": {"os": platform.system(), "python": platform.python_version()}}


def _write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def compare(path, baseline, output):
    expected = json.loads(baseline.read_text(encoding="utf-8"))
    actual = report(path)
    differences, max_delta = [], 0.0

    def visit(a, b, location):
        nonlocal max_delta
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            delta = abs(a - b)
            max_delta = max(max_delta, delta)
            if delta > 0.00001:
                differences.append(location)
        elif isinstance(a, dict) and isinstance(b, dict) and a.keys() == b.keys():
            for key in a:
                visit(a[key], b[key], location + "." + key)
        elif isinstance(a, (list, tuple)) and isinstance(b, (list, tuple)) and len(a) == len(b):
            for index, (left, right) in enumerate(zip(a, b)):
                visit(left, right, f"{location}[{index}]")
        elif a != b:
            differences.append(location)

    for key in ("fonts", "frames", "findings"):
        visit(expected[key], actual[key], key)
    result = {"deck": path.name, "origin": expected["environment"], "checked_on": actual["environment"],
              "max_numeric_delta": max_delta, "differences": differences,
              "frames": len(actual["frames"]), "findings": actual["findings"], "fonts": actual["fonts"]}
    _write(output, result)
    if differences:
        raise ValueError("同一PPTXの検査結果に差があります: " + ", ".join(differences[:10]))
    print(f"同一PPTXの文字検査一致: {path.name} / {len(actual['frames'])}枠 / 最大差{max_delta:.8f}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("prepare-font").add_argument("path", type=Path)
    sub.add_parser("build").add_argument("output", type=Path)
    command = sub.add_parser("compare")
    command.add_argument("pptx", type=Path)
    command.add_argument("baseline", type=Path)
    command.add_argument("output", type=Path)
    args = parser.parse_args()
    if args.command == "prepare-font":
        prepare_font(args.path)
    elif args.command == "build":
        build(args.output)
    else:
        compare(args.pptx, args.baseline, args.output)
