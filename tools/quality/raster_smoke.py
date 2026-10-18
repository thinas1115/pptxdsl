"""実Office描画で補助検査の正常文字・枠外文字・線欠けを確認するデッキ。"""
import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from slidegen.diagrams3 import plain_line
from slidegen.generate import NAVY, add_text
from slidegen.raster_qa import inspect


def build(root):
    root.mkdir(parents=True, exist_ok=True)
    for case in ("good", "text_oob", "line_gap"):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(4), Inches(2.25)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_text(slide, 0.4, 0.6, 0.35 if case == "text_oob" else 2.0, 0.45,
                 "Rendering QA", 18, color=NAVY, wrap=False)
        if case == "line_gap":
            for x0, x1 in ((0.4, 1.2), (2.0, 3.0)):
                plain_line(slide, x0, 1.4, x1, 1.4, color=NAVY, width=2)
        else:
            plain_line(slide, 0.4, 1.4, 3.0, 1.4, color=NAVY, width=2)
        prs.save(root / f"{case}.pptx")


def verify(root, backend, version):
    results = {}
    for case in ("good", "text_oob", "line_gap"):
        spec = {"version": 1, "backend": {"name": backend, "version": version}, "pages": [
            {"png": "slide_01.png", "size": [800, 450], "checks": [
                {"id": "text", "kind": "ink_bounds", "region": [20, 80, 760, 240],
                 "bounds": [80, 120, 150 if case == "text_oob" else 480, 210], "ink": "182C43"},
                {"id": "connection", "kind": "ink_presence", "region": [300, 276, 310, 285],
                 "ink": "182C43", "min_pixels": 10}]}]}
        result = inspect(spec, root / case)
        results[case] = result
        codes = {item["code"] for item in result["findings"]}
        expected = set() if case == "good" else {"PNG-INK-OOB"} if case == "text_oob" else {"PNG-INK-MISSING"}
        if codes != expected or result["visual_review"] != "not_performed":
            raise ValueError(f"実描画の既知NG検出に失敗しました: {case}")
    (root / "raster-report.json").write_text(json.dumps(results, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("実PNG補助検査: 正常合格・枠外文字拒否・線欠け拒否を確認")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("build", "verify"))
    parser.add_argument("root", type=Path)
    parser.add_argument("--backend", choices=("powerpoint", "libreoffice", "other"))
    parser.add_argument("--version")
    args = parser.parse_args()
    if args.command == "build":
        build(args.root)
    else:
        verify(args.root, args.backend, args.version)
