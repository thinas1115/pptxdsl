"""校正デッキのPowerPoint/Poppler描画幅をPillow実測と照合する。"""
import argparse
import json
from pathlib import Path
import re
import subprocess
import xml.etree.ElementTree as ET

from pptx import Presentation
from pptx.util import Inches

from slidegen.font_config import describe
from slidegen.generate import SLIDE_W, SLIDE_H, add_text
from slidegen.textfit import SAFETY, _font, text_width_in, width_safety


def build(output):
    output.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    cases = []
    for page, size in enumerate((8.2, 9, 10.5, 11, 13.5, 16, 24, 42), 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_text(slide, 0.7, 0.2, 11, 0.4, f"書体の幅測定 / {size}pt", 14, bold=True)
        for row, (weight, text) in enumerate((
            (weight, text) for weight in ("regular", "bold")
            for text in ("日本語の測定結果", "WWW iii 012.3", "申請API: 123件", "A日B本C語D")
        )):
            top = 0.85 + row * 0.70
            name = f"font-metric:{page}:{row}"
            shape = add_text(slide, 0.7, top, 11.9, 0.70, text, size, bold=weight == "bold", wrap=False, spacing=1.0)
            shape.name = name
            cases.append({"id": name, "page": page, "top_pt": top * 72,
                          "bottom_pt": (top + 0.70) * 72, "text": text,
                          "size": size, "weight": weight,
                          "safety": width_safety(weight),
                          "pillow_width_pt": _font(weight, size).getlength(text) * 72 / 96,
                          "estimated_width_pt": text_width_in(text, size, weight) / width_safety(weight) * 72})
    prs.save(output / "metrics.pptx")
    (output / "expected.json").write_text(json.dumps({"fonts": describe(), "cases": cases}, ensure_ascii=False, indent=2), encoding="utf-8")


def compare(expected_path, observed, backend):
    expected = json.loads(expected_path.read_text(encoding="utf-8"))
    if backend == "powerpoint":
        actual = {item["id"]: item["width_pt"] for item in json.loads(observed.read_text(encoding="utf-8-sig"))}
    else:
        pages = ET.parse(observed).getroot().findall(".//{*}page")
        actual = {}
        for case in expected["cases"]:
            words = [word for word in pages[case["page"] - 1].findall(".//{*}word")
                     if case["top_pt"] <= (float(word.get("yMin")) + float(word.get("yMax"))) / 2 < case["bottom_pt"]]
            found = "".join("".join("".join(word.itertext()).split()) for word in words)
            if found != "".join(case["text"].split()) or not words:
                raise ValueError(f"実描画の文字列を照合できません: {case['id']}")
            actual[case["id"]] = max(float(w.get("xMax")) for w in words) - min(float(w.get("xMin")) for w in words)
    limit = SAFETY
    results = [{**case, "safety": case.get("safety", limit), "actual_width_pt": actual[case["id"]],
                "raw_ratio": actual[case["id"]] / case["pillow_width_pt"],
                "ratio": actual[case["id"]] / case.get("estimated_width_pt", case["pillow_width_pt"])} for case in expected["cases"]]
    report = {"backend": backend, "fonts": expected["fonts"], "safety": limit,
              "max_ratio": max(row["ratio"] for row in results), "cases": results}
    if any(row["ratio"] > row["safety"] or row["ratio"] <= 0 for row in results):
        raise ValueError("描画幅が安全係数の範囲外です")
    return report


def pdf_fonts(pdf, family):
    result = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True, check=True)
    names = [line.split()[0].split("+")[-1] for line in result.stdout.splitlines()[2:] if line.strip()]
    normalize = lambda value: re.sub(r"[^a-z0-9]", "", value.lower())
    if not names or any(not normalize(name).startswith(normalize(family)) for name in names):
        raise ValueError("校正PDFに代替書体が使われています")
    return names


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    commands.add_parser("build").add_argument("output", type=Path)
    check = commands.add_parser("compare")
    check.add_argument("expected", type=Path)
    check.add_argument("observed", type=Path)
    check.add_argument("--backend", choices=("powerpoint", "poppler"), required=True)
    check.add_argument("--report", type=Path, required=True)
    check.add_argument("--pdf", type=Path, help="Poppler測定時に埋込familyも検証するPDF")
    args = parser.parse_args()
    if args.command == "build":
        build(args.output)
    else:
        report = compare(args.expected, args.observed, args.backend)
        if args.pdf:
            report["pdf_fonts"] = pdf_fonts(args.pdf, report["fonts"]["regular"]["family"])
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"幅測定: {len(report['cases'])}条件、最大比 {report['max_ratio']:.4f} <= {report['safety']}")


if __name__ == "__main__":
    main()
