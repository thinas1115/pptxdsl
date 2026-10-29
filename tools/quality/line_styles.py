"""破線の指定色が実OfficeのPNGへ反映されることを画素で確認する。"""
import argparse
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from slidegen import generate
from slidegen.diagrams import LINE, add_arrow, container
from slidegen.diagrams3 import plain_line
from slidegen.raster_qa import inspect
from slidegen.training_renderers import _plain_connector


ROWS = [(2.4, LINE), (3.1, LINE), (3.8, LINE), (4.5, generate.ACCENT), (5.2, LINE)]


def build(output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(generate.SLIDE_W), Inches(generate.SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    generate.header(slide, "共有部品の実描画", "破線と指定色の確認")
    for (y, _), text in zip(ROWS, ("矢印", "双方向の矢印", "補助線", "研修用の矢印", "配置境界")):
        generate.add_text(slide, 1, y - 0.12, 2.5, 0.4, text, 16, color=generate.NAVY)
    add_arrow(slide, 4, ROWS[0][0], 11, ROWS[0][0], color=LINE, dash="dash")
    add_arrow(slide, 4, ROWS[1][0], 11, ROWS[1][0], color=LINE, dash="dash", both=True)
    plain_line(slide, 4, ROWS[2][0], 11, ROWS[2][0], color=LINE, dash="dash")
    _plain_connector(slide, 4, ROWS[3][0], 11, ROWS[3][0], generate.ACCENT, arrow=True, dash="dash")
    container(slide, 4, ROWS[4][0], 7, 0.6, "境界", color=LINE, dash="dash")
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def verify(png_dir, backend, version):
    with Image.open(png_dir / "slide_01.png") as image:
        width, height = image.size
    checks = []
    for index, (y, color) in enumerate(ROWS):
        checks.append({"id": f"line-{index + 1}", "kind": "ink_presence",
                       "region": [round(6 / generate.SLIDE_W * width),
                                  round((y - 0.04) / generate.SLIDE_H * height),
                                  round(8 / generate.SLIDE_W * width),
                                  round((y + 0.04) / generate.SLIDE_H * height)],
                       "ink": str(color), "tolerance": 24, "min_pixels": 10})
    manifest = {"version": 1, "backend": {"name": backend, "version": version},
                "pages": [{"png": "slide_01.png", "size": [width, height], "checks": checks}]}
    result = inspect(manifest, png_dir)
    (png_dir / "color-report.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if result["automatic_status"] != "passed":
        raise ValueError("実PNGで指定した破線色を確認できません。color-report.jsonを確認してください。")
    print("実PNGの5領域で破線の指定色を確認（目視・人の承認とは別の補助検査）")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("build", "verify"))
    parser.add_argument("target", type=Path)
    parser.add_argument("--backend", choices=("powerpoint", "libreoffice", "other"))
    parser.add_argument("--version")
    args = parser.parse_args()
    if args.command == "build":
        build(args.target)
    else:
        verify(args.target, args.backend, args.version)
